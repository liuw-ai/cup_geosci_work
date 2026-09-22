"""Controlled processing for official recruitment announcement attachments.

This module deliberately stops at a private review queue.  A PDF/XLSX row is
not a public job until an administrator has inspected the official announcement
and explicitly verified the candidate.  Network access is limited to a
registered source, its declared attachment hosts, and robots-permitted public
URLs; no login, CAPTCHA, proxy or browser fingerprint workaround is attempted.
"""

from __future__ import annotations

import csv
import hashlib
import io
import mimetypes
import os
import re
import subprocess
import tempfile
from dataclasses import dataclass
from datetime import datetime, timezone
from pathlib import Path
from urllib.parse import urlparse
from urllib.robotparser import RobotFileParser

import requests
from bs4 import BeautifulSoup

from job_hub.config import Settings
from job_hub.contracts import is_http_url
from job_hub.db import Database
from job_hub.matching import (
    extract_deadline,
    extract_degree_levels,
    extract_major_tags,
    extract_published_date,
    score_relevance,
)
from job_hub.transport import configure_session, create_session


USER_AGENT = "cupb-geoscience-job-hub-attachment-fetcher/0.1 (+official-public-source)"
PARSER_VERSION = "attachments-v1"

SUPPORTED_SUFFIXES = {
    ".pdf",
    ".csv",
    ".xls",
    ".xlsx",
    ".xlsm",
    ".docx",
}
HTML_MARKERS = (b"<html", b"<!doctype html", b"<head", b"<script")


class AttachmentProcessingError(RuntimeError):
    """A compliant attachment could not be downloaded or parsed."""


class AttachmentSkipped(AttachmentProcessingError):
    """The source is public but cannot be processed under current policy."""


@dataclass(frozen=True)
class AttachmentResult:
    artifact_id: int
    status: str
    bytes_downloaded: int = 0
    rows_extracted: int = 0
    candidates_created: int = 0
    detail: str = ""

    def as_dict(self) -> dict[str, object]:
        return {
            "artifact_id": self.artifact_id,
            "status": self.status,
            "bytes_downloaded": self.bytes_downloaded,
            "rows_extracted": self.rows_extracted,
            "candidates_created": self.candidates_created,
            "detail": self.detail,
        }


class OfficialAttachmentProcessor:
    """Download and parse one registered official attachment at a time."""

    def __init__(
        self,
        settings: Settings,
        database: Database,
        session: requests.Session | None = None,
    ) -> None:
        self.settings = settings
        self.database = database
        headers = {"User-Agent": USER_AGENT, "Accept": "*/*"}
        self.session = session or create_session(
            settings.http_transport_mode,
            headers=headers,
        )
        if session is not None:
            configure_session(session, settings.http_transport_mode, headers=headers)
        self._robots: dict[str, RobotFileParser] = {}

    def discover_from_page(
        self,
        source_id: str,
        parent_url: str,
        *,
        html: str | None = None,
    ) -> list[dict[str, object]]:
        """Register file links from one approved public announcement page.

        Discovery never downloads an attachment. It only reads the page (or
        one public GET when ``html`` is omitted), checks the source host and
        robots policy, and places metadata in the private attachment ledger.
        """
        source = self.database.get_source(source_id)
        if source is None:
            raise ValueError("Source is not registered")
        if not is_http_url(parent_url):
            raise ValueError("parent_url must be a complete HTTP(S) URL")
        self._assert_official_hosts(source, parent_url, parent_url)
        self._robots_allowed(parent_url)
        if html is None:
            try:
                response = self.session.get(
                    parent_url,
                    timeout=self.settings.request_timeout_seconds,
                    allow_redirects=True,
                )
            except requests.RequestException as error:
                raise AttachmentProcessingError(
                    f"Announcement page request failed: {error}"
                ) from error
            if response.status_code in {401, 403, 407, 429}:
                raise AttachmentSkipped(
                    f"Announcement page access is restricted (HTTP {response.status_code})"
                )
            if response.status_code >= 400:
                raise AttachmentProcessingError(
                    f"Announcement page returned HTTP {response.status_code}"
                )
            if len(response.content) > self.settings.attachment_discovery_max_bytes:
                raise AttachmentSkipped("Announcement page exceeds the discovery size limit")
            self._assert_official_hosts(source, parent_url, response.url)
            self._robots_allowed(response.url)
            encoding = response.encoding
            if (
                not encoding
                or encoding.lower() in {"iso-8859-1", "ascii"}
            ) and response.apparent_encoding:
                encoding = response.apparent_encoding
            html = response.content.decode(encoding or "utf-8", errors="replace")
        soup = BeautifulSoup(html, "html.parser")
        page_text = " ".join(soup.get_text(" ", strip=True).split())
        page_metadata = {
            "employer": str(source.get("publisher") or ""),
            "category": str(source.get("category") or ""),
            "published_date": extract_published_date(page_text),
            "deadline_date": extract_deadline(page_text),
            "official_notice_title": (
                soup.title.get_text(" ", strip=True) if soup.title else ""
            ),
        }
        discovered: list[dict[str, object]] = []
        seen: set[str] = set()
        from urllib.parse import urljoin

        for anchor in soup.find_all("a", href=True):
            href = str(anchor.get("href") or "").strip()
            if not href or href.startswith(("javascript:", "mailto:", "#")):
                continue
            artifact_url = urljoin(parent_url, href)
            if artifact_url in seen or not is_http_url(artifact_url):
                continue
            label = " ".join(anchor.get_text(" ", strip=True).split())
            suffix = Path(urlparse(artifact_url).path).suffix.lower()
            if suffix not in SUPPORTED_SUFFIXES:
                continue
            try:
                self._assert_official_hosts(source, parent_url, artifact_url)
                self._robots_allowed(artifact_url)
            except AttachmentSkipped:
                continue
            seen.add(artifact_url)
            kind = "position_table" if any(
                word in f"{label} {artifact_url}"
                for word in ("职位", "岗位", "招聘", "position", "career")
            ) else "announcement_attachment"
            discovered.append(
                self.database.upsert_source_artifact(
                    {
                        "source_id": source_id,
                        "parent_url": parent_url,
                        "artifact_url": artifact_url,
                        "artifact_kind": kind,
                        "media_type": mimetypes.guess_type(artifact_url)[0],
                        "metadata": {
                            "display_name": label or Path(urlparse(artifact_url).path).name,
                            "discovered_from": parent_url,
                            "discovered_at": _utc_now(),
                            "discovery_only": True,
                            **page_metadata,
                        },
                    }
                )
            )
        return discovered

    def process(
        self,
        artifact_id: int,
        *,
        force_download: bool = False,
        extract: bool = True,
    ) -> AttachmentResult:
        """Download, hash, parse and queue candidates for one attachment."""
        artifact = self.database.get_source_artifact(artifact_id)
        if artifact is None:
            raise ValueError("Source artifact does not exist")
        try:
            if (
                artifact.get("extraction_status") in {"downloaded", "extracted"}
                and artifact.get("storage_path")
                and not force_download
            ):
                bytes_downloaded = self._stored_size(artifact)
            else:
                artifact, bytes_downloaded = self._download(artifact)
            if not extract:
                return AttachmentResult(
                    artifact_id=artifact_id,
                    status=str(artifact["extraction_status"]),
                    bytes_downloaded=bytes_downloaded,
                    detail="downloaded; extraction was not requested",
                )
            extracted = self._extract(artifact)
            rows = self.database.upsert_source_artifact_rows(
                artifact_id,
                extracted["rows"],
            )
            updated = self.database.update_source_artifact_processing(
                artifact_id,
                extraction_status="extracted",
                parser_version=PARSER_VERSION,
                metadata_updates=extracted["metadata"],
            )
            candidates = self._queue_candidates(updated, rows)
            return AttachmentResult(
                artifact_id=artifact_id,
                status="extracted",
                bytes_downloaded=bytes_downloaded,
                rows_extracted=len(rows),
                candidates_created=len(candidates),
                detail=str(extracted["metadata"].get("extraction_note") or ""),
            )
        except AttachmentSkipped as error:
            self._record_failure(artifact_id, "skipped", str(error))
            return AttachmentResult(artifact_id, "skipped", detail=str(error))
        except Exception as error:
            self._record_failure(artifact_id, "failed", str(error))
            raise AttachmentProcessingError(str(error)) from error

    def _download(
        self,
        artifact: dict[str, object],
    ) -> tuple[dict[str, object], int]:
        source = self.database.get_source(str(artifact["source_id"]))
        if source is None:
            raise AttachmentProcessingError("Attachment source is no longer registered")
        url = str(artifact["artifact_url"])
        parent_url = str(artifact["parent_url"])
        if not is_http_url(url) or not is_http_url(parent_url):
            raise AttachmentProcessingError("Attachment URLs must be complete HTTP(S) URLs")
        self._assert_official_hosts(source, parent_url, url)
        self._robots_allowed(parent_url)
        self._robots_allowed(url)
        try:
            response = self.session.get(
                url,
                stream=True,
                allow_redirects=True,
                timeout=self.settings.request_timeout_seconds,
            )
        except requests.RequestException as error:
            raise AttachmentProcessingError(f"Attachment request failed: {error}") from error
        try:
            if response.status_code in {401, 403, 407, 429}:
                raise AttachmentSkipped(
                    f"Attachment access is restricted (HTTP {response.status_code}); no bypass attempted"
                )
            if response.status_code >= 400:
                raise AttachmentProcessingError(
                    f"Attachment request returned HTTP {response.status_code}"
                )
            self._assert_official_hosts(source, parent_url, response.url)
            self._robots_allowed(response.url)
            content_type = self._content_type(response.headers.get("Content-Type"))
            suffix = self._suffix(response.url, content_type)
            if suffix not in SUPPORTED_SUFFIXES:
                raise AttachmentSkipped(
                    f"Unsupported attachment type: {content_type or 'unknown'}"
                )
            declared_length = response.headers.get("Content-Length")
            if declared_length and int(declared_length) > self.settings.attachment_max_bytes:
                raise AttachmentSkipped("Attachment exceeds the configured size limit")
            root = self.settings.managed_artifact_dir()
            root.mkdir(parents=True, exist_ok=True)
            fd, temporary_name = tempfile.mkstemp(prefix=".download-", suffix=".part", dir=root)
            total = 0
            digest = hashlib.sha256()
            try:
                with os.fdopen(fd, "wb") as handle:
                    for chunk in response.iter_content(chunk_size=64 * 1024):
                        if not chunk:
                            continue
                        total += len(chunk)
                        if total > self.settings.attachment_max_bytes:
                            raise AttachmentSkipped("Attachment exceeds the configured size limit")
                        if total == len(chunk) and self._looks_like_html(chunk):
                            raise AttachmentSkipped(
                                "Official endpoint returned an HTML/login page instead of a file"
                            )
                        digest.update(chunk)
                        handle.write(chunk)
                hexdigest = digest.hexdigest()
                relative_path = Path("sha256") / hexdigest[:2] / f"{hexdigest}{suffix}"
                destination = root / relative_path
                destination.parent.mkdir(parents=True, exist_ok=True)
                if destination.exists():
                    Path(temporary_name).unlink(missing_ok=True)
                else:
                    Path(temporary_name).replace(destination)
            except Exception:
                Path(temporary_name).unlink(missing_ok=True)
                raise
        finally:
            response.close()
        updated = self.database.update_source_artifact_processing(
            int(artifact["id"]),
            extraction_status="downloaded",
            media_type=content_type or str(artifact.get("media_type") or "") or None,
            content_sha256=hexdigest,
            storage_path=relative_path.as_posix(),
            parser_version=PARSER_VERSION,
            metadata_updates={
                "downloaded_url": response.url,
                "downloaded_bytes": total,
                "downloaded_at": _utc_now(),
                "content_type": content_type,
                "suffix": suffix,
            },
        )
        return updated, total

    def _extract(self, artifact: dict[str, object]) -> dict[str, object]:
        path = self._artifact_path(artifact)
        suffix = Path(path).suffix.lower()
        if suffix == ".pdf":
            return self._extract_pdf(path)
        if suffix == ".xls":
            return self._extract_legacy_workbook(path)
        if suffix in {".xlsx", ".xlsm"}:
            return self._extract_workbook(path)
        if suffix == ".csv":
            return self._extract_csv(path)
        if suffix == ".docx":
            return self._extract_docx(path)
        raise AttachmentSkipped(f"No parser is registered for {suffix}")

    def _extract_pdf(self, path: Path) -> dict[str, object]:
        try:
            from pypdf import PdfReader
        except ImportError as error:
            raise AttachmentSkipped("PDF parser pypdf is not installed") from error
        try:
            reader = PdfReader(str(path))
            page_texts = [(page.extract_text() or "").strip() for page in reader.pages]
        except Exception as error:
            raise AttachmentProcessingError(f"PDF text extraction failed: {error}") from error
        text = "\n\n".join(page_texts).strip()
        ocr_used = False
        extraction_note = "text layer extracted"
        if not text:
            if not self.settings.attachment_ocr_enabled:
                extraction_note = "PDF has no text layer; OCR is disabled"
            else:
                text, ocr_note = self._ocr_pdf(path, min(len(reader.pages), self.settings.attachment_ocr_max_pages))
                ocr_used = bool(text)
                extraction_note = ocr_note
        text = text[: self.settings.attachment_max_text_characters]
        text_storage_path = self._write_text_sidecar(path, text)
        rows = self._rows_from_pdf_pages(page_texts if not ocr_used else [text])
        if ocr_used:
            rows = [dict(row, extraction_confidence="low") for row in rows]
        return {
            "rows": rows[: self.settings.attachment_max_rows],
            "metadata": {
                "parser": "pypdf",
                "page_count": len(reader.pages),
                "text_characters": len(text),
                "text_sha256": hashlib.sha256(text.encode("utf-8")).hexdigest(),
                "extraction_mode": "ocr" if ocr_used else "text",
                "extraction_note": extraction_note,
                "ocr_enabled": self.settings.attachment_ocr_enabled,
                "text_storage_path": text_storage_path,
            },
        }

    def _extract_workbook(self, path: Path) -> dict[str, object]:
        try:
            from openpyxl import load_workbook
        except ImportError as error:
            raise AttachmentSkipped("Excel parser openpyxl is not installed") from error
        try:
            workbook_input = (
                io.BytesIO(path.read_bytes())
                if path.suffix.lower() == ".xls"
                else path
            )
            workbook = load_workbook(
                filename=workbook_input,
                read_only=True,
                data_only=True,
            )
        except Exception as error:
            raise AttachmentProcessingError(f"Excel extraction failed: {error}") from error
        rows: list[dict[str, object]] = []
        sheet_count = 0
        try:
            for worksheet in workbook.worksheets:
                sheet_count += 1
                raw_rows: list[list[str]] = []
                for values in worksheet.iter_rows(values_only=True):
                    cells = [self._cell_text(value) for value in values]
                    if any(cells):
                        raw_rows.append(cells)
                    if len(raw_rows) >= self.settings.attachment_max_rows + 1:
                        break
                if not raw_rows:
                    continue
                header, data_start = self._workbook_header_info(raw_rows)
                data_rows = raw_rows[data_start:] if header else raw_rows
                carry: dict[int, str] = {}
                for offset, values in enumerate(
                    data_rows,
                    start=data_start + 1 if header else 1,
                ):
                    values = self._expand_merged_values(values, header, carry)
                    cells = {
                        (header[index] if header and index < len(header) and header[index] else f"列{index + 1}"): value
                        for index, value in enumerate(values)
                        if value
                    }
                    if not cells:
                        continue
                    rows.append(
                        {
                            "sheet_name": worksheet.title,
                            "row_number": offset,
                            "row_kind": "tabular",
                            "cells": cells,
                            "row_text": "；".join(f"{key}：{value}" for key, value in cells.items()),
                            "extraction_confidence": "high",
                        }
                    )
                    if len(rows) >= self.settings.attachment_max_rows:
                        break
                if len(rows) >= self.settings.attachment_max_rows:
                    break
        finally:
            workbook.close()
        return {
            "rows": rows,
            "metadata": {
                "parser": "openpyxl",
                "sheet_count": sheet_count,
                "row_count": len(rows),
                "extraction_mode": "workbook_rows",
                "extraction_note": "表格行已保存为私有待核验记录",
            },
        }

    def _extract_legacy_workbook(self, path: Path) -> dict[str, object]:
        """Extract legacy BIFF ``.xls`` files used by government portals.

        A large share of provincial recruitment tables still uses the binary
        Excel format.  Keep it on the same private review path as XLSX: rows
        are evidence candidates, not automatically published vacancies.
        """
        # A number of government CMS instances keep an ``.xls`` display name
        # while serving an OOXML workbook.  Detect the ZIP signature before
        # handing the bytes to xlrd, otherwise a valid table is reported as a
        # parser failure.
        if path.read_bytes()[:2] == b"PK":
            extracted = self._extract_workbook(path)
            metadata = dict(extracted["metadata"])
            metadata["parser"] = "openpyxl (mislabelled .xls)"
            metadata["extraction_note"] = (
                "文件扩展名为 .xls，但内容为 OOXML；已用 openpyxl 提取并保留原始哈希"
            )
            extracted["metadata"] = metadata
            return extracted
        try:
            import xlrd
        except ImportError as error:
            raise AttachmentSkipped("Legacy Excel parser xlrd is not installed") from error
        try:
            workbook = xlrd.open_workbook(filename=str(path), on_demand=True)
        except Exception as error:
            raise AttachmentProcessingError(
                f"Legacy Excel extraction failed: {error}"
            ) from error
        rows: list[dict[str, object]] = []
        sheet_count = 0
        try:
            for sheet in workbook.sheets():
                sheet_count += 1
                raw_rows: list[list[str]] = []
                for row_number in range(sheet.nrows):
                    values = [
                        self._cell_text(sheet.cell_value(row_number, column))
                        for column in range(sheet.ncols)
                    ]
                    if any(values):
                        raw_rows.append(values)
                    if len(raw_rows) >= self.settings.attachment_max_rows + 1:
                        break
                if not raw_rows:
                    continue
                header, data_start = self._workbook_header_info(raw_rows)
                data_rows = raw_rows[data_start:] if header else raw_rows
                carry: dict[int, str] = {}
                for offset, values in enumerate(
                    data_rows,
                    start=data_start + 1 if header else 1,
                ):
                    values = self._expand_merged_values(values, header, carry)
                    cells = {
                        (
                            header[index]
                            if header and index < len(header) and header[index]
                            else f"列{index + 1}"
                        ): value
                        for index, value in enumerate(values)
                        if value
                    }
                    if not cells:
                        continue
                    rows.append(
                        {
                            "sheet_name": sheet.name,
                            "row_number": offset,
                            "row_kind": "tabular",
                            "cells": cells,
                            "row_text": "；".join(
                                f"{key}：{value}" for key, value in cells.items()
                            ),
                            "extraction_confidence": "high",
                        }
                    )
                    if len(rows) >= self.settings.attachment_max_rows:
                        break
                if len(rows) >= self.settings.attachment_max_rows:
                    break
        finally:
            release = getattr(workbook, "release_resources", None)
            if callable(release):
                release()
        return {
            "rows": rows,
            "metadata": {
                "parser": "xlrd",
                "sheet_count": sheet_count,
                "row_count": len(rows),
                "extraction_mode": "workbook_rows",
                "extraction_note": "XLS 表格行已保存为私有待核验记录",
            },
        }

    def _extract_csv(self, path: Path) -> dict[str, object]:
        raw = path.read_bytes()
        text = raw.decode("utf-8-sig", errors="replace")
        sample = text[:8192]
        try:
            dialect = csv.Sniffer().sniff(sample, delimiters=",\t;|")
        except csv.Error:
            dialect = csv.excel
        reader = csv.reader(io.StringIO(text), dialect)
        all_rows = [
            [self._cell_text(cell) for cell in values]
            for values in reader
        ]
        all_rows = [row for row in all_rows if any(row)]
        header = self._workbook_header(all_rows[0]) if all_rows else []
        rows: list[dict[str, object]] = []
        for number, values in enumerate(all_rows[1:] if header else all_rows, start=2 if header else 1):
            cells = {
                (header[index] if header and index < len(header) and header[index] else f"列{index + 1}"): value
                for index, value in enumerate(values)
                if value
            }
            if cells:
                rows.append(
                    {
                        "sheet_name": "csv",
                        "row_number": number,
                        "row_kind": "tabular",
                        "cells": cells,
                        "row_text": "；".join(f"{key}：{value}" for key, value in cells.items()),
                        "extraction_confidence": "high",
                    }
                )
            if len(rows) >= self.settings.attachment_max_rows:
                break
        return {
            "rows": rows,
            "metadata": {
                "parser": "csv",
                "row_count": len(rows),
                "text_characters": len(text),
                "extraction_mode": "workbook_rows",
                "extraction_note": "CSV 行已保存为私有待核验记录",
            },
        }

    def _extract_docx(self, path: Path) -> dict[str, object]:
        try:
            from docx import Document
        except ImportError as error:
            raise AttachmentSkipped("DOCX parser python-docx is not installed") from error
        document = Document(str(path))
        paragraphs = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]
        rows = [
            {
                "sheet_name": "docx",
                "row_number": index,
                "row_kind": "text_table",
                "cells": {"正文": value},
                "row_text": value,
                "extraction_confidence": "medium",
            }
            for index, value in enumerate(paragraphs[: self.settings.attachment_max_rows], start=1)
        ]
        return {
            "rows": rows,
            "metadata": {
                "parser": "python-docx",
                "row_count": len(rows),
                "extraction_mode": "paragraphs",
                "extraction_note": "DOCX 段落已保存为私有待核验记录",
            },
        }

    def _queue_candidates(
        self,
        artifact: dict[str, object],
        rows: list[dict[str, object]],
    ) -> list[dict[str, object]]:
        candidates: list[dict[str, object]] = []
        for row in rows:
            candidate = self._candidate_from_row(artifact, row)
            if candidate is None:
                continue
            candidates.append(
                self.database.upsert_artifact_job_candidate(candidate)
            )
        return candidates

    def _candidate_from_row(
        self,
        artifact: dict[str, object],
        row: dict[str, object],
    ) -> dict[str, object] | None:
        cells = row.get("cells") or {}
        if not isinstance(cells, dict):
            return None
        text = str(row.get("row_text") or "").strip()
        title = self._field(cells, ("岗位名称", "岗位", "职位", "招聘岗位", "岗位名称（岗位）"))
        employer = self._field(cells, ("用人单位", "招聘单位", "单位名称", "单位", "招聘机构"))
        location = self._field(cells, ("工作地点", "工作区域", "工作城市", "所在地", "地点"))
        degree = self._field(cells, ("学历", "学历要求", "学位要求", "面向对象"))
        major = self._field(cells, ("专业", "专业要求", "需求专业", "专业范围", "所学专业"))
        published = self._field(cells, ("发布日期", "发布时间", "公告日期", "发布日"))
        deadline = self._field(cells, ("报名截止", "截止日期", "报名截止日期", "截止时间", "报名时间"))
        if not title:
            for line in text.splitlines():
                if len(line.strip()) >= 4 and any(word in line for word in ("招聘", "岗位", "工程师", "研究员", "教师")):
                    title = line.strip()[:160]
                    break
        if title and major and title in {"专业技术", "专业技术岗", "专业技术岗位"}:
            title = f"{title}（{major[:100]}）"
        relevant_text = " ".join(filter(None, (title, major, degree, text)))
        if not title or not (extract_major_tags(relevant_text) or extract_degree_levels(relevant_text)):
            return None
        metadata = artifact.get("metadata") or {}
        source = self.database.get_source(str(artifact.get("source_id") or "")) or {}
        employer = employer or str(
            metadata.get("employer")
            or source.get("publisher")
            or "官方公告单位"
        )
        published = published or str(metadata.get("published_date") or "") or None
        deadline = deadline or str(metadata.get("deadline_date") or "") or None
        score, _, major_tags = score_relevance(
            relevant_text,
            "A",
            str(artifact.get("metadata", {}).get("category") or "能源、工程与地学拓展"),
        )
        field_evidence = {
            "row_locator": f"{row.get('sheet_name')}!{row.get('row_number')}",
            "artifact_url": artifact["artifact_url"],
            "row_text": text[:2000],
            "fields": {
                key: value
                for key, value in {
                    "title": title,
                    "employer": employer,
                    "location": location,
                    "degree": degree,
                    "major": major,
                    "published_date": published,
                    "deadline_date": deadline,
                }.items()
                if value
            },
        }
        return {
            "artifact_row_id": row["id"],
            "source_id": artifact["source_id"],
            "official_page_url": artifact["parent_url"],
            "title": title,
            "employer": employer,
            "application_url": None,
            "location": location,
            "published_date": _normalize_date(published),
            "deadline_date": _normalize_date(deadline),
            "degree_levels": extract_degree_levels(degree or relevant_text),
            "major_tags": major_tags,
            "summary": text[:420],
            "description": text[: self.settings.attachment_max_text_characters],
            "field_evidence": field_evidence,
            "relevance_score": score,
            "review_status": "needs_review",
            # Keep this empty: a non-empty note is reserved for a human's
            # post-extraction verification decision, not the parser's status.
            "review_note": "",
        }

    @staticmethod
    def _field(cells: dict[object, object], aliases: tuple[str, ...]) -> str | None:
        normalized = {
            re.sub(r"[\s:：（）()【】\[\]]", "", str(key)).lower(): str(value).strip()
            for key, value in cells.items()
            if str(value).strip()
        }
        for alias in aliases:
            key = re.sub(r"[\s:：（）()【】\[\]]", "", alias).lower()
            if key in normalized:
                return normalized[key][:500]
        return None

    @staticmethod
    def _workbook_header(row: list[str]) -> list[str]:
        if not row:
            return []
        keywords = ("岗位", "职位", "单位", "专业", "学历", "地点", "报名", "截止")
        score = sum(any(keyword in cell for keyword in keywords) for cell in row)
        if score < 1:
            return []
        return [cell or f"列{index + 1}" for index, cell in enumerate(row)]

    @classmethod
    def _workbook_header_info(cls, rows: list[list[str]]) -> tuple[list[str], int]:
        """Find a possibly multi-row header and return its first data index."""
        for index, row in enumerate(rows[:8]):
            header = cls._workbook_header(row)
            if not header:
                continue
            end = index + 1
            while end < len(rows):
                candidate = rows[end]
                score = sum(
                    any(
                        keyword in cell
                        for keyword in ("岗位", "职位", "单位", "专业", "学历", "地点", "报名", "截止", "学位")
                    )
                    for cell in candidate
                    if cell
                )
                # A data row can contain the word “专业”; continuation rows
                # usually expose at least two field labels and no job code.
                has_code = any(re.search(r"20\d{2}\d+", cell) for cell in candidate if cell)
                if score < 2 or has_code:
                    break
                for column, value in enumerate(candidate):
                    if not value:
                        continue
                    parent = header[column] if column < len(header) else ""
                    if not parent or parent.startswith("列"):
                        header[column] = value
                    elif value not in parent:
                        header[column] = value
                end += 1
            return header, end
        return [], 0

    @staticmethod
    def _expand_merged_values(
        values: list[str],
        header: list[str],
        carry: dict[int, str],
    ) -> list[str]:
        """Forward-fill only organizational columns merged in official tables."""
        expanded = list(values)
        carry_columns = {"主管部门", "招聘单位", "用人单位", "招聘机构", "单位名称"}
        for index, value in enumerate(expanded):
            label = header[index] if index < len(header) else ""
            if value:
                carry[index] = value
            elif label in carry_columns and index in carry:
                expanded[index] = carry[index]
        return expanded

    @staticmethod
    def _cell_text(value: object) -> str:
        if value is None:
            return ""
        if isinstance(value, datetime):
            return value.date().isoformat()
        return str(value).strip()

    @staticmethod
    def _rows_from_pdf_pages(page_texts: list[str]) -> list[dict[str, object]]:
        rows: list[dict[str, object]] = []
        for page_number, page in enumerate(page_texts, start=1):
            lines = [line.strip() for line in page.splitlines() if line.strip()]
            if not lines:
                continue
            for line_number, line in enumerate(lines, start=1):
                cells: dict[str, str]
                if "\t" in line:
                    parts = [part.strip() for part in line.split("\t") if part.strip()]
                    cells = {f"列{index + 1}": part for index, part in enumerate(parts)}
                else:
                    parts = [part.strip() for part in re.split(r"\s{2,}", line) if part.strip()]
                    cells = {f"列{index + 1}": part for index, part in enumerate(parts)}
                rows.append(
                    {
                        "sheet_name": f"page-{page_number}",
                        "row_number": line_number,
                        "row_kind": "text_table",
                        "cells": cells or {"正文": line},
                        "row_text": line,
                        "extraction_confidence": "medium" if len(cells) > 1 else "low",
                    }
                )
        return rows

    def _ocr_pdf(self, path: Path, page_count: int) -> tuple[str, str]:
        if page_count < 1:
            return "", "PDF has no pages for OCR"
        with tempfile.TemporaryDirectory(prefix="job-hub-ocr-") as directory:
            prefix = str(Path(directory) / "page")
            command = [
                "pdftoppm",
                "-f",
                "1",
                "-l",
                str(page_count),
                "-r",
                "160",
                "-png",
                str(path),
                prefix,
            ]
            try:
                subprocess.run(command, check=True, capture_output=True, timeout=180)
            except (OSError, subprocess.CalledProcessError, subprocess.TimeoutExpired) as error:
                return "", f"OCR rasterization unavailable: {error}"
            texts: list[str] = []
            for image in sorted(Path(directory).glob("page-*.png")):
                try:
                    result = subprocess.run(
                        ["tesseract", str(image), "stdout", "-l", self.settings.attachment_ocr_language],
                        check=True,
                        capture_output=True,
                        timeout=120,
                    )
                except (OSError, subprocess.CalledProcessError, subprocess.TimeoutExpired) as error:
                    return "", f"OCR engine unavailable: {error}"
                texts.append(result.stdout.decode("utf-8", errors="replace"))
            return "\n\n".join(texts)[: self.settings.attachment_max_text_characters], "OCR fallback extracted text with low confidence"

    def _artifact_path(self, artifact: dict[str, object]) -> Path:
        storage_path = str(artifact.get("storage_path") or "")
        if not storage_path:
            raise AttachmentProcessingError("Attachment has no managed storage path")
        root = self.settings.managed_artifact_dir().resolve()
        path = (root / storage_path).resolve()
        try:
            path.relative_to(root)
        except ValueError as error:
            raise AttachmentProcessingError("Attachment storage path escapes managed directory") from error
        if not path.is_file():
            raise AttachmentProcessingError("Managed attachment file does not exist")
        return path

    def _stored_size(self, artifact: dict[str, object]) -> int:
        return self._artifact_path(artifact).stat().st_size

    def _write_text_sidecar(self, artifact_path: Path, text: str) -> str | None:
        if not text:
            return None
        root = self.settings.managed_artifact_dir().resolve()
        sidecar = artifact_path.with_suffix(artifact_path.suffix + ".txt")
        try:
            sidecar.relative_to(root)
        except ValueError as error:  # pragma: no cover - storage path is contract validated
            raise AttachmentProcessingError(
                "Extracted text path escapes managed directory"
            ) from error
        sidecar.write_text(text, encoding="utf-8")
        return sidecar.relative_to(root).as_posix()

    def _record_failure(self, artifact_id: int, status: str, detail: str) -> None:
        try:
            self.database.update_source_artifact_processing(
                artifact_id,
                extraction_status=status,
                parser_version=PARSER_VERSION,
                metadata_updates={"last_error": detail[:2000], "last_error_at": _utc_now()},
            )
        except Exception:
            # Preserve the original processing error; the DB record is best effort.
            pass

    def _robots_allowed(self, url: str) -> None:
        parsed = urlparse(url)
        root = f"{parsed.scheme}://{parsed.netloc}"
        parser = self._robots.get(root)
        if parser is None:
            parser = RobotFileParser()
            robots_url = f"{root}/robots.txt"
            try:
                response = self.session.get(
                    robots_url,
                    timeout=min(self.settings.request_timeout_seconds, 10),
                    allow_redirects=True,
                )
            except requests.RequestException as error:
                raise AttachmentSkipped(f"Unable to verify robots.txt for {root}: {error}") from error
            if response.status_code == 404:
                parser.parse([])
            elif response.ok:
                parser.parse(response.text.splitlines())
            else:
                raise AttachmentSkipped(
                    f"Unable to verify robots.txt for {root}: HTTP {response.status_code}"
                )
            self._robots[root] = parser
        if not parser.can_fetch(USER_AGENT, url):
            raise AttachmentSkipped(f"robots.txt does not permit attachment access: {url}")

    @staticmethod
    def _assert_official_hosts(source: dict[str, object], parent_url: str, artifact_url: str) -> None:
        config = source.get("config") or {}
        if not isinstance(config, dict):
            raise AttachmentProcessingError("Registered source config is invalid")
        hosts = {
            str(value).lower().strip()
            for value in config.get("allowed_hosts", [])
            if str(value).strip()
        }
        homepage_host = urlparse(str(source.get("homepage_url") or "")).hostname
        if homepage_host:
            hosts.add(homepage_host.lower())
        hosts.update(
            str(value).lower().strip()
            for value in config.get("attachment_allowed_hosts", [])
            if str(value).strip()
        )
        for label, value in (("parent", parent_url), ("attachment", artifact_url)):
            host = urlparse(value).hostname
            if not host or host.lower() not in hosts:
                raise AttachmentSkipped(f"{label} URL host is outside the official source allowlist")

    @staticmethod
    def _content_type(value: str | None) -> str | None:
        if not value:
            return None
        return value.split(";", 1)[0].strip().lower() or None

    @staticmethod
    def _suffix(url: str, content_type: str | None) -> str:
        suffix = Path(urlparse(url).path).suffix.lower()
        if suffix in SUPPORTED_SUFFIXES:
            return suffix
        mapping = {
            "application/pdf": ".pdf",
            "text/csv": ".csv",
            "application/csv": ".csv",
            "application/vnd.ms-excel": ".xls",
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet": ".xlsx",
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document": ".docx",
            "application/msword": ".doc",
        }
        return mapping.get(content_type or "", suffix or ".bin")

    @staticmethod
    def _looks_like_html(chunk: bytes) -> bool:
        sample = chunk.lstrip().lower()[:256]
        return any(sample.startswith(marker) for marker in HTML_MARKERS)


def _utc_now() -> str:
    return datetime.now(timezone.utc).replace(microsecond=0).isoformat().replace("+00:00", "Z")


def _normalize_date(value: str | None) -> str | None:
    if not value:
        return None
    match = re.search(r"(20\d{2})\s*[年./-]\s*(\d{1,2})\s*[月./-]\s*(\d{1,2})", value)
    if not match:
        return value[:50]
    return f"{int(match.group(1)):04d}-{int(match.group(2)):02d}-{int(match.group(3)):02d}"
