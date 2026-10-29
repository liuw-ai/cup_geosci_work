"""Batch parser for official civil-service/public-institution position tables.

The parser is deliberately a *normalizer*, not a publisher. It converts rows
from common government spreadsheets into the compact ``position_batches``
shape already consumed by :mod:`job_hub.government_positions`. Missing or
ambiguous cells are returned in a rejection ledger with their original row
number; no value is inferred from a title or a neighboring row.
"""

from __future__ import annotations

import csv
import json
import re
import subprocess
from pathlib import Path
from typing import Any, Iterable

from job_hub.locations import extract_location_provinces


FIELD_ALIASES: dict[str, tuple[str, ...]] = {
    "position_code": ("岗位代码", "职位代码", "招考职位代码", "职位编号", "岗位编号", "代码"),
    "employer": (
        "用人单位",
        "招聘单位",
        "招录机关",
        "招录单位",
        "用人司局",
        "部门名称",
        "单位名称",
        "机构名称",
        "招考部门",
    ),
    "title": ("岗位名称", "职位名称", "招考职位", "招聘岗位", "招聘岗位类别", "具体岗位名称", "岗位类别", "岗位", "职位"),
    "major": ("专业要求", "专业条件", "专业条件要求", "专业", "招聘专业", "招聘条件", "需求专业", "岗位所需专业", "专业代码"),
    "degree": ("学历要求", "学历", "学历学位", "学历及学位", "学位要求", "学位", "面向对象"),
    "headcount": (
        "招聘人数",
        "招录人数",
        "招考人数",
        "计划人数",
        "计划招录人数",
        "人数",
        "招聘计划",
    ),
    "location": ("工作地点", "工作地", "工作城市", "所在地", "地点", "落户地点"),
}

_IMAGE_SUFFIXES = frozenset({".bmp", ".jpeg", ".jpg", ".png", ".tif", ".tiff", ".webp"})


def _compact(value: Any) -> str:
    return re.sub(r"[\s\u3000\n\r]+", "", str(value or "")).strip()


def _text(value: Any) -> str:
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return re.sub(r"\s+", " ", str(value or "")).strip()


def _header_map(row: dict[str, Any]) -> dict[str, str]:
    return {_compact(key): _text(value) for key, value in row.items() if _compact(key)}


def _field(row: dict[str, str], aliases: Iterable[str]) -> str:
    normalized = {_compact(key): value for key, value in row.items()}
    for alias in aliases:
        value = normalized.get(_compact(alias), "")
        if value:
            return value
    # Some provincial sheets append units or footnote markers to headers.
    for key, value in normalized.items():
        if value and any(len(_compact(alias)) >= 3 and _compact(alias) in key for alias in aliases):
            return value
    return ""


def _headcount(value: str) -> int | None:
    raw = _text(value).replace(",", "")
    if not raw:
        return None
    if re.search(r"[-—~至]|以上|以下|若干|不限|面议", raw):
        return None
    match = re.search(r"(?<!\d)(\d+)(?!\d)", raw)
    if not match:
        return None
    number = int(match.group(1))
    return number if number >= 0 else None


def parse_position_table_rows(
    rows: Iterable[dict[str, Any]],
    *,
    batch_id: str,
    source_id: str,
    position_type: str,
    province: str,
    official_notice_url: str,
    official_attachment_url: str,
    deadline_date: str = "",
    deadline_policy: str = "fixed_date",
    location: str = "",
    location_evidence_policy: str = "row_exact",
    sheet_name: str = "Sheet1",
    first_data_row: int = 2,
) -> dict[str, Any]:
    """Normalize spreadsheet dictionaries into a validated batch candidate."""
    accepted: list[list[Any]] = []
    accepted_evidence: list[dict[str, Any]] = []
    rejected: list[dict[str, Any]] = []
    for offset, raw in enumerate(rows):
        row_number = first_data_row + offset
        if not isinstance(raw, dict):
            rejected.append({"row_number": row_number, "reason": "row_not_object"})
            continue
        row_number = int(raw.get("__row_number", row_number))
        current_sheet = str(raw.get("__sheet_name") or sheet_name)
        cells = _header_map({key: value for key, value in raw.items() if not key.startswith("__")})
        extracted = {
            key: _field(cells, aliases) for key, aliases in FIELD_ALIASES.items()
        }
        degree_parts = [
            _field(cells, ("学历要求", "学历", "学历层次", "学历学位")),
            _field(cells, ("学位要求", "学位")),
        ]
        extracted["degree"] = "；".join(dict.fromkeys(item for item in degree_parts if item))
        missing = [key for key in ("position_code", "employer", "title", "major", "degree", "headcount") if not extracted[key]]
        count = _headcount(extracted["headcount"])
        if count is None:
            missing.append("headcount_integer")
        if missing:
            rejected.append(
                {
                    "row_number": row_number,
                    "sheet_name": current_sheet,
                    "reason": "missing_or_invalid_fields",
                    "missing": list(dict.fromkeys(missing)),
                    "cells": cells,
                }
            )
            continue
        row_location = extracted["location"] or location
        expected_provinces = set(extract_location_provinces(province))
        observed_provinces = set(extract_location_provinces(f"{current_sheet} {row_location}"))
        if expected_provinces and observed_provinces and observed_provinces != expected_provinces:
            rejected.append(
                {
                    "row_number": row_number,
                    "sheet_name": current_sheet,
                    "reason": "province_mismatch",
                    "expected_province": province,
                    "observed_provinces": sorted(observed_provinces),
                    "cells": cells,
                }
            )
            continue
        locator = f"{current_sheet}!{row_number}；岗位代码{extracted['position_code']}"
        accepted.append(
            [
                extracted["position_code"],
                extracted["employer"],
                extracted["title"],
                extracted["major"],
                extracted["degree"],
                count,
                locator,
                row_location,
                province,
            ]
        )
        accepted_evidence.append(
            {
                "row_number": row_number,
                "sheet_name": current_sheet,
                "position_code": extracted["position_code"],
                "cells": cells,
                "locator": locator,
            }
        )
    batch = {
        "id": batch_id,
        "source_id": source_id,
        "position_type": position_type,
        "province": province,
        "location": location,
        "location_evidence_policy": location_evidence_policy,
        "deadline_date": deadline_date,
        "deadline_policy": deadline_policy,
        "official_notice_url": official_notice_url,
        "official_attachment_url": official_attachment_url,
        "record_status": "manual_review_required",
        "match_status": "needs_review",
        "rows": accepted,
        "note": "批量解析器仅完成列映射和逐行定位；官方哈希、公告、专业/学历、地点、当前性和学生端门禁仍需后续核验。",
    }
    return {
        "version": 1,
        "batch": batch,
        "accepted_rows": accepted_evidence,
        "rejected_rows": rejected,
        "metrics": {
            "input_rows": len(accepted_evidence) + len(rejected),
            "accepted_rows": len(accepted_evidence),
            "rejected_rows": len(rejected),
            "field_contract": "position_code, employer, title, major, degree, headcount",
        },
    }


def _is_header(values: list[Any]) -> bool:
    headings = {_text(value): _text(value) for value in values if _text(value)}
    return all(_field(headings, FIELD_ALIASES[key]) for key in (
        "position_code", "employer", "title", "major", "degree", "headcount"
    ))


def _combined_header_rows(first: list[Any], second: list[Any]) -> list[str]:
    """Join adjacent multi-row headers while preserving physical columns."""
    width = max(len(first), len(second))
    return [
        _text(first[index]) if index < len(first) and _text(first[index])
        else _text(second[index]) if index < len(second)
        else ""
        for index in range(width)
    ]


def _sheet_rows(
    sheet_name: str,
    values: Iterable[list[Any]],
    *,
    header_row: int | None,
) -> tuple[list[dict[str, Any]], int | None]:
    detected_header = None
    headers: list[str] = []
    header_height = 1
    rows: list[dict[str, Any]] = []
    buffered = list(values)
    for physical_row, values_row in enumerate(buffered, start=1):
        if detected_header is None:
            if header_row is not None and physical_row != header_row:
                continue
            if header_row is None and physical_row > 30:
                break
            if _is_header(values_row):
                detected_header = physical_row
                headers = [_text(value) for value in values_row]
            elif header_row is None and physical_row < len(buffered):
                combined = _combined_header_rows(values_row, buffered[physical_row])
                if _is_header(combined):
                    detected_header = physical_row
                    headers = combined
                    header_height = 2
            continue
        if physical_row <= detected_header + header_height - 1:
            continue
        if not any(value not in (None, "") for value in values_row):
            continue
        row = {
            header: values_row[column] if column < len(values_row) else ""
            for column, header in enumerate(headers)
            if header
        }
        row["__sheet_name"] = sheet_name
        row["__row_number"] = physical_row
        rows.append(row)
    return rows, detected_header


def _flatten_extracted_rows(extracted_rows: Iterable[dict[str, Any]]) -> list[dict[str, Any]]:
    """Convert attachment-parser rows into the manifest parser's flat shape.

    The attachment pipeline keeps ``cells`` nested so it can preserve parser
    metadata.  The government-table normalizer uses header aliases directly;
    flattening here lets both paths share the same field contract while
    retaining sheet/page and physical row provenance.
    """
    rows: list[dict[str, Any]] = []
    for item in extracted_rows:
        if not isinstance(item, dict):
            continue
        cells = item.get("cells")
        if isinstance(cells, dict):
            row = {str(key): value for key, value in cells.items()}
        else:
            row = {
                str(key): value
                for key, value in item.items()
                if not str(key).startswith("__")
                and key not in {"sheet_name", "row_number", "row_kind", "row_text", "extraction_confidence"}
            }
        row["__sheet_name"] = str(item.get("sheet_name") or "document")
        row["__row_number"] = int(item.get("row_number") or len(rows) + 1)
        row["__extraction_confidence"] = str(item.get("extraction_confidence") or "unknown")
        rows.append(row)
    return rows


def _load_pdf_position_table(path: Path) -> dict[str, Any]:
    """Load text-layer PDF rows through the shared conservative PDF parser."""
    try:
        from pypdf import PdfReader
    except ImportError as error:
        raise ValueError("PDF position table requires pypdf") from error
    try:
        reader = PdfReader(str(path))
        page_texts: list[str] = []
        for page in reader.pages:
            try:
                text = page.extract_text(extraction_mode="layout") or ""
            except (TypeError, ValueError):
                text = page.extract_text() or ""
            page_texts.append(text)
    except Exception as error:  # pragma: no cover - parser-specific errors vary by PDF
        raise ValueError(f"PDF position table extraction failed: {error}") from error

    # Keep the same row boundary and conservative field inference as the
    # private attachment review queue.  Missing employer/headcount/etc. are
    # rejected later by parse_position_table_rows rather than inferred here.
    from job_hub.attachments import OfficialAttachmentProcessor

    extracted = OfficialAttachmentProcessor._rows_from_pdf_pages(page_texts)
    rows = _flatten_extracted_rows(extracted)
    sheets = [
        {
            "sheet_name": f"page-{index}",
            "header_row": None,
            "data_rows": sum(
                1 for row in rows if row.get("__sheet_name") == f"page-{index}"
            ),
            "parser": "pypdf",
        }
        for index in range(1, len(page_texts) + 1)
    ]
    return {
        "rows": rows,
        "sheets": sheets,
        "skipped_sheets": [],
        "metadata": {
            "parser": "pypdf",
            "page_count": len(page_texts),
            "extraction_mode": "layout_or_text",
            "extraction_confidence": "medium_or_low",
        },
    }


def _load_docx_position_table(path: Path) -> dict[str, Any]:
    """Load DOCX tables and preserve table/row provenance."""
    try:
        from docx import Document
    except ImportError as error:
        raise ValueError("DOCX position table requires python-docx") from error

    document = Document(str(path))
    extracted: list[dict[str, Any]] = []
    table_count = 0
    for table_index, table in enumerate(document.tables, start=1):
        if not table.rows:
            continue
        headers = [
            " ".join(cell.text.split()) or f"列{index}"
            for index, cell in enumerate(table.rows[0].cells, start=1)
        ]
        if not any(headers):
            continue
        table_count += 1
        for row_number, table_row in enumerate(table.rows[1:], start=2):
            values = [" ".join(cell.text.split()) for cell in table_row.cells]
            if not any(values):
                continue
            cells = {
                headers[index] if index < len(headers) else f"列{index + 1}": value
                for index, value in enumerate(values)
                if value
            }
            extracted.append(
                {
                    "sheet_name": f"table-{table_index}",
                    "row_number": row_number,
                    "cells": cells,
                    "extraction_confidence": "high",
                }
            )

    # Paragraph-only notices remain reviewable, but they will normally be
    # rejected by the required field contract until an operator maps them.
    if not table_count:
        for row_number, paragraph in enumerate(
            (item.text.strip() for item in document.paragraphs if item.text.strip()),
            start=1,
        ):
            extracted.append(
                {
                    "sheet_name": "docx",
                    "row_number": row_number,
                    "cells": {"正文": paragraph},
                    "extraction_confidence": "medium",
                }
            )

    rows = _flatten_extracted_rows(extracted)
    return {
        "rows": rows,
        "sheets": [
            {
                "sheet_name": f"table-{index}",
                "header_row": 1,
                "data_rows": sum(
                    1 for row in rows if row.get("__sheet_name") == f"table-{index}"
                ),
                "parser": "python-docx",
            }
            for index in range(1, table_count + 1)
        ] or [{"sheet_name": "docx", "header_row": None, "data_rows": len(rows), "parser": "python-docx"}],
        "skipped_sheets": [],
        "metadata": {
            "parser": "python-docx",
            "table_count": table_count,
            "extraction_mode": "tables" if table_count else "paragraphs",
        },
    }


def _load_image_position_table(path: Path) -> dict[str, Any]:
    """OCR an image into private review rows; never treat OCR as verified data."""
    try:
        completed = subprocess.run(
            ["tesseract", str(path), "stdout", "-l", "chi_sim+eng"],
            check=True,
            capture_output=True,
            timeout=120,
        )
    except (OSError, subprocess.CalledProcessError, subprocess.TimeoutExpired) as error:
        raise ValueError(f"image position table OCR unavailable: {error}") from error
    text = completed.stdout.decode("utf-8", errors="replace")
    from job_hub.attachments import OfficialAttachmentProcessor

    extracted = OfficialAttachmentProcessor._rows_from_pdf_pages([text])
    rows = _flatten_extracted_rows(
        [dict(row, extraction_confidence="low") for row in extracted]
    )
    return {
        "rows": rows,
        "sheets": [{"sheet_name": "image-1", "header_row": None, "data_rows": len(rows), "parser": "tesseract"}],
        "skipped_sheets": [],
        "metadata": {
            "parser": "tesseract",
            "extraction_mode": "ocr",
            "extraction_confidence": "low",
            "publication_policy": "review_only; OCR rows require independent official row verification",
        },
    }


def load_position_table_with_report(
    path: Path | str,
    *,
    sheet_name: str | None = None,
    header_row: int | None = None,
) -> dict[str, Any]:
    """Load CSV/JSON/Excel/PDF/DOCX/image tables with source-row provenance.

    PDF and OCR rows intentionally use the same conservative parser as the
    attachment review queue.  They remain review-only until every required
    field is independently verified against the official notice and artifact.
    """
    table_path = Path(path)
    suffix = table_path.suffix.casefold()
    if suffix == ".json":
        payload = json.loads(table_path.read_text(encoding="utf-8"))
        rows = payload.get("rows") if isinstance(payload, dict) else payload
        if not isinstance(rows, list) or any(not isinstance(item, dict) for item in rows):
            raise ValueError("JSON position table must be a list of row objects")
        return {"rows": rows, "sheets": [{"sheet_name": "JSON", "data_rows": len(rows)}], "skipped_sheets": []}
    if suffix == ".csv":
        with table_path.open("r", encoding="utf-8-sig", newline="") as handle:
            rows = [dict(row) for row in csv.DictReader(handle)]
        return {"rows": rows, "sheets": [{"sheet_name": "CSV", "data_rows": len(rows)}], "skipped_sheets": []}
    if suffix == ".pdf":
        return _load_pdf_position_table(table_path)
    if suffix == ".docx":
        return _load_docx_position_table(table_path)
    if suffix in _IMAGE_SUFFIXES:
        return _load_image_position_table(table_path)
    result: dict[str, Any] = {"rows": [], "sheets": [], "skipped_sheets": []}

    def add_sheet(name: str, values: Iterable[list[Any]]) -> None:
        rows, detected_header = _sheet_rows(name, values, header_row=header_row)
        if detected_header is None:
            result["skipped_sheets"].append({"sheet_name": name, "reason": "header_not_found"})
            return
        result["rows"].extend(rows)
        result["sheets"].append({"sheet_name": name, "header_row": detected_header, "data_rows": len(rows)})

    if suffix in {".xlsx", ".xlsm"}:
        from openpyxl import load_workbook

        workbook = load_workbook(table_path, read_only=False, data_only=True)
        try:
            if sheet_name and sheet_name not in workbook.sheetnames:
                raise ValueError(f"position table sheet not found: {sheet_name}")
            worksheets = [workbook[sheet_name]] if sheet_name else workbook.worksheets
            for worksheet in worksheets:
                merged_ranges = list(worksheet.merged_cells.ranges)

                def values_with_merged_cells() -> Iterable[list[Any]]:
                    for physical_row in range(1, worksheet.max_row + 1):
                        values = [
                            worksheet.cell(physical_row, column).value
                            for column in range(1, worksheet.max_column + 1)
                        ]
                        for merged in merged_ranges:
                            if merged.min_row <= physical_row <= merged.max_row:
                                source = worksheet.cell(merged.min_row, merged.min_col).value
                                for column in range(merged.min_col, merged.max_col + 1):
                                    values[column - 1] = source
                        yield values

                add_sheet(worksheet.title, values_with_merged_cells())
        finally:
            workbook.close()
    elif suffix == ".xls":
        import xlrd

        workbook = xlrd.open_workbook(table_path, on_demand=True)
        try:
            if sheet_name and sheet_name not in workbook.sheet_names():
                raise ValueError(f"position table sheet not found: {sheet_name}")
            worksheets = [workbook.sheet_by_name(sheet_name)] if sheet_name else workbook.sheets()
            for worksheet in worksheets:
                add_sheet(worksheet.name, (worksheet.row_values(index) for index in range(worksheet.nrows)))
        finally:
            workbook.release_resources()
    else:
        raise ValueError(
            "position table must be .csv, .json, .xlsx, .xlsm, .xls, .pdf, .docx or an OCR image"
        )

    if not result["sheets"]:
        raise ValueError("position table header not found in selected sheet(s)")
    return result


def load_position_table(
    path: Path | str, *, sheet_name: str | None = None, header_row: int | None = None
) -> list[dict[str, Any]]:
    """Compatibility wrapper returning rows; Excel defaults to all sheets."""
    return load_position_table_with_report(path, sheet_name=sheet_name, header_row=header_row)["rows"]


def parse_position_table_manifest(path: Path | str) -> dict[str, Any]:
    """Parse a manifest of official position-table files into review batches.

    The manifest is intentionally local and declarative: it may point at
    already downloaded official CSV/XLSX/XLS/JSON artifacts, but it never
    fetches URLs or changes the government registry. Relative file paths are
    resolved next to the manifest so a province-wide batch can be rerun from
    a versioned evidence directory.
    """

    manifest_path = Path(path)
    try:
        payload = json.loads(manifest_path.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError) as error:
        raise ValueError(f"cannot read position-table manifest: {manifest_path}") from error
    if not isinstance(payload, dict) or payload.get("version") != 1:
        raise ValueError("position-table manifest version must be 1")
    tables = payload.get("tables")
    if not isinstance(tables, list) or not tables:
        raise ValueError("position-table manifest requires a non-empty tables list")

    batches: list[dict[str, Any]] = []
    reports: list[dict[str, Any]] = []
    seen_batch_ids: set[str] = set()
    accepted_rows = rejected_rows = 0
    for index, raw in enumerate(tables):
        if not isinstance(raw, dict):
            raise ValueError(f"manifest table {index} must be an object")
        required = (
            "path",
            "batch_id",
            "source_id",
            "position_type",
            "province",
            "official_notice_url",
            "official_attachment_url",
        )
        missing = [key for key in required if not str(raw.get(key) or "").strip()]
        if missing:
            raise ValueError(f"manifest table {index} missing: {', '.join(missing)}")
        batch_id = str(raw["batch_id"]).strip()
        if batch_id in seen_batch_ids:
            raise ValueError(f"duplicate manifest batch_id: {batch_id}")
        seen_batch_ids.add(batch_id)
        table_path = Path(str(raw["path"]).strip())
        if not table_path.is_absolute():
            table_path = manifest_path.parent / table_path
        loaded = load_position_table_with_report(
            table_path,
            sheet_name=str(raw.get("sheet_name") or "").strip() or None,
            header_row=(int(raw["header_row"]) if raw.get("header_row") is not None else None),
        )
        parsed = parse_position_table_rows(
            loaded["rows"],
            batch_id=batch_id,
            source_id=str(raw["source_id"]).strip(),
            position_type=str(raw["position_type"]).strip(),
            province=str(raw["province"]).strip(),
            official_notice_url=str(raw["official_notice_url"]).strip(),
            official_attachment_url=str(raw["official_attachment_url"]).strip(),
            deadline_date=str(raw.get("deadline_date") or "").strip(),
            deadline_policy=str(raw.get("deadline_policy") or "fixed_date").strip(),
            location=str(raw.get("location") or "").strip(),
            location_evidence_policy=str(
                raw.get("location_evidence_policy") or "row_exact"
            ).strip(),
            sheet_name=str(raw.get("sheet_name") or "Sheet1").strip() or "Sheet1",
            first_data_row=int(raw.get("first_data_row") or 2),
        )
        batches.append(parsed["batch"])
        accepted_rows += int(parsed["metrics"]["accepted_rows"])
        rejected_rows += int(parsed["metrics"]["rejected_rows"])
        reports.append(
            {
                "batch_id": batch_id,
                "path": str(table_path),
                "load_report": {
                    "sheets": loaded["sheets"],
                    "skipped_sheets": loaded["skipped_sheets"],
                },
                "metrics": parsed["metrics"],
                "rejected_rows": parsed["rejected_rows"],
            }
        )
    return {
        "version": 1,
        "batches": batches,
        "reports": reports,
        "metrics": {
            "table_count": len(batches),
            "accepted_rows": accepted_rows,
            "rejected_rows": rejected_rows,
            "field_contract": "position_code, employer, title, major, degree, headcount",
        },
        "publication_policy": "review_only; no registry activation or student publication",
    }


__all__ = [
    "FIELD_ALIASES",
    "load_position_table",
    "load_position_table_with_report",
    "parse_position_table_rows",
    "parse_position_table_manifest",
]
