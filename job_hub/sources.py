from __future__ import annotations

import base64
import binascii
import hashlib
import io
import json
import re
import time
import xml.etree.ElementTree as ET
import zlib
from dataclasses import dataclass
from datetime import datetime, timezone
from html import unescape
from typing import Any
from urllib.parse import urlencode, urljoin, urlparse
from urllib.robotparser import RobotFileParser

import requests
from bs4 import BeautifulSoup

try:
    from cryptography.hazmat.primitives.ciphers import Cipher, algorithms, modes
    from cryptography.hazmat.primitives.padding import PKCS7
except ImportError:  # pragma: no cover - exercised only in minimal deployments
    Cipher = algorithms = modes = PKCS7 = None  # type: ignore[assignment]

from job_hub.config import Settings
from job_hub.contracts import ContractValidationError, validate_source_registry
from job_hub.locations import extract_location_hint
from job_hub.matching import (
    clean_text,
    extract_deadline,
    extract_published_date,
    extract_major_tags,
    looks_like_recruitment,
    normalize_url,
    parse_date_value,
    stable_hash,
)
from job_hub.transport import (
    RequestPolicy,
    ResponseCache,
    configure_session,
    create_session,
    request_exception_types,
)


REQUEST_ERRORS = request_exception_types()


USER_AGENT = (
    "CUPB-Geoscience-Employment-Information-Service/1.0 "
    "(official-public-source-crawler; contact: site-administrator)"
)

# Recruitment portals often leave assessment and appointment notices next to the
# original vacancy.  These are valuable to applicants who already applied, but
# they are not new employment opportunities and must never enter the public job
# corpus as active vacancies.  Match only announcement titles: a valid original
# vacancy may legitimately describe its written-test or interview process in
# the body text.
NON_VACANCY_TITLE_PATTERNS = (
    r"拟(?:聘|录用|聘用)",
    r"(?:进入|面试)(?:范围|名单)",
    r"递补",
    r"资格(?:审查|复审)",
    r"笔试(?:成绩|公告|结果)",
    r"面试(?:成绩|公告|结果)",
    r"体检(?:公告|名单|结果)",
    r"考察(?:公告|名单|结果)",
    r"录用(?:公示|名单|结果)",
)


class SourceCollectionError(RuntimeError):
    pass


class SourceSkipped(RuntimeError):
    pass


@dataclass(frozen=True)
class SourceHealthResult:
    status: str
    detail: str
    status_code: int | None = None
    successful: bool = False


class SourceHealthProbe:
    """Perform a light, compliant availability check for one public source.

    This is intentionally distinct from collection.  A healthy landing page is
    not proof that a source has matching vacancies, and a failed page must not
    be displayed as "no jobs" in province coverage.
    """

    def __init__(
        self,
        settings: Settings,
        session: requests.Session | None = None,
    ) -> None:
        self.settings = settings
        self.session = session or create_session(
            settings.http_transport_mode,
            headers={"User-Agent": USER_AGENT},
            client=settings.http_client,
        )
        if session is not None:
            configure_session(
                session,
                settings.http_transport_mode,
                headers={"User-Agent": USER_AGENT},
            )
        self.request_policy = RequestPolicy(
            self.session,
            timeout=settings.request_timeout_seconds,
            retries=settings.http_retry_attempts,
            backoff_seconds=settings.http_backoff_seconds,
            max_backoff_seconds=settings.http_max_backoff_seconds,
            jitter_seconds=settings.http_jitter_seconds,
            transport_mode=settings.http_transport_mode,
        )

    def check(self, source: dict[str, Any]) -> SourceHealthResult:
        if source.get("source_type") == "manual":
            return SourceHealthResult(
                "unknown",
                "人工补录来源不执行网络探测。",
            )
        entry_url = self._entry_url(source)
        parsed = urlparse(entry_url)
        root = f"{parsed.scheme}://{parsed.netloc}"
        robots_url = urljoin(root, "/robots.txt")
        try:
            robots_response = self.request_policy.request(
                "GET",
                robots_url,
                timeout=min(self.settings.request_timeout_seconds, 10),
                allow_redirects=True,
            )
        except REQUEST_ERRORS as error:
            return SourceHealthResult("source_degraded", f"robots.txt 无法核验：{error}")

        if robots_response.status_code == 404:
            robots_allowed = True
        elif robots_response.ok:
            parser = RobotFileParser()
            parser.parse(robots_response.text.splitlines())
            robots_allowed = parser.can_fetch(USER_AGENT, entry_url)
        else:
            return SourceHealthResult(
                "source_degraded",
                f"robots.txt 返回 HTTP {robots_response.status_code}，未访问来源页面。",
                robots_response.status_code,
            )
        if not robots_allowed:
            return SourceHealthResult(
                "source_blocked",
                "robots.txt 不允许本服务访问公开入口。",
                robots_response.status_code,
            )

        try:
            response = self.request_policy.request(
                "GET",
                entry_url,
                timeout=self.settings.request_timeout_seconds,
                allow_redirects=True,
            )
        except REQUEST_ERRORS as error:
            return SourceHealthResult("source_error", f"公开入口请求失败：{error}")
        if response.ok and self._is_soft_not_found(response):
            return SourceHealthResult(
                "source_degraded",
                "公开入口返回了站点的未找到页面，不能作为可用来源。",
                response.status_code,
            )
        if response.ok and not self._preserves_entry_path(entry_url, response.url, source):
            return SourceHealthResult(
                "source_degraded",
                "公开入口重定向后丢失了登记的栏目路径，不能作为可用来源。",
                response.status_code,
            )
        if response.ok:
            return SourceHealthResult(
                "source_active",
                "robots.txt 允许且登记的公开入口可访问；尚未代表有匹配岗位。",
                response.status_code,
                successful=True,
            )
        if response.status_code in {401, 403, 412, 429}:
            status = "source_blocked"
        elif response.status_code == 404:
            status = "source_degraded"
        else:
            status = "source_error"
        return SourceHealthResult(
            status,
            f"公开入口返回 HTTP {response.status_code}。",
            response.status_code,
        )

    @staticmethod
    def _entry_url(source: dict[str, Any]) -> str:
        """Probe the recruitment listing, not just an institution's home page."""
        config = source.get("config", {})
        configured = str(config.get("healthcheck_url") or "").strip()
        listing_urls = config.get("listing_urls") or []
        candidate = configured or (listing_urls[0] if listing_urls else source["homepage_url"])
        return normalize_url(str(candidate))

    @staticmethod
    def _preserves_entry_path(
        entry_url: str,
        resolved_url: str,
        source: dict[str, Any],
    ) -> bool:
        """Reject redirects from a registered recruitment column to a generic page."""
        if source.get("config", {}).get("require_path_stability") is False:
            return True
        expected = urlparse(entry_url)
        resolved = urlparse(resolved_url)
        expected_path = expected.path.rstrip("/")
        if not expected_path:
            return True
        if expected.hostname and resolved.hostname and expected.hostname.lower() != resolved.hostname.lower():
            return False
        return resolved.path.rstrip("/").startswith(expected_path)

    @staticmethod
    def _is_soft_not_found(response: requests.Response) -> bool:
        """Detect common HTTP-200 error pages returned by government sites."""
        path = urlparse(response.url).path.lower()
        if "/404/" in path or path.endswith("/404.html"):
            return True
        soup = BeautifulSoup(response.text[:160_000], "html.parser")
        title = clean_text(
            soup.title.get_text(" ", strip=True) if soup.title else ""
        ).lower()
        body = clean_text(soup.get_text(" ", strip=True))[:1_000].lower()
        signals = ("404", "页面不存在", "您访问的页面不存在", "找不到页面", "not found")
        return any(signal in title for signal in signals) or any(
            signal in body for signal in signals
        )


@dataclass(frozen=True)
class RawPosting:
    title: str
    employer: str
    source_url: str
    application_url: str | None
    text: str
    summary: str
    published_date: str | None
    deadline_date: str | None
    location: str | None
    external_id: str | None = None
    match_text: str | None = None
    official_evidence_url: str | None = None
    field_evidence: dict[str, str] | None = None
    qualification_text: str | None = None


class OfficialSourceCollector:
    """Fetches only explicitly configured, public and robots-permitted sources."""

    def __init__(
        self,
        settings: Settings,
        session: requests.Session | None = None,
    ) -> None:
        self.settings = settings
        headers = {
            "User-Agent": USER_AGENT,
            "Accept": "text/html,application/xhtml+xml,application/xml,"
            "application/json;q=0.9,*/*;q=0.5",
            "Accept-Language": "zh-CN,zh;q=0.9,en;q=0.6",
        }
        self.session = session or create_session(
            settings.http_transport_mode,
            headers=headers,
            client=settings.http_client,
        )
        if session is not None:
            configure_session(session, settings.http_transport_mode, headers=headers)
        cache = None
        if settings.http_cache_enabled:
            cache = ResponseCache(
                settings.data_dir / "http-cache",
                max_entries=settings.http_cache_max_entries,
            )
        self.request_policy = RequestPolicy(
            self.session,
            timeout=settings.request_timeout_seconds,
            retries=settings.http_retry_attempts,
            backoff_seconds=settings.http_backoff_seconds,
            max_backoff_seconds=settings.http_max_backoff_seconds,
            jitter_seconds=settings.http_jitter_seconds,
            cache=cache,
            transport_mode=settings.http_transport_mode,
        )
        self._robots_cache: dict[str, RobotFileParser] = {}

    def collect(self, source: dict[str, Any]) -> list[RawPosting]:
        source_type = source["source_type"]
        if source_type == "manual":
            return []
        if source_type == "rss":
            return self._collect_rss(source)
        if source_type == "json":
            return self._collect_json(source)
        if source_type == "cupb_career":
            return self._collect_cupb_career(source)
        if source_type == "cas_job_board":
            return self._collect_cas_job_board(source)
        if source_type == "successfactors_search":
            return self._collect_successfactors_search(source)
        if source_type == "mokahr_search":
            return self._collect_mokahr_search(source)
        if source_type == "zhaopin_campus":
            return self._collect_zhaopin_campus(source)
        if source_type == "mnr_recruitment":
            return self._collect_mnr_recruitment(source)
        if source_type == "slb_coveo_search":
            return self._collect_slb_coveo_search(source)
        if source_type == "structured_opening_page":
            return self._collect_structured_opening_page(source)
        if source_type == "official_xlsx_rows":
            return self._collect_official_xlsx_rows(source)
        if source_type == "official_table_rows":
            return self._collect_official_table_rows(source)
        if source_type in {"html_notice", "landing_page"}:
            return self._collect_html_notice(source)
        raise SourceCollectionError(f"Unsupported source type: {source_type}")

    def collect_with_fallback(self, source: dict[str, Any]) -> list[RawPosting]:
        """Try the registered official ladder for one source.

        A fallback is only a URL/configuration explicitly registered by the
        administrator.  It may be a unit's own announcement page or an
        official attachment route; arbitrary search results are never promoted
        into the ladder.  A successful empty scan is preserved as such unless
        ``fallback_on_empty`` is enabled for that source.
        """
        config = source.get("config", {})
        fallback_items = config.get("fallback_sources", []) or []
        candidates: list[dict[str, Any]] = [source]
        for item in fallback_items:
            if not isinstance(item, dict):
                continue
            url = str(item.get("url") or item.get("homepage_url") or "").strip()
            if not url:
                continue
            candidate = dict(source)
            candidate_config = dict(config)
            candidate_config.update(item.get("config", {}))
            candidate["homepage_url"] = url
            candidate["source_type"] = str(item.get("source_type") or source["source_type"])
            candidate["config"] = candidate_config
            candidates.append(candidate)

        failures: list[str] = []
        first_blocked: SourceSkipped | None = None
        for index, candidate in enumerate(candidates):
            try:
                postings = self.collect(candidate)
                if postings or index == 0 and not config.get("fallback_on_empty"):
                    return postings
                # A configured empty-scan ladder can still inspect a registered
                # alternate official page before concluding there are no matches.
                if index == len(candidates) - 1:
                    return postings
            except SourceSkipped as error:
                first_blocked = first_blocked or error
                failures.append(str(error))
            except SourceCollectionError as error:
                failures.append(str(error))
        if first_blocked is not None and all(
            "robots" in failure.lower() or "permit" in failure.lower()
            for failure in failures
        ):
            raise first_blocked
        detail = "; ".join(failures[-4:]) or "no registered official fallback succeeded"
        raise SourceCollectionError(f"Official source ladder exhausted: {detail}")

    def _collect_mnr_recruitment(self, source: dict[str, Any]) -> list[RawPosting]:
        """Collect public MNR recruitment rows through its own exposed API.

        The API powers the unauthenticated recruitment listing.  We only read
        announcement and position fields that the public web client itself
        requests; registration, user data, and any write endpoints are never
        called.
        """
        config = source["config"]
        api_base = str(config.get("api_base", "")).rstrip("/")
        public_detail_url = str(config.get("public_detail_url", "")).strip()
        if not api_base or not public_detail_url:
            raise SourceCollectionError("MNR source needs api_base and public_detail_url")
        announcement_limit = self._item_limit(source)
        self._wait(source)
        listing = self._post_json(
            f"{api_base}/Affiche/GetAfficheList",
            {
                "pageIndex": 1,
                "pageSize": announcement_limit,
                "type": int(config.get("notice_type", 0)),
            },
        )
        try:
            notices = listing.json().get("items", [])
        except (ValueError, AttributeError) as error:
            raise SourceCollectionError("MNR listing is not valid JSON") from error
        if not isinstance(notices, list):
            raise SourceCollectionError("MNR listing does not contain an items list")

        postings: list[RawPosting] = []
        seen_external_ids: set[str] = set()
        for notice in notices:
            if not isinstance(notice, dict):
                continue
            view_id = clean_text(str(notice.get("ViewId") or ""))
            title = clean_text(str(notice.get("Title") or ""))
            if not view_id or not title:
                continue
            self._wait(source)
            detail_response = self._post_json(
                f"{api_base}/Affiche/GetAfficheInfo",
                {"viewId": view_id, "type": int(config.get("notice_type", 0))},
            )
            try:
                detail_payload = detail_response.json()
            except ValueError as error:
                raise SourceCollectionError("MNR announcement detail is not valid JSON") from error
            detail = (
                detail_payload[0]
                if isinstance(detail_payload, list) and detail_payload
                else detail_payload
            )
            if not isinstance(detail, dict):
                continue
            announcement_url = f"{public_detail_url}?ViewId={view_id}"
            published_date = parse_date_value(
                str(detail.get("FbDate") or notice.get("FbDate") or "")
            )
            deadline_date = parse_date_value(str(detail.get("BmjsDate") or ""))
            announcement_text = clean_text(
                BeautifulSoup(str(detail.get("AnncCont") or ""), "html.parser").get_text(
                    " ", strip=True
                )
            )
            positions = self._mnr_positions(
                api_base,
                view_id,
                source,
                int(config.get("notice_type", 0)),
                announcement_limit,
            )
            if positions:
                for position in positions:
                    posting = self._mnr_position_posting(
                        position,
                        view_id=view_id,
                        announcement_title=title,
                        announcement_url=announcement_url,
                        published_date=published_date,
                        deadline_date=deadline_date,
                        source=source,
                    )
                    if posting is None or posting.external_id in seen_external_ids:
                        continue
                    seen_external_ids.add(str(posting.external_id))
                    postings.append(posting)
                    if len(postings) >= announcement_limit:
                        return postings
                continue
            if config.get("include_announcement_fallback", True) and self._accept_candidate(
                f"{title} {announcement_text}", source
            ):
                postings.append(
                    RawPosting(
                        title=title,
                        employer=clean_text(str(detail.get("Fbdw") or source["publisher"])),
                        source_url=announcement_url,
                        application_url=None,
                        text=announcement_text,
                        summary=announcement_text[:500],
                        published_date=published_date,
                        deadline_date=deadline_date,
                        location=None,
                        external_id=view_id,
                        match_text=f"{title} {announcement_text}",
                    )
                )
        return postings[:announcement_limit]

    def _mnr_positions(
        self,
        api_base: str,
        view_id: str,
        source: dict[str, Any],
        notice_type: int,
        item_limit: int,
    ) -> list[dict[str, Any]]:
        self._wait(source)
        response = self._post_json(
            f"{api_base}/Affiche/GetPostSelectFyList",
            {
                "viewId": view_id,
                "zpdw": "",
                "zpgw": "",
                "xwxlyq": "",
                "pageIndex": 1,
                "pageSize": min(item_limit, 100),
                "type": notice_type,
            },
        )
        try:
            items = response.json().get("items", [])
        except (ValueError, AttributeError) as error:
            raise SourceCollectionError("MNR position list is not valid JSON") from error
        return [item for item in items if isinstance(item, dict)]

    def _mnr_position_posting(
        self,
        position: dict[str, Any],
        *,
        view_id: str,
        announcement_title: str,
        announcement_url: str,
        published_date: str | None,
        deadline_date: str | None,
        source: dict[str, Any],
    ) -> RawPosting | None:
        employer = clean_text(str(position.get("zpdw") or source["publisher"]))
        role = clean_text(str(position.get("zpgw") or ""))
        if not role:
            return None
        details = " ".join(
            clean_text(str(position.get(field) or ""))
            for field in (
                "gwbh",
                "zpryfw",
                "zy",
                "xwxlyq",
                "gwyq",
                "yjfx",
                "gzdd",
                "remark",
            )
        )
        # Eligibility filtering must not use the employer name.  A finance role
        # at a geological institution is not a geology role merely because the
        # institution's name contains a geological keyword.
        eligibility_text = clean_text(f"{role} {details}")
        if not self._accept_candidate(eligibility_text, source):
            return None
        matching_text = clean_text(f"{employer} {eligibility_text}")
        external_position_id = clean_text(
            str(position.get("gwbh") or position.get("gwbm") or role)
        )
        return RawPosting(
            title=f"{employer} - {role}",
            employer=employer,
            source_url=announcement_url,
            application_url=None,
            text=clean_text(f"{announcement_title} {details}"),
            summary=clean_text(
                f"岗位：{role}；专业：{position.get('zy') or '未注明'}；"
                f"学历：{position.get('xwxlyq') or '未注明'}；"
                f"地点：{position.get('gzdd') or '未注明'}"
            ),
            published_date=published_date,
            deadline_date=deadline_date,
            location=clean_text(str(position.get("gzdd") or "")) or None,
            external_id=f"{view_id}:{external_position_id}",
            match_text=matching_text,
        )

    def _collect_slb_coveo_search(self, source: dict[str, Any]) -> list[RawPosting]:
        """Collect SLB's unauthenticated job search results and public details.

        The search page itself exposes a short-lived public search credential to
        its browser client.  It is read at runtime and never stored in source
        configuration or the database.  Registration, login and application
        endpoints are deliberately out of scope.
        """
        config = source["config"]
        listing_url = str(config.get("listing_url") or source["homepage_url"])
        listing_response = self._get(listing_url, source)
        listing_soup = BeautifulSoup(listing_response.text, "html.parser")
        organization_id = self._hidden_input_value(listing_soup, "organizationId")
        access_token = self._hidden_input_value(listing_soup, "accessToken")
        search_hub = self._hidden_input_value(listing_soup, "searchHub")
        source_name = self._hidden_input_value(listing_soup, "searchsource")
        if not organization_id or not access_token:
            raise SourceCollectionError(
                "SLB public job page is missing its browser search configuration"
            )

        api_base = str(
            config.get("search_api_url", "https://platform.cloud.coveo.com/rest/search/v2")
        ).rstrip("/")
        api_host = (urlparse(api_base).hostname or "").lower()
        allowed_api_hosts = {
            str(host).lower() for host in config.get("api_allowed_hosts", [])
        }
        if api_host not in allowed_api_hosts:
            raise SourceCollectionError("SLB search API host is not allowlisted")

        queries = [str(query).strip() for query in config.get("queries", []) if str(query).strip()]
        if not queries:
            raise SourceCollectionError("SLB source requires at least one query")
        item_limit = self._item_limit(source)
        page_size = max(1, min(int(config.get("query_page_size", 20)), 50))
        pipeline = str(config.get("pipeline", "ATSJobsPipeline"))
        fields = config.get(
            "fields_to_include",
            ["title", "date", "country", "city", "category", "jobposteddate"],
        )
        source_filter = f'@source=="{source_name}"' if source_name else ""
        search_url = f"{api_base}?{urlencode({'organizationId': organization_id})}"
        headers = {
            "Authorization": f"Bearer {access_token}",
            "Content-Type": "application/json",
            "Origin": f"{urlparse(listing_response.url).scheme}://{urlparse(listing_response.url).netloc}",
        }

        # A query can return hundreds of overlapping roles. Bound detail-page
        # requests before fetching any role, because the detail is required for
        # reliable location and qualification evidence.
        detail_limit = self._slb_detail_limit(source, item_limit)
        candidates: list[dict[str, Any]] = []
        seen_candidates: set[str] = set()
        for query in queries:
            self._wait(source)
            response = self._post_json(
                search_url,
                {
                    "q": query,
                    "numberOfResults": page_size,
                    "searchHub": search_hub or "CoveoJobsHub",
                    "pipeline": pipeline,
                    "cq": source_filter,
                    "fieldsToInclude": fields,
                },
                headers=headers,
            )
            try:
                results = response.json().get("results", [])
            except (ValueError, AttributeError) as error:
                raise SourceCollectionError("SLB public search did not return JSON") from error
            if not isinstance(results, list):
                raise SourceCollectionError("SLB public search did not return a result list")

            for item in results:
                if not isinstance(item, dict):
                    continue
                candidate_id = self._slb_candidate_id(item, source)
                if not candidate_id or candidate_id in seen_candidates:
                    continue
                seen_candidates.add(candidate_id)
                candidates.append(item)
                if len(candidates) >= detail_limit:
                    break
            if len(candidates) >= detail_limit:
                break

        postings: list[RawPosting] = []
        for item in candidates:
            posting = self._slb_posting_from_result(item, source)
            if posting is None:
                continue
            postings.append(posting)
            if len(postings) >= item_limit:
                break
        return postings

    def _slb_detail_limit(self, source: dict[str, Any], item_limit: int) -> int:
        """Keep public detail requests bounded even when searches overlap."""
        value = source["config"].get("max_detail_candidates", item_limit)
        try:
            return max(1, min(int(value), self.settings.max_source_items))
        except (TypeError, ValueError):
            return item_limit

    def _slb_candidate_id(
        self,
        item: dict[str, Any],
        source: dict[str, Any],
    ) -> str | None:
        """Return a safe unique key only for a relevant official result row."""
        config = source["config"]
        raw = item.get("raw") if isinstance(item.get("raw"), dict) else {}
        title = clean_text(str(item.get("title") or raw.get("title") or ""))
        source_url = clean_text(
            str(
                item.get("clickUri")
                or raw.get("clickableuri")
                or item.get("uri")
                or raw.get("uri")
                or ""
            )
        )
        if not title or not source_url or not self._title_matches_filters(title, config):
            return None
        source_host = (urlparse(source_url).hostname or "").lower()
        allowed_hosts = {
            str(host).lower() for host in config.get("allowed_hosts", [])
        }
        if source_host not in allowed_hosts:
            return None
        return clean_text(
            str(raw.get("sysurihash") or raw.get("urihash") or source_url)
        ) or None

    def _slb_posting_from_result(
        self,
        item: dict[str, Any],
        source: dict[str, Any],
    ) -> RawPosting | None:
        config = source["config"]
        raw = item.get("raw") if isinstance(item.get("raw"), dict) else {}
        title = clean_text(str(item.get("title") or raw.get("title") or ""))
        source_url = clean_text(
            str(
                item.get("clickUri")
                or raw.get("clickableuri")
                or item.get("uri")
                or raw.get("uri")
                or ""
            )
        )
        if not title or not source_url or not self._title_matches_filters(title, config):
            return None
        source_host = (urlparse(source_url).hostname or "").lower()
        if source_host not in {
            str(host).lower() for host in config.get("allowed_hosts", [])
        }:
            return None

        city = self._slb_text_value(raw.get("city"))
        country = self._slb_text_value(raw.get("country"))
        location = ", ".join(part for part in (city, country) if part) or None
        published_date = self._epoch_date(
            raw.get("jobposteddate") or raw.get("date")
        )
        external_id = clean_text(
            str(raw.get("sysurihash") or raw.get("urihash") or source_url)
        )

        if not config.get("fetch_detail_pages", True):
            return RawPosting(
                title=title,
                employer=source["publisher"],
                source_url=normalize_url(source_url),
                application_url=normalize_url(source_url),
                text=clean_text(f"{title} {location or ''}"),
                summary=clean_text(f"SLB 官方职位检索结果：{title}。"),
                published_date=published_date,
                deadline_date=None,
                location=location,
                external_id=external_id,
                match_text=title,
            )

        self._wait(source)
        try:
            detail_response = self._get(source_url, source)
        except (SourceSkipped, SourceCollectionError):
            return None
        return self._extract_slb_detail(
            detail_response.text,
            detail_response.url,
            source,
            title,
            location,
            published_date,
            external_id,
        )

    def _extract_slb_detail(
        self,
        document: str,
        source_url: str,
        source: dict[str, Any],
        title_hint: str,
        location_hint: str | None,
        published_date: str | None,
        external_id: str,
    ) -> RawPosting | None:
        soup = BeautifulSoup(document, "html.parser")
        page_text = clean_text(soup.get_text(" ", strip=True))
        role_match = re.search(
            r"Job Name:\s*(.+?)(?=\s+(?:City|Country|Nationality|Job Summary|Requirements|Responsibilities):)",
            page_text,
            re.IGNORECASE,
        )
        title = clean_text(role_match.group(1)) if role_match else title_hint
        if len(page_text) < int(source["config"].get("minimum_detail_characters", 250)):
            return None
        if not self._title_matches_filters(title, source["config"]):
            return None
        if (
            source["config"].get("require_location_evidence", True)
            and location_hint
            and not self._slb_location_is_evidenced(location_hint, page_text)
        ):
            # Some generic early-career records share a role page across countries.
            # A page that describes a different country must not be labeled with the
            # search result's country on the student-facing site.
            return None

        city_match = re.search(
            r"City:\s*(.+?)(?=\s+(?:Nationality|Job Summary|Requirements|Responsibilities):)",
            page_text,
            re.IGNORECASE,
        )
        detail_location = clean_text(city_match.group(1)) if city_match else None
        location = detail_location or location_hint
        summary_start = max(
            page_text.lower().find("job summary:"),
            page_text.lower().find("requirements:"),
        )
        summary = page_text[summary_start : summary_start + 500] if summary_start >= 0 else page_text[:500]
        return RawPosting(
            title=title,
            employer=source["publisher"],
            source_url=normalize_url(source_url),
            application_url=self._find_application_url(soup, source_url)
            or normalize_url(source_url),
            text=page_text,
            summary=clean_text(summary),
            published_date=published_date,
            deadline_date=extract_deadline(page_text),
            location=location,
            external_id=external_id,
            match_text=clean_text(f"{title} {page_text}"),
        )

    @staticmethod
    def _hidden_input_value(soup: BeautifulSoup, input_id: str) -> str:
        node = soup.select_one(f"input#{input_id}")
        return clean_text(str(node.get("value") or "")) if node else ""

    @staticmethod
    def _slb_text_value(value: Any) -> str:
        if isinstance(value, list):
            return clean_text(", ".join(str(item) for item in value if str(item).strip()))
        return clean_text(str(value or ""))

    @staticmethod
    def _epoch_date(value: Any) -> str | None:
        try:
            numeric = float(value)
            if numeric <= 0:
                return None
            return datetime.fromtimestamp(numeric / 1000, tz=timezone.utc).date().isoformat()
        except (TypeError, ValueError, OverflowError, OSError):
            return None

    @staticmethod
    def _slb_location_is_evidenced(location: str, page_text: str) -> bool:
        expected = [
            part.strip().lower()
            for part in location.split(",")
            if part.strip() and part.strip().lower() not in {"multi-location", "multiple locations"}
        ]
        return not expected or any(part in page_text.lower() for part in expected)

    def _collect_html_notice(self, source: dict[str, Any]) -> list[RawPosting]:
        config = source["config"]
        listing_urls = config.get("listing_urls") or [source["homepage_url"]]
        links: list[tuple[str, str]] = []
        seen_links: set[str] = set()
        item_limit = self._item_limit(source)

        # Some official units publish a single announcement page without a
        # listing page.  The URL is explicitly registered by an administrator;
        # it is never discovered from a search engine or a third-party index.
        for direct_url in config.get("direct_notice_urls", []) or []:
            normalized = normalize_url(str(direct_url))
            if normalized and normalized not in seen_links:
                seen_links.add(normalized)
                links.append((str(config.get("direct_title_hint") or ""), normalized))

        if not config.get("direct_only"):
            for listing_url in listing_urls:
                response = self._get(listing_url, source)
                soup = BeautifulSoup(response.text, "html.parser")
                for title_hint, detail_url in self._notice_links(
                    soup, response.url, source
                ):
                    if detail_url in seen_links:
                        continue
                    seen_links.add(detail_url)
                    links.append((title_hint, detail_url))
                    if len(links) >= item_limit:
                        break
                if len(links) >= item_limit:
                    break

        postings: list[RawPosting] = []
        for title_hint, detail_url in links:
            try:
                detail_response = self._get(detail_url, source)
                if source["config"].get("split_role_lines"):
                    postings.extend(
                        self._extract_line_split_details(
                            detail_response.text,
                            detail_response.url,
                            source,
                            title_hint,
                        )
                    )
                elif source["config"].get("split_role_blocks"):
                    postings.extend(
                        self._extract_role_split_details(
                            detail_response.text,
                            detail_response.url,
                            source,
                            title_hint,
                        )
                    )
                else:
                    posting = self._extract_html_detail(
                        detail_response.text,
                        detail_response.url,
                        source,
                        title_hint,
                    )
                    if posting:
                        postings.append(posting)
            except SourceSkipped:
                continue
            except (*REQUEST_ERRORS, SourceCollectionError):
                continue
            self._wait(source)
        return postings

    def _extract_line_split_details(
        self,
        document: str,
        source_url: str,
        source: dict[str, Any],
        title_hint: str,
    ) -> list[RawPosting]:
        """Split a verified announcement whose roles are separated by ``<br>``.

        A number of government CMS pages do not use table rows or paragraphs:
        the whole notice is one text container with numbered role headings.  A
        configured heading expression is required, and each resulting record
        keeps the complete section text as its professional evidence.  This is
        deliberately stricter than guessing every line containing “岗位”.
        """
        config = source["config"]
        soup = BeautifulSoup(document, "html.parser")
        content_node = self._select_first(
            soup,
            str(config.get("role_content_selector") or config.get("content_selector") or "body"),
        )
        if content_node is None:
            raise SourceCollectionError("Line-split notice has no configured content container")
        for selector in config.get("remove_selectors", []):
            for node in content_node.select(selector):
                node.decompose()
        lines = [clean_text(line) for line in content_node.get_text("\n", strip=True).splitlines()]
        lines = [line for line in lines if line]
        role_pattern = str(config.get("role_line_pattern") or "").strip()
        if not role_pattern:
            raise SourceCollectionError("Line-split notice requires role_line_pattern")
        try:
            heading = re.compile(role_pattern, re.IGNORECASE)
        except re.error as error:
            raise SourceCollectionError(f"Invalid role_line_pattern: {error}") from error
        role_indexes = [index for index, line in enumerate(lines) if heading.search(line)]
        if not role_indexes:
            raise SourceCollectionError("Line-split notice exposed no matching role headings")

        title_node = self._select_first(soup, str(config.get("title_selector") or "h1, title"))
        page_title = clean_text(
            title_node.get_text(" ", strip=True) if title_node is not None else title_hint
        )
        page_text = clean_text(" ".join(lines))
        published_date = (
            str(config.get("published_date") or "").strip()
            or self._published_date_from_meta(soup)
            or extract_published_date(page_text)
        )
        source_url = normalize_url(source_url)
        postings: list[RawPosting] = []
        section_end_pattern = str(
            config.get("role_section_end_pattern") or r"^\s*[一二三四五六七八九十]+、"
        )
        try:
            section_heading = re.compile(section_end_pattern)
        except re.error as error:
            raise SourceCollectionError(
                f"Invalid role_section_end_pattern: {error}"
            ) from error
        for position, start in enumerate(role_indexes, start=1):
            end = role_indexes[position] if position < len(role_indexes) else len(lines)
            for candidate in range(start + 1, end):
                if section_heading.search(lines[candidate]):
                    end = candidate
                    break
            role_lines = lines[start:end]
            role_text = clean_text(" ".join(role_lines))
            role_excludes = [
                str(pattern).strip()
                for pattern in config.get("role_exclude_patterns", [])
                if str(pattern).strip()
            ]
            if any(re.search(pattern, role_text, re.IGNORECASE) for pattern in role_excludes):
                continue
            role_includes = [
                str(pattern).strip()
                for pattern in config.get("role_include_patterns", [])
                if str(pattern).strip()
            ]
            if role_includes and not any(
                re.search(pattern, role_text, re.IGNORECASE)
                for pattern in role_includes
            ):
                continue
            if not self._accept_candidate(role_text, source):
                continue
            raw_title = role_lines[0]
            title = re.sub(
                str(config.get("role_line_prefix_pattern") or r"^\s*\d+[、.．)]\s*"),
                "",
                raw_title,
            ).strip()
            postings.append(
                RawPosting(
                    title=title or raw_title,
                    employer=clean_text(str(config.get("employer_hint") or source["publisher"])),
                    source_url=source_url,
                    application_url=self._find_application_url(soup, source_url),
                    text=clean_text(f"{page_title} {role_text}"),
                    summary=role_text[:420],
                    published_date=published_date,
                    deadline_date=extract_deadline(page_text),
                    location=clean_text(str(config.get("location_hint") or "")) or None,
                    external_id=(
                        f"{self._external_id_from_url(source_url)}#role-{position}-"
                        f"{stable_hash(title or raw_title)[:16]}"
                    ),
                    match_text=role_text,
                    official_evidence_url=source_url,
                )
            )
            if len(postings) >= self._item_limit(source):
                break
        return postings

    def _collect_official_xlsx_rows(self, source: dict[str, Any]) -> list[RawPosting]:
        """Read reviewed rows from an official XLSX or legacy XLS file.

        This adapter is intentionally hash-locked and row-allowlisted.  It is
        not a general spreadsheet importer: if the official file changes, the
        source fails closed until its hash and reviewed row list are updated.
        OOXML workbooks use ``openpyxl``; legacy BIFF workbooks use ``xlrd``.
        """
        config = source["config"]
        page_url = normalize_url(str(config.get("notice_url") or source["homepage_url"]))
        attachment_url = normalize_url(str(config.get("attachment_url") or ""))
        if not attachment_url:
            raise SourceCollectionError("official_xlsx_rows requires attachment_url")
        response = self._get(attachment_url, source)
        content = bytes(getattr(response, "content", b""))
        expected_hash = str(config.get("sha256") or "").strip().lower()
        actual_hash = hashlib.sha256(content).hexdigest()
        if expected_hash and actual_hash != expected_hash:
            raise SourceCollectionError(
                "Official attachment hash changed; review the new file before enabling rows"
            )
        allowed_rows = {
            int(value)
            for value in (config.get("verified_rows") or [])
            if str(value).strip().isdigit()
        }
        if not allowed_rows:
            raise SourceCollectionError("official_xlsx_rows requires verified_rows")
        sheet_name = str(config.get("sheet_name") or "").strip()

        workbook = None
        worksheet = None
        legacy_book = None
        legacy_sheet = None
        row_iterator: Any
        try:
            try:
                from openpyxl import load_workbook

                workbook = load_workbook(
                    io.BytesIO(content), read_only=True, data_only=True
                )
                worksheet = (
                    workbook[sheet_name]
                    if sheet_name and sheet_name in workbook.sheetnames
                    else workbook.active
                )
                row_iterator = worksheet.iter_rows(values_only=True)
            except Exception as openpyxl_error:
                # A number of government portals label BIFF8 workbooks as
                # .xlsx or expose them with an .xls suffix.  Fall back only
                # after the OOXML reader rejects the bytes; never reinterpret
                # an arbitrary HTML error page as a spreadsheet.
                try:
                    import xlrd

                    legacy_book = xlrd.open_workbook(
                        file_contents=content,
                        on_demand=True,
                    )
                    if sheet_name and sheet_name in legacy_book.sheet_names():
                        legacy_sheet = legacy_book.sheet_by_name(sheet_name)
                    else:
                        legacy_sheet = legacy_book.sheet_by_index(0)
                    row_iterator = (
                        legacy_sheet.row_values(row_number)
                        for row_number in range(legacy_sheet.nrows)
                    )
                except ImportError as error:
                    raise SourceCollectionError(
                        "official_xlsx_rows requires openpyxl or xlrd"
                    ) from error
                except Exception as legacy_error:
                    raise SourceCollectionError(
                        "Official spreadsheet extraction failed: "
                        f"OOXML={openpyxl_error}; legacy={legacy_error}"
                    ) from legacy_error
        except SourceCollectionError:
            raise
        except Exception as error:
            raise SourceCollectionError(
                f"Official spreadsheet extraction failed: {error}"
            ) from error

        postings: list[RawPosting] = []
        try:
            for row_number, values in enumerate(row_iterator, start=1):
                if row_number not in allowed_rows:
                    continue
                cells = [clean_text(str(value)) if value is not None else "" for value in values]
                # The reviewed row contract is positional and versioned with
                # the file hash, so a column shift cannot silently publish data.
                title_index = int(config.get("title_column", 2))
                duties_index = int(config.get("duties_column", 3))
                requirements_index = int(config.get("requirements_column", 4))
                quantity_index = int(config.get("quantity_column", 5))
                employer_index = int(config.get("employer_column", 0))
                location_index = int(config.get("location_column", 0))
                title = cells[title_index - 1] if len(cells) >= title_index else ""
                duties = cells[duties_index - 1] if len(cells) >= duties_index else ""
                requirements = cells[requirements_index - 1] if len(cells) >= requirements_index else ""
                quantity = cells[quantity_index - 1] if len(cells) >= quantity_index else ""
                employer = (
                    cells[employer_index - 1]
                    if employer_index > 0 and len(cells) >= employer_index
                    else ""
                )
                location = (
                    cells[location_index - 1]
                    if location_index > 0 and len(cells) >= location_index
                    else ""
                )
                evidence_text = clean_text(" ".join(filter(None, (title, duties, requirements))))
                if not title or not self._accept_candidate(evidence_text, source):
                    continue
                postings.append(
                    RawPosting(
                        title=title,
                        employer=clean_text(
                            employer
                            or str(config.get("employer_hint") or source["publisher"])
                        ),
                        source_url=page_url,
                        application_url=None,
                        text=clean_text(
                            f"{evidence_text} 招聘人数：{quantity}；官方附件：{attachment_url}"
                        ),
                        summary=evidence_text[:420],
                        published_date=str(config.get("published_date") or "") or None,
                        deadline_date=str(config.get("deadline_date") or "") or None,
                        location=clean_text(
                            location or str(config.get("location_hint") or "")
                        ) or None,
                        external_id=f"{self._external_id_from_url(attachment_url)}#row-{row_number}",
                        match_text=evidence_text,
                        official_evidence_url=attachment_url,
                    )
                )
                if len(postings) >= self._item_limit(source):
                    break
        finally:
            if workbook is not None:
                workbook.close()
            if legacy_book is not None:
                legacy_book.release_resources()
        return postings

    def _collect_official_table_rows(self, source: dict[str, Any]) -> list[RawPosting]:
        """Publish explicitly evidenced rows from an official HTML table.

        State-owned research institutes frequently publish a single official
        announcement containing several tables of positions.  Treating that
        page as one vacancy loses the degree, location and role evidence for
        every row.  This adapter is deliberately configuration-driven: it
        visits only registered notice pages, maps named header columns, and
        requires professional evidence in each row before creating a posting.
        It never guesses rows from arbitrary prose or from a search result.
        """
        config = source["config"]
        item_limit = self._item_limit(source)
        links: list[tuple[str, str]] = []
        seen_links: set[str] = set()
        for raw_url in config.get("direct_notice_urls", []) or []:
            detail_url = normalize_url(str(raw_url))
            if detail_url and detail_url not in seen_links:
                seen_links.add(detail_url)
                links.append((str(config.get("direct_title_hint") or ""), detail_url))

        if not config.get("direct_only"):
            for listing_url in config.get("listing_urls", []) or []:
                response = self._get(str(listing_url), source)
                soup = BeautifulSoup(response.text, "html.parser")
                for title_hint, detail_url in self._notice_links(
                    soup, response.url, source
                ):
                    if detail_url in seen_links:
                        continue
                    seen_links.add(detail_url)
                    links.append((title_hint, detail_url))
                    if len(links) >= int(config.get("max_notice_pages", 20)):
                        break
                if len(links) >= int(config.get("max_notice_pages", 20)):
                    break

        postings: list[RawPosting] = []
        for title_hint, detail_url in links:
            try:
                response = self._get(detail_url, source)
                postings.extend(
                    self._extract_official_table_rows(
                        response.text,
                        response.url,
                        source,
                        title_hint,
                    )
                )
            except SourceSkipped:
                continue
            except (*REQUEST_ERRORS, SourceCollectionError):
                continue
            if len(postings) >= item_limit:
                return postings[:item_limit]
            self._wait(source)
        return postings[:item_limit]

    def _extract_official_table_rows(
        self,
        document: str,
        source_url: str,
        source: dict[str, Any],
        title_hint: str,
    ) -> list[RawPosting]:
        config = source["config"]
        soup = BeautifulSoup(document, "html.parser")
        title_node = self._select_first(
            soup,
            str(config.get("title_selector") or "h1, title"),
        )
        page_title = clean_text(
            title_node.get_text(" ", strip=True)
            if title_node is not None
            else title_hint
        )
        if self._is_non_vacancy_notice_title(page_title):
            return []
        content_node = self._select_first(
            soup,
            str(config.get("content_selector") or "article, main, body"),
        ) or soup
        for selector in config.get("remove_selectors", []):
            for node in content_node.select(str(selector)):
                node.decompose()
        page_text = clean_text(content_node.get_text(" ", strip=True))
        if not self._accept_candidate(
            f"{page_title} {page_text}",
            source,
            exclude_patterns=[],
            require_recruitment_word=False,
        ) and config.get("require_page_recruitment_word", True):
            raise SourceCollectionError(
                "Official table page does not contain recruitment evidence"
            )
        published_date = (
            str(config.get("published_date") or "").strip()
            or self._published_date_from_meta(soup)
            or extract_published_date(page_text)
        )
        deadline_date = (
            str(config.get("deadline_date") or "").strip()
            or extract_deadline(page_text)
        )
        contexts: dict[int, dict[str, str]] = {}
        for item in config.get("table_contexts", []) or []:
            if not isinstance(item, dict):
                continue
            try:
                index = int(item.get("table_index"))
            except (TypeError, ValueError):
                continue
            contexts[index] = {
                "employer": clean_text(str(item.get("employer") or "")),
                "location": clean_text(str(item.get("location") or "")),
            }

        aliases = {
            "title": ("招聘岗位", "需求岗位", "岗位名称", "岗位", "职位"),
            "quantity": ("招聘人数", "需求人数", "人数", "数量"),
            "degree": ("学历要求", "学历", "学位"),
            "major": ("需求专业", "专业要求", "专业", "专业方向"),
            "location": ("工作地点", "工作城市", "工作地区", "地点"),
            "employer": ("招聘单位", "用人单位", "单位", "需求单位"),
        }
        configured_aliases = config.get("column_aliases", {})
        for field, values in configured_aliases.items():
            if isinstance(values, list) and values:
                aliases[str(field)] = tuple(clean_text(str(value)) for value in values)

        tables = content_node.select(str(config.get("table_selector") or "table"))
        if not tables:
            raise SourceCollectionError("Official table page exposes no configured tables")
        postings: list[RawPosting] = []
        for table_index, table in enumerate(tables, start=1):
            rows: list[list[str]] = []
            for row in table.select("tr"):
                cells = [clean_text(cell.get_text(" ", strip=True)) for cell in row.select("th, td")]
                if any(cells):
                    rows.append(cells)
            if len(rows) < 2:
                continue
            header_index = self._table_header_index(rows, aliases)
            if header_index is None:
                continue
            header = rows[header_index]
            column_indexes = self._table_column_indexes(header, aliases)
            context = contexts.get(table_index, {})
            for row_number, cells in enumerate(rows[header_index + 1 :], start=header_index + 2):
                fields = {
                    field: self._table_cell(cells, index)
                    for field, index in column_indexes.items()
                }
                title = clean_text(fields.get("title", ""))
                if not title or self._is_non_vacancy_notice_title(title):
                    continue
                row_evidence = clean_text(
                    " ".join(
                        value
                        for value in (
                            title,
                            fields.get("major", ""),
                            fields.get("degree", ""),
                            fields.get("location", ""),
                        )
                        if value
                    )
                )
                row_include_patterns = [
                    str(pattern).strip()
                    for pattern in config.get("row_include_patterns", [])
                    if str(pattern).strip()
                ]
                row_exclude_patterns = [
                    str(pattern).strip()
                    for pattern in config.get("row_exclude_patterns", [])
                    if str(pattern).strip()
                ]
                if any(
                    re.search(pattern, row_evidence, re.IGNORECASE)
                    for pattern in row_exclude_patterns
                ):
                    continue
                explicit_major_match = bool(
                    row_include_patterns
                    and any(
                        re.search(pattern, row_evidence, re.IGNORECASE)
                        for pattern in row_include_patterns
                    )
                )
                if (
                    config.get("require_major_match", True)
                    and not extract_major_tags(row_evidence)
                    and not explicit_major_match
                ):
                    continue
                if not self._accept_candidate(
                    clean_text(f"{page_title} 招聘 {row_evidence}"),
                    source,
                    exclude_patterns=[],
                    require_recruitment_word=False,
                    require_major_match=(
                        bool(config.get("require_major_match", True))
                        and not explicit_major_match
                    ),
                ):
                    continue
                employer = clean_text(fields.get("employer", "")) or context.get("employer")
                employer = employer or clean_text(
                    str(config.get("employer_hint") or source["publisher"])
                )
                location = clean_text(fields.get("location", "")) or context.get("location")
                evidence_text = clean_text(
                    "；".join(
                        part
                        for part in (
                            f"岗位：{title}",
                            f"人数：{fields.get('quantity')}" if fields.get("quantity") else "",
                            f"学历：{fields.get('degree')}" if fields.get("degree") else "",
                            f"专业：{fields.get('major')}" if fields.get("major") else "",
                            f"地点：{location}" if location else "",
                        )
                        if part
                    )
                )
                postings.append(
                    RawPosting(
                        title=title,
                        employer=employer,
                        source_url=normalize_url(source_url),
                        application_url=self._find_application_url(soup, source_url),
                        text=clean_text(f"{page_title} {evidence_text} {page_text}"),
                        summary=evidence_text[:420],
                        published_date=published_date,
                        deadline_date=deadline_date,
                        location=location or None,
                        external_id=(
                            f"{self._external_id_from_url(source_url)}"
                            f"#table-{table_index}-row-{row_number}"
                        ),
                        match_text=clean_text(f"{title} {fields.get('major', '')} {fields.get('degree', '')}"),
                        official_evidence_url=normalize_url(source_url),
                    )
                )
        return postings

    @staticmethod
    def _table_header_index(
        rows: list[list[str]],
        aliases: dict[str, tuple[str, ...]],
    ) -> int | None:
        required = ("title",)
        for index, row in enumerate(rows[:4]):
            joined = " ".join(row)
            if all(
                any(alias in joined for alias in aliases[field])
                for field in required
            ):
                return index
        return None

    @staticmethod
    def _table_column_indexes(
        header: list[str],
        aliases: dict[str, tuple[str, ...]],
    ) -> dict[str, int]:
        indexes: dict[str, int] = {}
        for field, names in aliases.items():
            for index, cell in enumerate(header):
                if any(name in cell for name in names):
                    indexes[field] = index
                    break
        return indexes

    @staticmethod
    def _table_cell(cells: list[str], index: int | None) -> str:
        if index is None or index < 0 or index >= len(cells):
            return ""
        return clean_text(cells[index])

    def _extract_role_split_details(
        self,
        document: str,
        source_url: str,
        source: dict[str, Any],
        title_hint: str,
    ) -> list[RawPosting]:
        """Split one official notice into independently searchable vacancies.

        Many state-owned units publish a single HTML notice containing a list of
        roles instead of one page per role.  Treating that page as one vacancy
        loses the unit, degree and major evidence for every other role.  This
        adapter only splits explicitly configured paragraph blocks; it does not
        guess rows from arbitrary prose.  A source is enabled only after its
        selectors and role-heading pattern have been verified against a fixture.
        """
        soup = BeautifulSoup(document, "html.parser")
        config = source["config"]
        content_selector = str(
            config.get("role_content_selector")
            or config.get("content_selector")
            or "article, main, body"
        )
        content_node = self._select_first(soup, content_selector)
        if content_node is None:
            raise SourceCollectionError(
                "Role-split notice has no configured content container"
            )
        for selector in config.get("remove_selectors", []):
            for node in content_node.select(selector):
                node.decompose()

        block_selector = str(config.get("role_block_selector") or "p")
        blocks = [
            clean_text(node.get_text(" ", strip=True))
            for node in content_node.select(block_selector)
        ]
        blocks = [block for block in blocks if block]
        if not blocks:
            raise SourceCollectionError(
                "Role-split notice exposed no non-empty configured blocks"
            )

        role_pattern = str(config.get("role_title_pattern") or "").strip()
        if not role_pattern:
            raise SourceCollectionError("Role-split notice requires role_title_pattern")
        try:
            role_heading = re.compile(role_pattern, re.IGNORECASE)
        except re.error as error:
            raise SourceCollectionError(
                f"Invalid role_title_pattern: {error}"
            ) from error
        role_indexes = [
            index for index, block in enumerate(blocks) if role_heading.search(block)
        ]
        if not role_indexes:
            raise SourceCollectionError(
                "Role-split notice exposed no matching role headings"
            )

        page_text = clean_text(" ".join(blocks))
        page_title_node = self._select_first(
            soup,
            str(config.get("title_selector") or "h1, title"),
        )
        page_title = clean_text(
            page_title_node.get_text(" ", strip=True)
            if page_title_node is not None
            else title_hint
        )
        published_date = self._published_date_from_meta(soup) or extract_published_date(
            page_text
        )
        source_url = normalize_url(source_url)
        include_patterns = [
            str(pattern).strip()
            for pattern in config.get("role_include_patterns", [])
            if str(pattern).strip()
        ]
        postings: list[RawPosting] = []
        for position, start in enumerate(role_indexes, start=1):
            end = role_indexes[position] if position < len(role_indexes) else len(blocks)
            role_blocks = blocks[start:end]
            role_heading_text = role_blocks[0]
            role_text = clean_text(" ".join(role_blocks))
            if include_patterns and not any(
                re.search(pattern, role_text, re.IGNORECASE)
                for pattern in include_patterns
            ):
                continue
            combined = clean_text(f"{page_title} {role_text}")
            if not self._accept_candidate(combined, source):
                continue
            title = re.sub(
                str(config.get("role_title_prefix_pattern") or r"^\s*\d+[、.．)]\s*"),
                "",
                role_heading_text,
            ).strip()
            if not title:
                title = role_heading_text
            unit = ""
            for role_block in role_blocks:
                if "工作单位" not in role_block:
                    continue
                unit = clean_text(
                    re.sub(r"^.*?工作单位\s*[:：]\s*", "", role_block)
                )
                # Some government CMS pages concatenate the next section into
                # the same paragraph. Keep only the value of the labelled
                # field; otherwise salary/benefit prose becomes part of the
                # employer name shown to students.
                unit = re.split(
                    r"\s*(?:[一二三四五六七八九十]+、|\d+[、.．)]\s*)?"
                    r"(?:薪资待遇|福利待遇|报名方式|联系方式|招聘程序)\b",
                    unit,
                    maxsplit=1,
                )[0].strip(" ：:；;")
                break
            employer = clean_text(
                str(config.get("employer_hint") or source["publisher"])
            )
            if unit and config.get("append_unit_to_employer", True):
                employer = f"{employer} - {unit}"
            location = config.get("location_hint")
            if not location:
                location = self._label_value(
                    role_text,
                    "工作地点",
                    ("岗位职责", "任职要求", "学历要求", "工作单位", "报名"),
                )
            postings.append(
                RawPosting(
                    title=title,
                    employer=employer,
                    source_url=source_url,
                    application_url=self._find_application_url(soup, source_url),
                    text=combined,
                    summary=role_text[:420],
                    published_date=published_date,
                    deadline_date=extract_deadline(role_text),
                    location=clean_text(str(location or "")) or None,
                    external_id=(
                        f"{self._external_id_from_url(source_url)}#role-{position}-"
                        f"{stable_hash(title)[:16]}"
                    ),
                    match_text=role_text,
                    official_evidence_url=source_url,
                )
            )
            if len(postings) >= self._item_limit(source):
                break
        return postings

    def _collect_cupb_career(self, source: dict[str, Any]) -> list[RawPosting]:
        """Collect only concrete CUPB vacancy and recruitment-announcement pages."""
        config = source["config"]
        listing_urls = config.get("listing_urls") or [source["homepage_url"]]
        detail_patterns = config.get(
            "detail_path_patterns",
            [r"/(?:campus|job)/view/", r"/news/view/.+tag/xwzp"],
        )
        item_limit = self._item_limit(source)
        # A listing page is usually ordered by publication time, not by
        # relevance to geoscience.  Do not stop after the first few cards: the
        # first page can contain banking, teaching, events and energy postings
        # in an arbitrary mix.  Collect a bounded candidate window, then apply
        # the professional/official-content gates on each detail page.
        try:
            candidate_limit = max(
                item_limit,
                min(int(config.get("candidate_limit", item_limit * 4)), 600),
            )
        except (TypeError, ValueError):
            candidate_limit = min(max(item_limit, item_limit * 4), 600)
        candidates: list[tuple[str, str]] = []
        seen: set[str] = set()

        # Keep a small, administrator-reviewed set of high-value official
        # detail pages as a resilience layer.  The portal homepage is still
        # the primary discovery route; these URLs prevent a homepage redesign
        # or short pagination window from hiding a known official notice.
        for item in config.get("direct_notice_urls", []) or []:
            if isinstance(item, dict):
                raw_url = str(item.get("url") or "").strip()
                title_hint = clean_text(str(item.get("title_hint") or ""))
            else:
                raw_url = str(item).strip()
                title_hint = ""
            if not raw_url:
                continue
            detail_url = normalize_url(raw_url)
            parsed = urlparse(detail_url)
            path_and_query = f"{parsed.path}?{parsed.query}"
            if parsed.hostname != "career.cup.edu.cn":
                continue
            if not any(
                re.search(pattern, path_and_query, re.IGNORECASE)
                for pattern in detail_patterns
            ):
                continue
            if detail_url not in seen:
                seen.add(detail_url)
                candidates.append((title_hint, detail_url))
                if len(candidates) >= candidate_limit:
                    break

        # CUPB renders each listing body in the same compressed JavaScript
        # fragment used by detail pages.  The outer HTML contains the shell and
        # form only, so parse both documents and follow the rendered "next page"
        # link.  Pagination is deliberately bounded: the source is a discovery
        # channel, not a licence to crawl its entire historical archive.
        listing_page_limit = max(1, min(int(config.get("listing_page_limit", 6)), 20))
        listing_excludes = config.get(
            "listing_exclude_patterns",
            config.get("exclude_patterns", []),
        )
        listing_queue = [str(url).strip() for url in listing_urls if str(url).strip()]
        seen_listing_urls: set[str] = set()
        while listing_queue and len(candidates) < candidate_limit:
            listing_url = listing_queue.pop(0)
            # The portal's pagination router distinguishes the documented
            # ``//page/N`` path from a collapsed ``/page/N`` path.  Preserve
            # that public URL spelling for listing requests; detail URLs still
            # use the normal canonicalizer below.
            normalized_listing_url = self._cupb_listing_url(listing_url)
            if normalized_listing_url in seen_listing_urls:
                continue
            parsed_listing = urlparse(normalized_listing_url)
            if parsed_listing.hostname != "career.cup.edu.cn":
                continue
            # Keep a per-entry page budget.  The URL itself is the stable key;
            # page links generated by the portal retain a path prefix.
            entry_root = self._cupb_listing_root(normalized_listing_url)
            pages_seen_for_entry = sum(
                1
                for value in seen_listing_urls
                if self._cupb_listing_root(value) == entry_root
            )
            if pages_seen_for_entry >= listing_page_limit:
                continue
            seen_listing_urls.add(normalized_listing_url)
            try:
                response = self._get(normalized_listing_url, source)
            except (SourceSkipped, *REQUEST_ERRORS, SourceCollectionError):
                continue
            outer_soup = BeautifulSoup(response.text, "html.parser")
            decoded = self._decode_cupb_embedded_content(response.text)
            soups = [outer_soup]
            if decoded:
                soups.append(BeautifulSoup(decoded, "html.parser"))
            next_urls: list[str] = []
            for soup in soups:
                for anchor in soup.find_all("a"):
                    href = anchor.get("href")
                    title = clean_text(
                        anchor.get("title") or anchor.get_text(" ", strip=True)
                    )
                    if not href or not title:
                        continue
                    detail_url = normalize_url(urljoin(response.url, href))
                    parsed = urlparse(detail_url)
                    path_and_query = f"{parsed.path}?{parsed.query}"
                    if parsed.hostname != "career.cup.edu.cn":
                        continue
                    if any(
                        re.search(pattern, path_and_query, re.IGNORECASE)
                        for pattern in detail_patterns
                    ):
                        # Listing titles are discovery hints only.  Apply
                        # exclusions here, then require professional evidence
                        # on the detail page itself.
                        if self._is_non_vacancy_notice_title(title) or any(
                            re.search(pattern, title, re.IGNORECASE)
                            for pattern in listing_excludes
                        ):
                            continue
                        if detail_url not in seen:
                            seen.add(detail_url)
                            candidates.append((title, detail_url))
                            if len(candidates) >= candidate_limit:
                                break
                    elif self._cupb_is_next_page_anchor(anchor):
                        next_url = self._cupb_listing_url(urljoin(response.url, href))
                        if next_url not in seen_listing_urls and next_url not in next_urls:
                            next_urls.append(next_url)
                if len(candidates) >= candidate_limit:
                    break
            # Only follow the first next-page link.  This preserves the source's
            # publication order and avoids multiplying requests through page
            # number links that all point to the same archive.
            if len(candidates) < candidate_limit and next_urls:
                listing_queue.append(next_urls[0])
            self._wait(source)

        # The portal mixes banks, general employers and energy/geoscience
        # notices in one feed.  Sort discovery candidates by an explicit title
        # hint before opening detail pages so a bounded run spends its detail
        # budget on the target sector, while still retaining generic notices
        # as a lower-priority fallback.
        candidates = self._cupb_deduplicate_candidates(candidates)
        candidates.sort(key=lambda item: self._cupb_candidate_priority(item[0]))

        postings: list[RawPosting] = []
        for title_hint, detail_url in candidates:
            try:
                response = self._get(detail_url, source)
                postings_for_notice = self._extract_cupb_details(
                    response.text,
                    response.url,
                    source,
                    title_hint,
                )
                for posting in postings_for_notice:
                    postings.append(posting)
                    if len(postings) >= item_limit:
                        break
                if len(postings) >= item_limit:
                    break
            except SourceSkipped:
                continue
            except (*REQUEST_ERRORS, SourceCollectionError):
                continue
            self._wait(source)
        return postings

    def _collect_cas_job_board(self, source: dict[str, Any]) -> list[RawPosting]:
        """Collect public job cards from the Chinese Academy of Sciences job board."""
        response = self._get(source["homepage_url"], source)
        soup = BeautifulSoup(response.text, "html.parser")
        selector = source["config"].get(
            "job_link_selector",
            "a[href*='toRecruitment.do'][href*='postInfo.id']",
        )
        item_limit = self._item_limit(source)
        candidates: list[tuple[str, str, str, str | None]] = []
        seen: set[str] = set()

        for anchor in soup.select(selector):
            href = anchor.get("href")
            title = clean_text(anchor.get("title") or anchor.get_text(" ", strip=True))
            if not href or not title:
                continue
            detail_url = normalize_url(urljoin(response.url, href))
            if detail_url in seen:
                continue
            seen.add(detail_url)
            row = anchor.find_parent("tr")
            cells = row.find_all("td", recursive=False) if row else []
            employer = (
                clean_text(cells[1].get_text(" ", strip=True))
                if len(cells) >= 2
                else source["publisher"]
            )
            location = (
                clean_text(cells[2].get_text(" ", strip=True))
                if len(cells) >= 3
                else None
            )
            candidates.append((title, detail_url, employer, location))
            if len(candidates) >= item_limit:
                break

        postings: list[RawPosting] = []
        for title, detail_url, employer, location in candidates:
            try:
                detail_response = self._get(detail_url, source)
                posting = self._extract_cas_job_detail(
                    detail_response.text,
                    detail_response.url,
                    title,
                    employer,
                    location,
                )
                postings.append(posting)
            except SourceSkipped:
                continue
            except (*REQUEST_ERRORS, SourceCollectionError):
                continue
            self._wait(source)
        return postings

    def _collect_successfactors_search(
        self, source: dict[str, Any]
    ) -> list[RawPosting]:
        """Collect keyword-filtered public results from a SuccessFactors job board."""
        config = source["config"]
        search_url = config.get("search_url", source["homepage_url"])
        locale = config.get("locale", "en_US")
        queries = config.get("queries", [])
        if not queries:
            raise SourceCollectionError("successfactors_search requires configured queries")
        item_limit = self._item_limit(source)
        postings: list[RawPosting] = []
        seen: set[str] = set()

        for query in queries:
            params = urlencode({"q": query, "locale": locale})
            separator = "&" if "?" in search_url else "?"
            response = self._get(f"{search_url}{separator}{params}", source)
            soup = BeautifulSoup(response.text, "html.parser")
            for row in soup.select("tr.data-row"):
                anchor = row.select_one("a.jobTitle-link[href*='/job/']")
                if anchor is None:
                    continue
                title = clean_text(anchor.get_text(" ", strip=True))
                detail_url = normalize_url(urljoin(response.url, anchor["href"]))
                if not title or detail_url in seen:
                    continue
                if not self._title_matches_filters(title, config):
                    continue
                seen.add(detail_url)
                location_node = row.select_one(
                    ".colLocation .jobLocation, .jobLocation"
                )
                location = clean_text(
                    location_node.get_text(" ", strip=True) if location_node else ""
                ) or None
                row_text = clean_text(row.get_text(" ", strip=True))
                detail_posting: RawPosting | None = None
                if config.get("fetch_detail_pages"):
                    self._wait(source)
                    try:
                        detail_response = self._get(detail_url, source)
                        detail_posting = self._extract_successfactors_detail(
                            detail_response.text,
                            detail_response.url,
                            source,
                            title,
                            location,
                        )
                    except (SourceSkipped, SourceCollectionError):
                        # Keep the public search-row fallback when a detail page
                        # is temporarily unavailable or has stricter robots rules.
                        detail_posting = None
                    if detail_posting is not None:
                        postings.append(detail_posting)
                        if len(postings) >= item_limit:
                            return postings
                        continue
                postings.append(
                    RawPosting(
                        title=title,
                        employer=source["publisher"],
                        source_url=detail_url,
                        application_url=detail_url,
                        text=(
                            f"{title} {source['publisher']} {location or ''} "
                            f"{row_text} Search topic: {query}"
                        ),
                        match_text=f"{title} {row_text}",
                        summary=(
                            f"{source['publisher']} 官方招聘岗位"
                            f"{'，地点：' + location if location else ''}"
                        ),
                        published_date=extract_published_date(row_text),
                        deadline_date=extract_deadline(row_text),
                        location=location,
                        external_id=self._external_id_from_url(detail_url),
                    )
                )
                if len(postings) >= item_limit:
                    return postings
            self._wait(source)
        return postings

    def _extract_successfactors_detail(
        self,
        document: str,
        source_url: str,
        source: dict[str, Any],
        title_hint: str,
        location_hint: str | None,
    ) -> RawPosting:
        """Read a public SuccessFactors detail page when the source permits it."""
        soup = BeautifulSoup(document, "html.parser")
        config = source["config"]
        title_node = self._select_first(
            soup,
            config.get("detail_title_selector", "h1, .job-title, .jobTitle")
        )
        title = clean_text(title_node.get_text(" ", strip=True)) if title_node else title_hint
        content_node = self._select_first(
            soup,
            config.get("detail_content_selector", ".jobDisplayShell, .jobDisplay, #content")
        ) or soup.body
        if content_node:
            for selector in (
                "script",
                "style",
                "noscript",
                "form.jobAlertsSearchForm",
                ".cookie-banner",
            ):
                for node in content_node.select(selector):
                    node.decompose()
        body_text = clean_text(
            content_node.get_text(" ", strip=True) if content_node else document
        )
        location_node = soup.select_one(
            config.get(
                "detail_location_selector",
                "#job-location, .job-location, [id*='job-location']",
            )
        )
        location = clean_text(
            location_node.get_text(" ", strip=True) if location_node else ""
        ) or location_hint
        if not location:
            location_match = re.search(
                r"Location\s*:\s*(.+?)(?=\s+(?:Date|Job Duties|Qualifications|We are|$))",
                body_text,
                re.IGNORECASE,
            )
            location = clean_text(location_match.group(1)) if location_match else None
        summary_sections: list[str] = []
        for marker in ("Job Duties", "Qualifications", "Job Details"):
            marker_position = body_text.lower().find(marker.lower())
            if marker_position >= 0:
                summary_sections.append(body_text[marker_position : marker_position + 210])
        summary = clean_text("；".join(summary_sections))[:420] or body_text[:420] or title
        return RawPosting(
            title=title or title_hint,
            employer=source["publisher"],
            source_url=normalize_url(source_url),
            application_url=self._find_application_url(soup, source_url)
            or normalize_url(source_url),
            text=clean_text(f"{title or title_hint} {body_text}"),
            summary=summary,
            published_date=extract_published_date(body_text),
            deadline_date=extract_deadline(body_text),
            location=location,
            external_id=self._external_id_from_url(source_url),
            match_text=clean_text(f"{title or title_hint} {body_text}"),
        )

    def _collect_mokahr_search(self, source: dict[str, Any]) -> list[RawPosting]:
        """Collect public jobs from an employer's official MokaHR portal."""
        config = source["config"]
        listing_url = config.get("listing_url", source["homepage_url"])
        listing_response = self._get(listing_url, source)
        initial = self._mokahr_initial_data(listing_response.text)
        organization = initial.get("org") if isinstance(initial.get("org"), dict) else {}
        org_id = str(config.get("org_id") or organization.get("id") or "").strip()
        site_id = config.get("site_id") or initial.get("siteId") or organization.get("siteId")
        if not org_id or site_id is None:
            raise SourceCollectionError("MokaHR public page is missing organization/site id")
        try:
            site_id = int(site_id)
        except (TypeError, ValueError) as error:
            raise SourceCollectionError("MokaHR site id is invalid") from error
        iv = str(initial.get("aesIv") or "").strip()
        if not iv:
            raise SourceCollectionError("MokaHR public page is missing its AES IV")
        mode = str(config.get("mode") or initial.get("mode") or "").lower()
        mode = "campus" if mode == "camp" else mode
        if mode not in {"social", "campus"}:
            raise SourceCollectionError("MokaHR source mode must be social or campus")
        locale = str(config.get("locale", "zh-CN"))
        parsed = urlparse(listing_response.url)
        api_base = f"{parsed.scheme}://{parsed.netloc}"
        list_url = urljoin(api_base, "/api/outer/ats-apply/website/jobs/v2")
        result = self._mokahr_api_json(
            list_url,
            {
                "orgId": org_id,
                "siteId": site_id,
                "limit": self._mokahr_page_size(source),
                "offset": 0,
                "needStat": True,
                "site": mode,
                "locale": locale,
            },
            iv,
        )
        result_data = result.get("data")
        jobs = result_data.get("jobs") if isinstance(result_data, dict) else None
        if not isinstance(jobs, list):
            raise SourceCollectionError("MokaHR job list payload does not contain jobs")
        detail_url = urljoin(api_base, "/api/outer/ats-apply/website/job")
        postings: list[RawPosting] = []
        seen: set[str] = set()
        for listed_job in jobs:
            if not isinstance(listed_job, dict):
                continue
            job_id = str(listed_job.get("id") or "").strip()
            if not job_id or job_id in seen:
                continue
            seen.add(job_id)
            if not self._mokahr_candidate_allowed(
                self._mokahr_job_text(listed_job), config
            ):
                continue
            job = listed_job
            if config.get("fetch_detail_pages", True):
                self._wait(source)
                try:
                    detail_result = self._mokahr_api_json(
                        detail_url,
                        {
                            "orgId": org_id,
                            "siteId": site_id,
                            "jobId": job_id,
                            "locale": locale,
                        },
                        iv,
                    )
                    detail_data = detail_result.get("data")
                    if not isinstance(detail_data, dict):
                        raise SourceCollectionError("MokaHR job detail payload is invalid")
                    job = {**listed_job, **detail_data}
                except (SourceSkipped, SourceCollectionError):
                    if config.get("require_detail_pages"):
                        continue
            if not self._mokahr_candidate_allowed(
                self._mokahr_job_text(job),
                config,
                exclude_patterns=config.get("detail_exclude_patterns", []),
            ):
                continue
            postings.append(self._mokahr_posting(job, source, listing_url))
            if len(postings) >= self._item_limit(source):
                break
        return postings

    @staticmethod
    def _cupb_listing_root(url: str) -> str:
        """Return a stable root for bounding one CUPB listing's pagination."""
        parsed = urlparse(url)
        path = parsed.path
        if path.startswith("/campus"):
            path = "/campus"
        elif path.startswith("/news/index/tag/xwzp") or (
            path.startswith("/news/index/domain/") and "/tag/xwzp" in path
        ):
            path = "/news/index/tag/xwzp"
        elif path.startswith("/job/search"):
            path = "/job/search"
        else:
            path = re.sub(r"/page/\d+/?$", "", path)
        return f"{parsed.scheme}://{parsed.netloc}{path}?{parsed.query}"

    @staticmethod
    def _cupb_listing_url(url: str) -> str:
        """Canonicalize a listing URL without collapsing meaningful slashes."""
        parsed = urlparse(str(url).strip())
        query = f"?{parsed.query}" if parsed.query else ""
        return f"{parsed.scheme.lower()}://{parsed.netloc.lower()}{parsed.path}{query}"

    @staticmethod
    def _cupb_is_next_page_anchor(anchor: Any) -> bool:
        text = clean_text(anchor.get_text(" ", strip=True))
        if "下一页" in text or text.lower() in {"next", "next page"}:
            return True
        href = str(anchor.get("href") or "")
        return bool(re.search(r"/page/2(?:/|$)", href)) and not text.isdigit()

    @staticmethod
    def _cupb_candidate_priority(title: str) -> tuple[int, int]:
        """Prioritize target-sector notices without filtering broad employers."""
        target_terms = (
            "中国石油",
            "中国石化",
            "中国海油",
            "国家管网",
            "油田",
            "石油",
            "石化",
            "油气",
            "地质",
            "地球物理",
            "勘探",
            "勘查",
            "矿产",
            "地勘",
            "地调",
            "自然资源",
            "测绘",
            "水文",
            "能源",
            "研究院",
            "地理信息",
        )
        normalized = clean_text(title)
        return (0 if any(term in normalized for term in target_terms) else 1, len(normalized))

    @staticmethod
    def _cupb_employer_from_title(title: str) -> str | None:
        """Infer a real unit from a CUPB title when the portal account is a relay."""
        normalized = clean_text(title).strip("\u200b \t\r\n")
        if not normalized:
            return None
        normalized = re.sub(
            r"(?:20\d{2}(?:年|届)?|招聘公告|招聘简章|招聘启事|校园招聘|高校毕业生招聘|秋季招聘|春季招聘).*$",
            "",
            normalized,
            flags=re.IGNORECASE,
        ).strip(" -|丨：:，,。. ")
        if not re.search(
            r"(?:有限公司|有限责任公司|集团|油田|研究院|实验室|大学|学院|银行|医院|矿业|地质局)$",
            normalized,
        ):
            return None
        return normalized if len(normalized) >= 3 else None

    @staticmethod
    def _cupb_deduplicate_candidates(
        candidates: list[tuple[str, str]],
    ) -> list[tuple[str, str]]:
        """Collapse the portal's announcement/position-page duplicates."""
        selected: dict[str, tuple[str, str]] = {}
        for title, url in candidates:
            key = re.sub(r"\s+", "", clean_text(title)).casefold()
            if not key:
                continue
            existing = selected.get(key)
            if existing is None or OfficialSourceCollector._cupb_url_rank(url) < OfficialSourceCollector._cupb_url_rank(existing[1]):
                selected[key] = (title, url)
        return list(selected.values())

    @staticmethod
    def _cupb_url_rank(url: str) -> int:
        parsed = urlparse(url)
        if "/job/view/" in parsed.path:
            return 2
        if "/campus/view/" in parsed.path:
            return 0
        if "/news/view/" in parsed.path:
            return 1
        return 3

    def _collect_structured_opening_page(
        self, source: dict[str, Any]
    ) -> list[RawPosting]:
        """Read one public page whose opening blocks have stable field labels.

        This adapter is intentionally strict.  A page with several opening
        headings must expose the same number of non-empty content blocks; a
        positional mismatch could attach one unit's location or qualification
        text to another unit's title and is therefore a collection failure.
        """
        config = source["config"]
        opening_url = normalize_url(
            str(config.get("opening_url") or source["homepage_url"])
        )
        response = self._get(opening_url, source)
        allowed_hosts = {
            str(host).strip().lower().rstrip(".")
            for host in config.get("allowed_hosts", [])
            if str(host).strip()
        }
        final_host = (urlparse(response.url).hostname or "").lower().rstrip(".")
        if allowed_hosts and not any(
            final_host == host or final_host.endswith(f".{host}")
            for host in allowed_hosts
        ):
            raise SourceCollectionError(
                f"Structured opening page redirected outside official hosts: {final_host}"
            )

        title_selector = str(config.get("opening_title_selector") or "").strip()
        content_selector = str(config.get("opening_content_selector") or "").strip()
        if not title_selector or not content_selector:
            raise SourceCollectionError(
                "structured_opening_page requires opening title and content selectors"
            )
        try:
            soup = BeautifulSoup(response.text, "html.parser")
            title_pattern = str(config.get("opening_title_pattern") or "").strip()
            title_nodes = [
                node
                for node in soup.select(title_selector)
                if not title_pattern
                or re.search(
                    title_pattern,
                    clean_text(node.get_text(" ", strip=True)),
                    re.IGNORECASE,
                )
            ]
            content_nodes = []
            seen_nodes: set[int] = set()
            for node in soup.select(content_selector):
                marker = id(node)
                if marker in seen_nodes or not clean_text(node.get_text(" ", strip=True)):
                    continue
                seen_nodes.add(marker)
                content_nodes.append(node)
        except re.error as error:
            raise SourceCollectionError(
                f"Invalid structured opening title pattern: {error}"
            ) from error

        if not title_nodes:
            raise SourceCollectionError(
                "Structured opening page exposed no matching opening title"
            )
        if len(title_nodes) != len(content_nodes):
            raise SourceCollectionError(
                "Structured opening title/content block count mismatch: "
                f"{len(title_nodes)} titles versus {len(content_nodes)} non-empty contents"
            )

        labels = tuple(
            str(label).strip()
            for label in (
                config.get("opening_field_labels")
                or (
                    "Job Title",
                    "Quantity",
                    "Work Location",
                    "Job Type",
                    "Job Responsibilities",
                    "Job Requirements",
                    "How to Apply",
                    "Date",
                )
            )
            if str(label).strip()
        )
        title_label = str(config.get("opening_title_field_label") or "").strip()
        location_label = str(config.get("opening_location_field_label") or "").strip()
        date_label = str(config.get("opening_date_field_label") or "").strip()
        required_fields = {
            str(field).strip()
            for field in config.get("required_opening_fields", [])
            if str(field).strip()
        }
        postings: list[RawPosting] = []
        for index, (title_node, content_node) in enumerate(
            zip(title_nodes, content_nodes, strict=True),
            start=1,
        ):
            heading = clean_text(title_node.get_text(" ", strip=True))
            lines = self._structured_opening_lines(content_node)
            body_text = clean_text(content_node.get_text(" ", strip=True))
            title = self._structured_opening_field(lines, title_label, labels)
            if not title:
                title = re.sub(
                    str(config.get("opening_title_prefix_pattern") or r"^Job Opening\s*:\s*"),
                    "",
                    heading,
                    flags=re.IGNORECASE,
                ).strip()
            location = self._structured_opening_field(lines, location_label, labels)
            published_date = parse_date_value(
                self._structured_opening_field(lines, date_label, labels)
            ) or extract_published_date(body_text)
            combined = clean_text(f"{heading} {title} {body_text}")
            if not self._accept_candidate(combined, source):
                continue
            missing = {
                field
                for field, value in (("title", title), ("location", location))
                if field in required_fields and not value
            }
            if missing:
                raise SourceCollectionError(
                    "Structured opening is missing required fields: "
                    + ", ".join(sorted(missing))
                )
            source_url = normalize_url(response.url)
            postings.append(
                RawPosting(
                    title=title or heading,
                    employer=clean_text(
                        str(config.get("employer_hint") or source["publisher"])
                    ),
                    source_url=source_url,
                    application_url=self._find_application_url(
                        soup, source_url
                    ),
                    text=combined,
                    summary=clean_text(
                        "; ".join(
                            f"{label}: {value}"
                            for label, value in (
                                (title_label, title),
                                (location_label, location),
                            )
                            if label and value
                        )
                    )[:420]
                    or body_text[:420],
                    published_date=published_date,
                    deadline_date=extract_deadline(body_text),
                    location=location,
                    external_id=(
                        f"{self._external_id_from_url(source_url)}#opening-"
                        f"{index}-{stable_hash(title or heading)[:16]}"
                    ),
                    match_text=clean_text(f"{title or heading} {body_text}"),
                    official_evidence_url=source_url,
                )
            )
        return postings

    @staticmethod
    def _structured_opening_lines(node: Any) -> list[str]:
        return [
            clean_text(line)
            for line in node.get_text("\n", strip=True).splitlines()
            if clean_text(line)
        ]

    @staticmethod
    def _structured_opening_field(
        lines: list[str], label: str, labels: tuple[str, ...]
    ) -> str | None:
        if not label:
            return None
        normalized_label = label.rstrip("：:").strip().casefold()
        normalized_labels = {
            item.rstrip("：:").strip().casefold() for item in labels
        }
        for index, line in enumerate(lines):
            inline = re.match(
                rf"^{re.escape(label)}\s*[:：]\s*(.+)$", line, re.IGNORECASE
            )
            if inline:
                return clean_text(inline.group(1)) or None
            if line.rstrip("：:").strip().casefold() != normalized_label:
                continue
            values: list[str] = []
            for following in lines[index + 1 :]:
                if following.rstrip("：:").strip().casefold() in normalized_labels:
                    break
                values.append(following)
            return clean_text(" ".join(values)) or None
        return None

    def _collect_zhaopin_campus(self, source: dict[str, Any]) -> list[RawPosting]:
        """Collect a public Zhaopin campus campaign without login or writes.

        Zhaopin's campaign pages are a static shell around a public, read-only
        JSON request.  The adapter follows the request shape used by the page
        itself, but treats a business error, malformed payload, or missing job
        evidence as a collection failure.  Only an explicit ``code=200`` with
        a valid ``jobList`` can produce the legitimate empty result ``[]``.
        """
        config = source["config"]
        listing_url = str(config.get("listing_url") or source["homepage_url"])
        landing = self._get(listing_url, source)
        metadata = self._zhaopin_campaign_metadata(
            landing.text,
            config,
            landing.url,
        )
        company_id = str(
            config.get("company_id")
            or metadata.get("companyId")
            or metadata.get("xiaozhaoId")
            or ""
        ).strip()
        if not company_id:
            raise SourceCollectionError(
                "Zhaopin public campaign is missing companyId/xiaozhaoId"
            )
        scene = str(config.get("scene") or metadata.get("scene") or "cam").lower()
        if scene not in {"cam", "social"}:
            raise SourceCollectionError(f"Unsupported Zhaopin campaign scene: {scene}")
        api_host = str(config.get("api_host") or "https://fe.zhaopin.com").rstrip("/")
        api_url = urljoin(api_host + "/", str(
            config.get("api_path") or "/grace/api/dsc/search-job-list"
        ).lstrip("/"))
        parsed_api_host = urlparse(api_url).hostname
        allowed_api_hosts = {
            str(value).lower()
            for value in config.get("api_allowed_hosts", [])
            if str(value).strip()
        }
        if not parsed_api_host or parsed_api_host.lower() not in allowed_api_hosts:
            raise SourceCollectionError(
                f"Zhaopin API host is not allowlisted: {parsed_api_host or api_url}"
            )
        job_source = 2 if scene == "cam" else 1
        page_size = min(self._item_limit(source), 100)
        request_payload = {
            "orgNumbers": [company_id],
            "jobSource": job_source,
            "pageIndex": 1,
            "pageSize": page_size,
            "orgDepartmentIds": [],
            "workRegionIds": "",
            "jobTypes": "",
            "priorityMajors": "",
            "customTags": "",
        }
        self._wait(source)
        response = self._post_json(
            api_url,
            request_payload,
            headers={
                "Accept": "application/json",
                "Content-Type": "application/json",
                "Origin": f"{urlparse(landing.url).scheme}://{urlparse(landing.url).netloc}",
                "Referer": landing.url,
            },
        )
        try:
            payload = response.json()
        except ValueError as error:
            raise SourceCollectionError(
                "Zhaopin public job API did not return JSON"
            ) from error
        if not isinstance(payload, dict):
            raise SourceCollectionError("Zhaopin public job API returned an unexpected payload")
        code = payload.get("code")
        if str(code) != "200":
            message = clean_text(str(payload.get("message") or payload.get("msg") or ""))
            detail = f"; message={message}" if message else ""
            raise SourceCollectionError(
                f"Zhaopin public job API business error code={code}{detail}"
            )
        data = payload.get("data")
        if not isinstance(data, dict):
            raise SourceCollectionError("Zhaopin public job API data is not an object")
        jobs = data.get("jobList")
        page_info = data.get("pageInfo")
        if not isinstance(jobs, list) or not isinstance(page_info, dict):
            raise SourceCollectionError(
                "Zhaopin public job API is missing jobList/pageInfo"
            )
        # An empty list is only a valid no-match result when the public page
        # explicitly reports a zero total.  Missing or inconsistent counts are
        # treated as an adapter regression instead of silently publishing none.
        try:
            total = int(page_info.get("totalNum", -1))
        except (TypeError, ValueError):
            total = -1
        if total < 0 or (total == 0 and jobs) or (total > 0 and not jobs):
            raise SourceCollectionError(
                "Zhaopin public job API returned inconsistent pageInfo/jobList"
            )
        postings: list[RawPosting] = []
        for item in jobs:
            if not isinstance(item, dict):
                continue
            job = item.get("job") if isinstance(item.get("job"), dict) else item
            posting = self._zhaopin_posting(job, source, listing_url)
            if posting is not None:
                postings.append(posting)
            if len(postings) >= page_size:
                break
        if jobs and not postings:
            raise SourceCollectionError(
                "Zhaopin public job API returned rows but no row had publishable evidence"
            )
        return postings

    def _zhaopin_campaign_metadata(
        self,
        document: str,
        config: dict[str, Any],
        base_url: str,
    ) -> dict[str, str]:
        """Read public campaign identifiers from configured values or assets."""
        metadata: dict[str, str] = {}
        for key in ("companyId", "companyNumber", "scene", "xiaozhaoId"):
            configured_value = config.get(
                {
                    "companyId": "company_id",
                    "companyNumber": "company_number",
                    "scene": "scene",
                    "xiaozhaoId": "campaign_id",
                }[key]
            )
            configured = (
                str(configured_value).strip() if configured_value is not None else ""
            )
            if configured:
                metadata[key] = configured

        candidates = [document]
        if not metadata.get("companyId") and not metadata.get("xiaozhaoId"):
            soup = BeautifulSoup(document, "html.parser")
            asset_hosts = {
                str(value).lower()
                for value in config.get(
                    "metadata_allowed_hosts",
                    [urlparse(base_url).netloc, "webapp.zhaopin.com", "common-bucket.zhaopin.cn"],
                )
                if str(value).strip()
            }
            for script in soup.find_all("script", src=True)[: int(config.get("metadata_asset_limit", 8))]:
                asset_url = normalize_url(urljoin(base_url, str(script.get("src"))))
                if urlparse(asset_url).hostname not in asset_hosts:
                    continue
                try:
                    candidates.append(self._get(asset_url, {"config": {"request_interval_seconds": 0}}).text)
                except (SourceSkipped, SourceCollectionError):
                    continue
        for candidate in candidates:
            for key in ("companyId", "companyNumber", "scene", "xiaozhaoId"):
                if metadata.get(key):
                    continue
                match = re.search(
                    rf"[\"']?{key}[\"']?\s*:\s*[\"']([^\"']*)[\"']",
                    candidate,
                )
                if match and match.group(1).strip():
                    metadata[key] = clean_text(match.group(1))
        return metadata

    def _zhaopin_posting(
        self,
        job: dict[str, Any],
        source: dict[str, Any],
        listing_url: str,
    ) -> RawPosting | None:
        title = clean_text(str(job.get("title") or ""))
        detail_html = str(job.get("detail") or job.get("jobDetail") or "")
        detail = clean_text(BeautifulSoup(detail_html, "html.parser").get_text(" ", strip=True))
        category_value = job.get("jobCategories")
        if isinstance(category_value, list):
            categories = "、".join(clean_text(str(value)) for value in category_value)
        else:
            categories = clean_text(str(category_value or ""))
        city = clean_text(str(job.get("cityName") or job.get("workCity") or ""))
        evidence_url = normalize_url(str(job.get("url") or "")) if job.get("url") else ""
        if evidence_url:
            allowed_hosts = {
                str(value).lower()
                for value in source["config"].get("allowed_hosts", [])
            }
            if urlparse(evidence_url).hostname not in allowed_hosts:
                evidence_url = ""
        if not title or not detail or not evidence_url:
            return None
        match_text = clean_text(f"{title} {categories} {city} {detail}")
        include_patterns = source["config"].get("include_patterns", [])
        if include_patterns and not any(
            re.search(pattern, match_text, re.IGNORECASE)
            for pattern in include_patterns
        ):
            return None
        if not self._accept_candidate(match_text, source):
            return None
        job_number = clean_text(str(job.get("jobNumber") or job.get("id") or ""))
        fields = [
            f"职位类别：{categories}" if categories else "",
            f"工作地点：{city}" if city else "",
        ]
        summary = clean_text("；".join(value for value in fields if value))
        if detail:
            summary = clean_text(f"{summary}；{detail}" if summary else detail)[:420]
        return RawPosting(
            title=title,
            employer=source["publisher"],
            source_url=evidence_url,
            application_url=evidence_url,
            text=match_text,
            summary=summary or title,
            published_date=extract_published_date(detail),
            deadline_date=extract_deadline(detail),
            location=city or extract_location_hint(detail),
            external_id=job_number or self._external_id_from_url(evidence_url),
            match_text=match_text,
            official_evidence_url=evidence_url,
        )

    def _mokahr_api_json(
        self,
        url: str,
        payload: dict[str, Any],
        initialization_vector: str,
        source: dict[str, Any] | None = None,
    ) -> dict[str, Any]:
        response = self._post_json(url, payload)
        try:
            encrypted = response.json()
        except ValueError as error:
            raise SourceCollectionError("MokaHR API did not return JSON") from error
        if not isinstance(encrypted, dict):
            raise SourceCollectionError("MokaHR API returned an unexpected payload")
        return self._decrypt_mokahr_payload(encrypted, initialization_vector)

    @staticmethod
    def _mokahr_initial_data(document: str) -> dict[str, Any]:
        node = BeautifulSoup(document, "html.parser").select_one("#init-data")
        value = node.get("value") if node else None
        if not isinstance(value, str) or not value.strip():
            raise SourceCollectionError("MokaHR public page is missing #init-data")
        try:
            parsed = json.loads(unescape(value))
        except json.JSONDecodeError as error:
            raise SourceCollectionError("MokaHR init-data is invalid JSON") from error
        if not isinstance(parsed, dict):
            raise SourceCollectionError("MokaHR init-data is not an object")
        return parsed

    @staticmethod
    def _decrypt_mokahr_payload(
        payload: dict[str, Any],
        initialization_vector: str,
    ) -> dict[str, Any]:
        if isinstance(payload.get("data"), dict):
            return payload
        if Cipher is None or algorithms is None or modes is None or PKCS7 is None:
            raise SourceCollectionError("MokaHR support requires cryptography")
        encrypted_data = payload.get("data")
        response_key = payload.get("necromancer")
        if not isinstance(encrypted_data, str) or not isinstance(response_key, str):
            raise SourceCollectionError("MokaHR encrypted response is incomplete")
        try:
            cipher = Cipher(
                algorithms.AES(response_key.encode("utf-8")),
                modes.CBC(initialization_vector.encode("utf-8")),
            )
            decryptor = cipher.decryptor()
            padded = decryptor.update(base64.b64decode(encrypted_data, validate=True))
            padded += decryptor.finalize()
            unpadder = PKCS7(algorithms.AES.block_size).unpadder()
            decoded = json.loads((unpadder.update(padded) + unpadder.finalize()).decode("utf-8"))
        except (ValueError, TypeError, UnicodeDecodeError, binascii.Error, json.JSONDecodeError) as error:
            raise SourceCollectionError("Unable to decrypt MokaHR public response") from error
        if not isinstance(decoded, dict):
            raise SourceCollectionError("MokaHR decrypted response is not an object")
        if decoded.get("success") is False:
            raise SourceCollectionError(clean_text(str(decoded.get("msg") or "MokaHR API rejected the request")))
        return decoded

    @staticmethod
    def _mokahr_job_text(job: dict[str, Any]) -> str:
        values: list[str] = []
        for field in ("title", "education", "commitment", "jobDescription"):
            value = job.get(field)
            if isinstance(value, str):
                values.append(value)
        for field in ("department", "zhineng"):
            value = job.get(field)
            if isinstance(value, dict) and isinstance(value.get("name"), str):
                values.append(value["name"])
        return clean_text(
            " ".join(
                BeautifulSoup(value, "html.parser").get_text(" ", strip=True)
                if value.lstrip().startswith("<")
                else value
                for value in values
            )
        )

    @staticmethod
    def _mokahr_candidate_allowed(
        value: str,
        config: dict[str, Any],
        *,
        exclude_patterns: list[str] | None = None,
    ) -> bool:
        normalized = clean_text(value)
        patterns = config.get("exclude_patterns", []) if exclude_patterns is None else exclude_patterns
        if any(
            re.search(pattern, normalized, re.IGNORECASE)
            for pattern in patterns
        ):
            return False
        include_patterns = config.get("include_patterns", [])
        if include_patterns:
            return any(
                re.search(pattern, normalized, re.IGNORECASE)
                for pattern in include_patterns
            )
        return bool(extract_major_tags(normalized))

    @staticmethod
    def _mokahr_page_size(source: dict[str, Any]) -> int:
        value = source["config"].get(
            "page_size", source["config"].get("max_items", 40)
        )
        try:
            return max(1, min(int(value), 100))
        except (TypeError, ValueError):
            return 40

    @staticmethod
    def _mokahr_date(value: Any) -> str | None:
        match = re.search(r"(20\d{2}-\d{2}-\d{2})", str(value or ""))
        return match.group(1) if match else None

    @staticmethod
    def _mokahr_location(job: dict[str, Any]) -> str | None:
        def labels(value: Any) -> list[str]:
            if isinstance(value, str):
                return [clean_text(value)]
            if isinstance(value, list):
                return [item for entry in value for item in labels(entry)]
            if isinstance(value, dict):
                named = clean_text(str(value.get("name") or value.get("label") or ""))
                if named:
                    return [named]
                return [
                    item
                    for key in ("country", "province", "city", "district")
                    for item in labels(value.get(key))
                ]
            return []

        values: list[str] = []
        for key in (
            "locations",
            "location",
            "workLocation",
            "workplace",
            "workPlace",
            "city",
            "address",
            "place",
        ):
            values.extend(labels(job.get(key)))
        # Different MokaHR tenants use different field names.  Preserve the
        # first explicit location-like value and avoid treating a department
        # or business function as a place.
        unique = list(dict.fromkeys(value for value in values if value))
        return "、".join(unique) or None

    @staticmethod
    def _mokahr_detail_url(listing_url: str, job_id: str) -> str:
        return f"{listing_url.split('#', 1)[0]}#/job/{job_id}"

    @staticmethod
    def _mokahr_named_value(value: Any) -> str | None:
        if isinstance(value, dict):
            value = value.get("name")
        return clean_text(str(value or "")) or None

    def _mokahr_posting(
        self,
        job: dict[str, Any],
        source: dict[str, Any],
        listing_url: str,
    ) -> RawPosting:
        title = clean_text(str(job.get("title") or "未命名岗位"))
        description = clean_text(
            BeautifulSoup(
                str(job.get("jobDescription") or ""), "html.parser"
            ).get_text(" ", strip=True)
        )
        education = clean_text(str(job.get("education") or ""))
        commitment = clean_text(str(job.get("commitment") or ""))
        function = self._mokahr_named_value(job.get("zhineng"))
        department = self._mokahr_named_value(job.get("department"))
        location = self._mokahr_location(job)
        if not location:
            location = extract_location_hint(
                " ".join(
                    value
                    for value in (
                        description,
                        clean_text(str(job.get("content") or "")),
                        clean_text(str(job.get("detail") or "")),
                    )
                    if value
                )
            )
        fields = [
            f"学历：{education}" if education else "",
            f"职位性质：{commitment}" if commitment else "",
            f"职位类别：{function}" if function else "",
            f"所属部门：{department}" if department else "",
            f"工作地点：{location}" if location else "",
        ]
        field_text = clean_text("；".join(value for value in fields if value))
        match_text = clean_text(f"{title} {field_text} {description}")
        detail_url = self._mokahr_detail_url(listing_url, str(job.get("id") or ""))
        return RawPosting(
            title=title,
            employer=source["publisher"],
            source_url=detail_url,
            application_url=detail_url,
            text=match_text,
            summary=clean_text(f"{field_text}；{description}")[:420] or field_text or title,
            published_date=self._mokahr_date(job.get("publishedAt")),
            deadline_date=self._mokahr_date(job.get("closedAt")) or extract_deadline(description),
            location=location,
            external_id=str(job.get("id") or "") or None,
            match_text=match_text,
        )

    def _collect_rss(self, source: dict[str, Any]) -> list[RawPosting]:
        response = self._get(source["homepage_url"], source)
        try:
            root = ET.fromstring(response.content)
        except ET.ParseError as error:
            raise SourceCollectionError(f"Invalid RSS XML: {error}") from error

        config = source["config"]
        entries = list(root.findall(".//item"))
        if not entries:
            entries = list(root.findall(".//{http://www.w3.org/2005/Atom}entry"))
        postings: list[RawPosting] = []
        for entry in entries[: self.settings.max_source_items]:
            title = clean_text(
                self._xml_text(entry, "title")
                or self._xml_text(entry, "{http://www.w3.org/2005/Atom}title")
            )
            link = self._rss_link(entry)
            description = clean_text(
                self._xml_text(entry, "description")
                or self._xml_text(entry, "content")
                or self._xml_text(entry, "{http://www.w3.org/2005/Atom}summary")
                or self._xml_text(entry, "{http://www.w3.org/2005/Atom}content")
            )
            combined = f"{title} {description}"
            if not title or not link:
                continue
            if not config.get("accept_all_entries") and not self._accept_candidate(
                combined, source
            ):
                continue
            postings.append(
                RawPosting(
                    title=title,
                    employer=config.get("employer_hint", source["publisher"]),
                    source_url=normalize_url(urljoin(response.url, link)),
                    application_url=None,
                    text=combined,
                    summary=description[:420],
                    published_date=extract_published_date(
                        self._xml_text(entry, "pubDate")
                        or self._xml_text(entry, "published")
                        or self._xml_text(entry, "updated")
                    ),
                    deadline_date=extract_deadline(combined),
                    location=config.get("location_hint"),
                    external_id=self._xml_text(entry, "guid") or None,
                )
            )
        return postings

    def _collect_json(self, source: dict[str, Any]) -> list[RawPosting]:
        response = self._get(source["homepage_url"], source)
        try:
            body = response.json()
        except ValueError as error:
            raise SourceCollectionError(f"Invalid JSON feed: {error}") from error
        config = source["config"]
        items = self._nested_value(body, config.get("items_path", "items"))
        if not isinstance(items, list):
            raise SourceCollectionError("Configured JSON items_path did not resolve to a list")
        postings: list[RawPosting] = []
        for item in items[: self.settings.max_source_items]:
            if not isinstance(item, dict):
                continue
            title = clean_text(str(self._nested_value(item, config.get("title_key", "title")) or ""))
            link_value = self._nested_value(item, config.get("url_key", "url"))
            if not title or not link_value:
                continue
            link = normalize_url(urljoin(response.url, str(link_value)))
            body_text = clean_text(
                str(self._nested_value(item, config.get("description_key", "description")) or "")
            )
            combined = f"{title} {body_text}"
            if not config.get("accept_all_entries") and not self._accept_candidate(
                combined, source
            ):
                continue
            postings.append(
                RawPosting(
                    title=title,
                    employer=clean_text(
                        str(
                            self._nested_value(
                                item,
                                config.get("employer_key", "employer"),
                            )
                            or config.get("employer_hint", source["publisher"])
                        )
                    ),
                    source_url=link,
                    application_url=normalize_url(
                        urljoin(
                            response.url,
                            str(
                                self._nested_value(
                                    item, config.get("application_url_key", "application_url")
                                )
                                or link
                            ),
                        )
                    ),
                    text=combined,
                    summary=body_text[:420],
                    published_date=extract_published_date(
                        str(
                            self._nested_value(
                                item, config.get("published_date_key", "published_at")
                            )
                            or ""
                        )
                    ),
                    deadline_date=extract_deadline(combined),
                    location=clean_text(
                        str(
                            self._nested_value(item, config.get("location_key", "location"))
                            or config.get("location_hint", "")
                        )
                    )
                    or None,
                    external_id=str(
                        self._nested_value(item, config.get("id_key", "id")) or ""
                    )
                    or None,
                )
            )
        return postings

    def _notice_links(
        self,
        soup: BeautifulSoup,
        base_url: str,
        source: dict[str, Any],
    ) -> list[tuple[str, str]]:
        config = source["config"]
        selector = config.get("listing_selector")
        anchors = list(soup.select(selector) if selector else soup.find_all("a"))
        # A number of government CMS portals render their list as HTML nested
        # inside ``<script type="text/xml">`` or CDATA blocks.  A normal CSS
        # query cannot see those anchors, which silently turns a real source
        # into a false "no matching jobs" result.  Parse only these explicitly
        # marked markup blocks and run the same host/pattern checks below.
        for script in soup.select("script[type='text/xml'], script[type='application/xml']"):
            raw = script.string or script.get_text()
            if not raw:
                continue
            fragments = re.findall(r"<!\[CDATA\[(.*?)\]\]>", raw, re.DOTALL)
            if not fragments:
                fragments = [raw]
            for fragment in fragments:
                embedded = BeautifulSoup(fragment, "html.parser")
                anchors.extend(
                    embedded.select(selector) if selector else embedded.find_all("a")
                )
        allowed_hosts = set(config.get("allowed_hosts", []))
        base_parts = urlparse(base_url)
        allowed_hosts.add(base_parts.netloc.lower())
        if base_parts.hostname:
            allowed_hosts.add(base_parts.hostname.lower())
        include_patterns = config.get("include_patterns", [])
        detail_path_patterns = config.get("detail_path_patterns", [])
        exclude_patterns = config.get("exclude_patterns", [])
        links: list[tuple[str, str]] = []
        seen: set[str] = set()
        for anchor in anchors:
            href = anchor.get("href")
            title = clean_text(anchor.get("title") or anchor.get_text(" ", strip=True))
            if not href or not title:
                continue
            if self._is_non_vacancy_notice_title(title):
                continue
            if not self._title_matches_filters(title, config):
                continue
            detail_url = normalize_url(urljoin(base_url, href))
            parsed = urlparse(detail_url)
            if parsed.scheme not in {"http", "https"}:
                continue
            if (
                parsed.netloc.lower() not in allowed_hosts
                and (parsed.hostname or "").lower() not in allowed_hosts
            ):
                continue
            if re.search(r"\.(?:pdf|docx?|xlsx?|zip|rar)$", parsed.path, re.IGNORECASE):
                continue
            candidate = f"{title} {detail_url}"
            if any(
                re.search(pattern, candidate, re.IGNORECASE)
                for pattern in exclude_patterns
            ):
                continue
            if not self._accept_candidate(candidate, source):
                continue
            if include_patterns and not any(
                re.search(pattern, candidate, re.IGNORECASE) for pattern in include_patterns
            ):
                continue
            path_and_query = f"{parsed.path}?{parsed.query}"
            if detail_path_patterns and not any(
                re.search(pattern, path_and_query, re.IGNORECASE)
                for pattern in detail_path_patterns
            ):
                continue
            if detail_url in seen:
                continue
            seen.add(detail_url)
            links.append((title, detail_url))
            if len(links) >= self._item_limit(source):
                break
        return links

    def _extract_html_detail(
        self,
        document: str,
        source_url: str,
        source: dict[str, Any],
        title_hint: str,
    ) -> RawPosting | None:
        soup = BeautifulSoup(document, "html.parser")
        config = source["config"]
        title_node = self._select_first(soup, config.get("title_selector", "h1"))
        title = clean_text(title_node.get_text(" ", strip=True) if title_node else "")
        if not title:
            meta_title = soup.select_one("meta[name='ArticleTitle'], meta[property='og:title']")
            title = clean_text(str(meta_title.get("content") or "")) if meta_title else ""
        if not title:
            title = clean_text(soup.title.get_text(" ", strip=True) if soup.title else title_hint)
        if self._is_non_vacancy_notice_title(title):
            return None
        if not self._title_matches_filters(title, config):
            return None
        content_selector = config.get("content_selector")
        content_node = self._select_first(soup, content_selector) if content_selector else None
        if not content_node:
            content_node = soup.find("article") or soup.find("main") or soup.body
        if content_node:
            for selector in config.get("remove_selectors", []):
                for node in content_node.select(selector):
                    node.decompose()
        body_text = clean_text(
            content_node.get_text(" ", strip=True) if content_node else document
        )
        combined = f"{title} {body_text}"
        # Exclusion phrases such as “拟录用” and “采购” are reliable for a
        # listing title, but they routinely occur in the explanatory body of a
        # genuine vacancy (for example, the sentence describing the final
        # hiring step).  Apply them to the title only and keep full-body text
        # for professional relevance evidence.
        detail_exclude_patterns = config.get(
            "detail_exclude_patterns",
            config.get("exclude_patterns", []),
        )
        if any(
            re.search(str(pattern), title, re.IGNORECASE)
            for pattern in detail_exclude_patterns
        ):
            return None
        if not self._accept_candidate(combined, source, exclude_patterns=[]):
            return None
        application_url = self._find_application_url(soup, source_url)
        published_date = self._published_date_from_meta(soup) or extract_published_date(
            body_text
        )
        employer_selector = config.get("employer_selector")
        employer_node = (
            self._select_first(soup, employer_selector) if employer_selector else None
        )
        employer = clean_text(
            employer_node.get_text(" ", strip=True) if employer_node else ""
        ) or config.get("employer_hint", source["publisher"])
        if config.get("infer_employer_from_title"):
            employer = self._employer_from_title(title, employer)

        location = config.get("location_hint")
        if not location:
            location = self._label_value(
                body_text,
                "工作地点",
                ("岗位职责", "任职要求", "学历要求", "招聘人数", "薪酬", "报名", "截止", "联系"),
            )
        return RawPosting(
            title=title,
            employer=employer,
            source_url=normalize_url(source_url),
            application_url=application_url,
            text=combined,
            summary=body_text[:420],
            published_date=published_date,
            deadline_date=extract_deadline(body_text),
            location=location,
        )

    def _extract_cupb_details(
        self,
        document: str,
        source_url: str,
        source: dict[str, Any],
        title_hint: str,
    ) -> list[RawPosting]:
        soup = BeautifulSoup(document, "html.parser")
        title_node = soup.select_one(
            ".details-title h5, .title-message h5, h1, h2"
        )
        title = clean_text(
            title_node.get_text(" ", strip=True)
            if title_node
            else (soup.title.get_text(" ", strip=True) if soup.title else title_hint)
        )
        embedded_html = self._decode_cupb_embedded_content(document)
        content_soup = BeautifulSoup(embedded_html or document, "html.parser")
        content_candidates = content_soup.select(
            ".aContent, .zp-details, .common-view, main, article"
        )
        # Portal templates often render a short metadata block before the real
        # article body.  ``select_one`` would choose that first block and lose
        # the professional qualification and location evidence.  Choose the
        # substantive candidate while retaining the same allowed selectors.
        content_node = max(
            content_candidates,
            key=lambda node: len(node.get_text(" ", strip=True)),
            default=content_soup,
        )
        if content_node:
            for selector in (
                ".common-view-tips",
                ".operation",
                ".share",
                ".footer",
                "script",
                "style",
            ):
                for node in content_node.select(selector):
                    node.decompose()
        body_text = clean_text(content_node.get_text(" ", strip=True))
        metadata_node = soup.select_one(".zp-details, .common-view")
        metadata_text = clean_text(
            metadata_node.get_text(" ", strip=True) if metadata_node else ""
        )
        # Header text contains site-wide anti-fraud and employment-guidance copy.
        # Keep it for dates, but do not let those unrelated words reject a genuine
        # job notice through source-level exclusion patterns.
        combined = clean_text(f"{title} {body_text}")
        listing_excludes = source["config"].get(
            "listing_exclude_patterns",
            source["config"].get("exclude_patterns", []),
        )
        if self._is_non_vacancy_notice_title(title) or any(
            re.search(pattern, title, re.IGNORECASE) for pattern in listing_excludes
        ):
            return []
        # Detail-level exclusions are intentionally narrower than discovery
        # exclusions.  Words such as “通知” and “活动” occur in the portal's
        # shared header and must not erase a valid official announcement.
        if not self._accept_candidate(
            combined,
            source,
            exclude_patterns=source["config"].get("detail_exclude_patterns", []),
        ):
            return []
        if (
            source["config"].get("require_detail_content")
            and len(body_text) < int(source["config"].get("minimum_detail_characters", 80))
        ):
            # A listing title alone cannot establish the target major, degree,
            # deadline or official application route. Keep it out until its
            # public detail page exposes substantive recruitment content.
            return []
        employer_node = soup.select_one(".title-message .name")
        employer_candidate = clean_text(
            employer_node.get_text(" ", strip=True) if employer_node else ""
        )
        # The CUPB portal may display the posting account (for example a
        # third-party HR service) in this field rather than the hiring unit.
        # Prefer the unit named in the official title when the account clearly
        # looks like an intermediary; never expose the intermediary as the
        # employer merely because it owns the portal account.
        intermediary_markers = (
            "人力资源",
            "招聘服务",
            "就业服务",
            "人才服务",
            "前锦",
            "智联",
            "同道",
            "猎聘",
            "网络信息技术",
            "月生信息技术",
        )
        if not employer_candidate or any(
            marker in employer_candidate for marker in intermediary_markers
        ):
            employer_candidate = (
                self._cupb_employer_from_title(title)
                or self._employer_from_title(title, source["publisher"])
            )
        employer = employer_candidate or source["publisher"]
        fields = self._cupb_table_fields(content_soup)
        # The portal keeps the job-list table in the outer document while the
        # long announcement body is embedded in a compressed fragment.  Read
        # both public fragments so that a valid major/degree/location row is
        # not lost when the article body is decoded.
        for key, value in self._cupb_table_fields(soup).items():
            fields.setdefault(key, value)
        matching_fields = " ".join(
            value
            for key, value in fields.items()
            if key in {"岗位", "专业范围", "面向对象", "学历要求"}
        )
        records = self._cupb_table_records(content_soup)
        outer_records = self._cupb_table_records(soup)
        if len(outer_records) > len(records):
            records = outer_records
        # For a multi-row table, only the current row can establish a
        # qualification.  A single-row notice may still put the precise major
        # in the prose body, so retain its bounded evidence as a supplement.
        major_evidence = (
            self._cupb_major_evidence(body_text)
            if len(records) <= 1
            else ""
        )
        # CUPB notices often start with a long employer introduction. When their
        # structured job table is present, use it as matching evidence so a unit's
        # industry description cannot masquerade as a candidate's qualification.
        # If a notice has no table, retain only bounded sentences that explicitly
        # mention a recognized geoscience major/degree instead of the entire
        # employer introduction.
        # Keep both sources of evidence.  The compact job table may describe
        # only a broad discipline while the official announcement body lists
        # the precise geoscience majors; dropping either side creates false
        # negatives for otherwise eligible students.
        match_evidence = clean_text(" ".join(filter(None, (matching_fields, major_evidence))))
        match_text = clean_text(f"{title} {match_evidence}") or combined
        summary_fields = [
            f"{key}：{value}"
            for key, value in fields.items()
            if key in {"岗位", "专业范围", "面向对象", "工作地点"}
        ]
        base_posting = RawPosting(
            title=title,
            employer=employer,
            source_url=normalize_url(source_url),
            application_url=(
                self._find_application_url(content_soup, source_url)
                or self._find_application_url(soup, source_url)
            ),
            text=combined,
            summary=clean_text("；".join(summary_fields))[:420] or body_text[:420],
            published_date=(
                extract_published_date(metadata_text)
                or extract_published_date(body_text)
            ),
            deadline_date=extract_deadline(clean_text(f"{metadata_text} {body_text}")),
            location=(
                self._label_value(
                    body_text,
                    "工作地点",
                    (
                        "岗位职责",
                        "任职要求",
                        "学历要求",
                        "招聘人数",
                        "薪酬",
                        "报名",
                        "报名方式",
                        "栏目分类",
                        "需求学科",
                        "截止",
                        "联系",
                        "培养机制",
                        "员工福利",
                        "福利待遇",
                    ),
                )
                # The outer job-board table may contain the publisher's
                # default city rather than the hiring unit's location. Prefer
                # an explicit address/location statement in the notice body.
                or self._cupb_location_evidence(body_text)
                or fields.get("工作地点")
            ),
            external_id=self._external_id_from_url(source_url),
            match_text=match_text,
            qualification_text=match_evidence,
            field_evidence={
                key: value
                for key, value in fields.items()
                if key in {"岗位", "专业范围", "面向对象", "学历要求", "工作地点"}
            },
        )
        if len(records) <= 1:
            return [base_posting]

        # A structured table is authoritative at row level.  Do not reuse the
        # whole announcement body for every row: that would attach another
        # position's major and degree requirements to this one.
        row_postings: list[RawPosting] = []
        for row_number, record in enumerate(records, start=1):
            row_title = record.get("岗位") or title
            row_major = record.get("专业范围", "")
            row_degree = record.get("学历要求", "") or record.get("面向对象", "")
            row_location = record.get("工作地点", "") or fields.get("工作地点", "")
            row_evidence = clean_text(
                "；".join(
                    part
                    for part in (
                        f"岗位：{row_title}" if row_title else "",
                        f"专业范围：{row_major}" if row_major else "",
                        f"学历要求：{row_degree}" if row_degree else "",
                        f"工作地点：{row_location}" if row_location else "",
                    )
                    if part
                )
            )
            row_postings.append(
                RawPosting(
                    title=(
                        row_title
                        if row_title != title
                        else f"{title}（第{row_number}项）"
                    ),
                    employer=base_posting.employer,
                    source_url=base_posting.source_url,
                    application_url=base_posting.application_url,
                    text=row_evidence,
                    summary=row_evidence[:420] or base_posting.summary,
                    published_date=base_posting.published_date,
                    deadline_date=base_posting.deadline_date,
                    location=row_location or base_posting.location,
                    external_id=f"{base_posting.external_id or self._external_id_from_url(source_url)}#row-{row_number}",
                    match_text=clean_text(f"{row_title} {row_major} {row_degree}"),
                    qualification_text=row_evidence,
                    official_evidence_url=base_posting.official_evidence_url,
                    field_evidence={
                        "table_row": str(row_number),
                        **{
                            key: value
                            for key, value in record.items()
                            if key in {"岗位", "专业范围", "面向对象", "学历要求", "工作地点"}
                        },
                    },
                )
            )
        return row_postings

    def _extract_cupb_detail(
        self,
        document: str,
        source_url: str,
        source: dict[str, Any],
        title_hint: str,
    ) -> RawPosting | None:
        """Backward-compatible single-notice helper used by existing callers."""
        postings = self._extract_cupb_details(document, source_url, source, title_hint)
        return postings[0] if postings else None

    @staticmethod
    def _decode_cupb_embedded_content(document: str) -> str | None:
        """Decode public UEditor content embedded by the CUPB announcement page.

        Some public detail pages ship their article body in a zlib-compressed,
        Base64-encoded JavaScript literal. This decodes only data already returned
        by that page; it does not use a login, private endpoint, or browser bypass.
        """
        match = re.search(
            r"""Base64\.decode\s*\(\s*unzip\s*\(\s*[\"'](?P<payload>[^\"']+)[\"']""",
            document,
        )
        if not match:
            return None
        try:
            compressed = base64.b64decode(match.group("payload"), validate=True)
            packed_text = zlib.decompress(compressed).decode("utf-8")
            _, encoded_html = packed_text.split(None, 1)
            html = base64.b64decode(encoded_html, validate=True).decode("utf-8")
        except (
            ValueError,
            UnicodeDecodeError,
            binascii.Error,
            zlib.error,
        ):
            return None
        # The page's helper leaves a short view marker before the HTML fragment.
        return re.sub(r"^view\d+d\s*", "", html).strip() or None

    @staticmethod
    def _cupb_table_fields(soup: BeautifulSoup) -> dict[str, str]:
        """Read label/value rows from public CUPB recruitment-announcement tables."""
        records = OfficialSourceCollector._cupb_table_records(soup)
        return records[0] if records else {}

    @staticmethod
    def _cupb_table_records(soup: BeautifulSoup) -> list[dict[str, str]]:
        """Return every structured CUPB table row as an independent record."""
        aliases = {
            "岗位": "岗位",
            "岗位需求": "岗位",
            "职位": "岗位",
            "专业范围": "专业范围",
            "专业要求": "专业范围",
            "需求专业": "专业范围",
            "面向对象": "面向对象",
            "学历要求": "学历要求",
            "工作地点": "工作地点",
        }
        label_fields: dict[str, str] = {}
        for row in soup.select("table tr"):
            cells = [clean_text(cell.get_text(" ", strip=True)) for cell in row.select("th, td")]
            cells = [cell for cell in cells if cell]
            if len(cells) < 2:
                continue
            label = aliases.get(cells[0].rstrip("：:"))
            if not label:
                continue
            value = clean_text(" ".join(cells[1:]))
            if value:
                label_fields[label] = value
        # Job-board announcements also use a conventional header row followed
        # by one or more records (序号 / 职位信息 / 需求专业 / ...), rather than
        # label-value rows.  Read the first complete record without guessing
        # beyond the table's named columns.
        column_aliases = {
            "岗位": ("职位信息", "岗位名称", "岗位", "职位"),
            "专业范围": ("需求专业", "专业要求", "专业范围"),
            "学历要求": ("学历要求", "学历", "学位"),
            "工作地点": ("工作地点", "工作城市", "工作地区", "地点"),
        }
        records: list[dict[str, str]] = []
        for table in soup.select("table"):
            rows: list[list[str]] = []
            for row in table.select("tr"):
                cells = [
                    clean_text(cell.get_text(" ", strip=True))
                    for cell in row.select("th, td")
                ]
                if any(cells):
                    rows.append(cells)
            if len(rows) < 2:
                continue
            header_index = next(
                (
                    index
                    for index, row in enumerate(rows[:3])
                    if any("职位信息" in cell or "岗位名称" in cell for cell in row)
                ),
                None,
            )
            if header_index is None:
                continue
            header = rows[header_index]
            indexes = {
                field: next(
                    (
                        column
                        for column, cell in enumerate(header)
                        if any(alias in cell for alias in aliases)
                    ),
                    None,
                )
                for field, aliases in column_aliases.items()
            }
            for data_row in rows[header_index + 1 :]:
                record = {
                    field: clean_text(data_row[index])
                    for field, index in indexes.items()
                    if index is not None and index < len(data_row) and data_row[index]
                }
                if record.get("岗位") or record.get("专业范围"):
                    records.append(record)
            if records:
                break
        if records:
            for record in records:
                inline_job = record.get("岗位", "")
                if inline_job and "学历要求" not in record:
                    degree_match = re.search(
                        r"(?P<degree>大专|本科|硕士|博士)(?:研究生)?(?:及以上)?",
                        inline_job,
                    )
                    if degree_match:
                        record["学历要求"] = degree_match.group("degree")
                if inline_job and "工作地点" not in record:
                    location_match = re.search(
                        r"(?:\d[\d,]*(?:\.\d+)?\s*(?:[-~至]\s*\d[\d,]*(?:\.\d+)?)?\s*)"
                        r"(?P<location>[^\s]+(?:市|区|县|省|自治区|特别行政区))\s*"
                        r"(?:全职|实习|兼职)",
                        inline_job,
                    )
                    if location_match:
                        record["工作地点"] = clean_text(location_match.group("location"))
            return records
        fields = dict(label_fields)
        inline_job = fields.get("岗位", "")
        if inline_job:
            if "学历要求" not in fields:
                degree_match = re.search(
                    r"(?P<degree>大专|本科|硕士|博士)(?:研究生)?(?:及以上)?",
                    inline_job,
                )
                if degree_match:
                    fields["学历要求"] = degree_match.group("degree")
            if "工作地点" not in fields:
                # CUPB's compact job row is commonly formatted as:
                # title + salary + city + employment type + degree.
                location_match = re.search(
                    r"(?:\d[\d,]*(?:\.\d+)?\s*(?:[-~至]\s*\d[\d,]*(?:\.\d+)?)?\s*)"
                    r"(?P<location>[^\s]+(?:市|区|县|省|自治区|特别行政区))\s*"
                    r"(?:全职|实习|兼职)",
                    inline_job,
                )
                if location_match:
                    fields["工作地点"] = clean_text(location_match.group("location"))
        return [fields] if fields else []

    @staticmethod
    def _cupb_major_evidence(text: str) -> str:
        """Keep only explicit qualification snippets from a free-form notice."""
        segments = [
            clean_text(segment)
            for segment in re.split(r"[。！？；;\n]", text)
            if clean_text(segment)
        ]
        selected = []
        for segment in segments:
            if not extract_major_tags(segment):
                continue
            # Unit introductions often mention geophysics, resources or
            # energy as business context.  They are not qualification
            # evidence unless the same sentence carries a hiring/eligibility
            # marker.
            if re.search(r"单位|公司|企业|业务范围|主营|成立于|简介", segment) and not re.search(
                r"专业|要求|学历|学位|岗位|招聘|应聘|博士后|毕业生", segment
            ):
                continue
            selected.append(segment)
        if not selected:
            return ""
        return clean_text("；".join(selected))[:1800]

    @staticmethod
    def _cupb_location_evidence(text: str) -> str | None:
        """Extract an explicitly stated institution/work location from prose."""
        patterns = (
            r"(?:注册在|位于|坐落于)\s*[^，。；;\n]{0,120}?[—-]\s*"
            r"([^，。；;\n]{2,20})",
            r"(?:注册在|位于|坐落于|驻地为|工作地点为|工作地点是)\s*"
            r"([^，。；;\n]{2,40})",
            r"(?:通讯地址|通信地址|办公地址|联系地址|地址)\s*[：:]\s*"
            r"([^，。；;\n]{2,80})",
        )
        for pattern in patterns:
            match = re.search(pattern, text)
            if match:
                value = clean_text(match.group(1)).strip(" ：:—-")
                # Do not treat a following process sentence as a location.
                value = re.split(r"\s+(?:5年|招聘|培养|员工|岗位|报名)", value, maxsplit=1)[0]
                # Address evidence often continues with a street, postal code,
                # or institution name. Keep the city/province prefix so the
                # public listing remains concise and filterable.
                city_match = re.search(
                    r"((?:北京|上海|天津|重庆|香港|澳门)[市区]?|"
                    r"[^，。；;\s]{2,12}(?:省|自治区|特别行政区)?"
                    r"[^，。；;\s]{2,12}(?:市|州|盟|县|区))",
                    value,
                )
                if city_match:
                    value = city_match.group(1)
                if value:
                    return value[:80]
        return None

    def _extract_cas_job_detail(
        self,
        document: str,
        source_url: str,
        title_hint: str,
        employer: str,
        location_hint: str | None,
    ) -> RawPosting:
        soup = BeautifulSoup(document, "html.parser")
        page_text = clean_text(soup.get_text(" ", strip=True))
        title_node = soup.select_one("h1, h2, .post_title, .content_title")
        candidate_title = clean_text(
            title_node.get_text(" ", strip=True) if title_node else ""
        )
        generic_titles = {
            "岗位描述",
            "岗位职责",
            "应聘条件",
            "招聘岗位",
            "联系方式",
        }
        title = (
            candidate_title
            if candidate_title and candidate_title not in generic_titles
            else title_hint
        )
        if title not in page_text:
            page_text = f"{title} {page_text}"
        location = self._label_value(
            page_text,
            "工作地点",
            (
                "工作经验",
                "学历要求",
                "招聘人数",
                "薪资水平",
                "发布时间",
                "截止时间",
                "学科领域",
                "专业描述",
                "岗位描述",
            ),
        ) or location_hint
        discipline = self._label_value(
            page_text,
            "学科领域",
            ("专业描述", "专业技术等级", "招聘类型", "岗位描述", "岗位职责"),
        )
        professional_requirements = self._label_value(
            page_text,
            "专业描述",
            (
                "专业技术等级",
                "招聘类型",
                "岗位描述",
                "岗位职责",
                "应聘条件",
            ),
        )
        summary = self._cas_summary(page_text, title)
        return RawPosting(
            title=title,
            employer=employer,
            source_url=normalize_url(source_url),
            application_url=normalize_url(source_url),
            text=page_text,
            summary=summary,
            published_date=extract_published_date(page_text),
            deadline_date=extract_deadline(page_text),
            location=location,
            external_id=self._external_id_from_url(source_url),
            match_text=clean_text(
                f"{title} {discipline or ''} {professional_requirements or ''}"
            ),
        )

    def _get(self, url: str, source: dict[str, Any]) -> requests.Response:
        # CUPB's public router uses the documented ``//page/N`` spelling for
        # pagination.  Its path is semantically significant, so preserve it
        # only for this source and this route; all other requests retain the
        # normal canonical URL policy.
        if source.get("source_type") == "cupb_career" and re.search(
            r"//page/\d+(?:/|$)", str(url)
        ):
            normalized = self._cupb_listing_url(url)
        else:
            normalized = normalize_url(url)
        if not self._robots_allowed(normalized):
            raise SourceSkipped(f"robots.txt does not permit collection: {normalized}")
        try:
            response = self.request_policy.request(
                "GET",
                normalized,
                timeout=self.settings.request_timeout_seconds,
                allow_redirects=True,
            )
            response.raise_for_status()
            if (
                not response.encoding
                or response.encoding.lower() in {"iso-8859-1", "ascii"}
            ) and response.apparent_encoding:
                response.encoding = response.apparent_encoding
            return response
        except REQUEST_ERRORS as error:
            raise SourceCollectionError(f"Request failed for {normalized}: {error}") from error

    def _post_json(
        self,
        url: str,
        payload: dict[str, Any],
        *,
        headers: dict[str, str] | None = None,
    ) -> requests.Response:
        normalized = normalize_url(url)
        if not self._robots_allowed(normalized):
            raise SourceSkipped(f"robots.txt does not permit collection: {normalized}")
        try:
            response = self.request_policy.request(
                "POST",
                normalized,
                json=payload,
                headers=headers,
                timeout=self.settings.request_timeout_seconds,
                allow_redirects=True,
            )
            response.raise_for_status()
            return response
        except REQUEST_ERRORS as error:
            raise SourceCollectionError(f"Request failed for {normalized}: {error}") from error

    def _robots_allowed(self, url: str) -> bool:
        parsed = urlparse(url)
        root = f"{parsed.scheme}://{parsed.netloc}"
        parser = self._robots_cache.get(root)
        if parser is None:
            parser = RobotFileParser()
            robots_url = urljoin(root, "/robots.txt")
            try:
                response = self.request_policy.request(
                    "GET",
                    robots_url,
                    timeout=min(self.settings.request_timeout_seconds, 10),
                    allow_redirects=True,
                )
                if response.status_code == 404:
                    parser.parse([])
                elif response.ok:
                    parser.parse(response.text.splitlines())
                else:
                    raise SourceSkipped(
                        f"Unable to verify robots.txt for {root}: HTTP {response.status_code}"
                    )
            except REQUEST_ERRORS as error:
                raise SourceSkipped(
                    f"Unable to verify robots.txt for {root}: {error}"
                ) from error
            self._robots_cache[root] = parser
        return parser.can_fetch(USER_AGENT, url)

    @staticmethod
    def _select_first(soup: BeautifulSoup, selector: str) -> Any | None:
        """Return the first match in the configured selector order.

        BeautifulSoup treats a comma-separated selector as one CSS query and
        returns the first document-order match, which can make a broad fallback
        such as ``body`` win over the intended article container. Splitting the
        configured alternatives preserves their explicit priority.
        """
        for candidate in (part.strip() for part in selector.split(",")):
            if not candidate:
                continue
            node = soup.select_one(candidate)
            if node is not None:
                return node
        return None

    @staticmethod
    def _title_matches_filters(title: str, config: dict[str, Any]) -> bool:
        normalized = clean_text(title)
        if any(
            re.search(pattern, normalized, re.IGNORECASE)
            for pattern in config.get("excluded_title_patterns", [])
        ):
            return False
        required = config.get("required_title_patterns", [])
        return not required or any(
            re.search(pattern, normalized, re.IGNORECASE) for pattern in required
        )

    @staticmethod
    def _is_non_vacancy_notice_title(title: str) -> bool:
        normalized = clean_text(title)
        return any(
            re.search(pattern, normalized, re.IGNORECASE)
            for pattern in NON_VACANCY_TITLE_PATTERNS
        )

    @staticmethod
    def _accept_candidate(
        value: str,
        source: dict[str, Any],
        *,
        exclude_patterns: list[str] | None = None,
        require_recruitment_word: bool | None = None,
        require_major_match: bool | None = None,
    ) -> bool:
        config = source["config"]
        if config.get("accept_all_entries"):
            return True
        normalized = clean_text(value)
        if exclude_patterns is None:
            exclude_patterns = config.get("exclude_patterns", [])
        if any(
            re.search(pattern, normalized, re.IGNORECASE)
            for pattern in exclude_patterns
        ):
            return False
        has_recruitment_word = looks_like_recruitment(normalized)
        has_major_match = bool(extract_major_tags(normalized))
        recruitment_required = (
            config.get("require_recruitment_word")
            if require_recruitment_word is None
            else require_recruitment_word
        )
        major_required = (
            config.get("require_major_match")
            if require_major_match is None
            else require_major_match
        )
        if recruitment_required and not has_recruitment_word:
            return False
        if major_required and not has_major_match:
            return False
        include_keywords = config.get("include_keywords", [])
        if include_keywords and not any(
            keyword.lower() in normalized.lower() for keyword in include_keywords
        ):
            return False
        return has_recruitment_word or has_major_match

    def _item_limit(self, source: dict[str, Any]) -> int:
        value = source["config"].get("max_items", self.settings.max_source_items)
        try:
            return max(1, min(int(value), self.settings.max_source_items))
        except (TypeError, ValueError):
            return self.settings.max_source_items

    @staticmethod
    def _external_id_from_url(url: str) -> str:
        parsed = urlparse(url)
        return f"{parsed.path}?{parsed.query}".rstrip("?")

    def _wait(self, source: dict[str, Any]) -> None:
        try:
            seconds = float(
                source["config"].get("request_interval_seconds", 0.75)
            )
        except (TypeError, ValueError):
            seconds = 0.75
        if seconds > 0:
            time.sleep(min(seconds, 5.0))

    @staticmethod
    def _employer_from_title(title: str, fallback: str) -> str:
        markers = (
            "校园招聘",
            "社会招聘",
            "公开招聘",
            "招聘公告",
            "招聘简章",
            "招聘启事",
            "招聘",
            "招录",
            "招考",
            "人才引进",
            "202",
        )
        positions = [title.find(marker) for marker in markers if title.find(marker) > 1]
        if not positions:
            return fallback
        candidate = title[: min(positions)].strip(" -|丨：:，,。.")
        return candidate if len(candidate) >= 3 else fallback

    @staticmethod
    def _label_value(
        text: str,
        label: str,
        end_labels: tuple[str, ...],
    ) -> str | None:
        match = re.search(
            rf"{re.escape(label)}\s*[:：]\s*(.+?)(?="
            + "|".join(re.escape(item) + r"\s*[:：]" for item in end_labels)
            + r"|$)",
            text,
        )
        if not match:
            return None
        value = clean_text(match.group(1))
        return value[:200] or None

    @staticmethod
    def _cas_summary(page_text: str, title: str) -> str:
        fields = []
        for label in ("所在部门", "学历要求", "工作地点", "招聘人数", "截止时间"):
            match = re.search(
                rf"{re.escape(label)}\s*[:：]\s*(.+?)(?="
                r"(?:工作经验|学历要求|工作地点|招聘人数|薪资水平|发布时间|"
                r"截止时间|学科领域|专业描述|招聘类型|岗位描述)\s*[:：]|$)",
                page_text,
            )
            if match:
                fields.append(f"{label}：{clean_text(match.group(1))}")
        return clean_text("；".join(fields))[:420] or title

    @staticmethod
    def _find_application_url(soup: BeautifulSoup, base_url: str) -> str | None:
        for anchor in soup.find_all("a"):
            label = clean_text(anchor.get_text(" ", strip=True) or anchor.get("title"))
            href = anchor.get("href")
            context_node = anchor.find_parent("p") or anchor.parent
            context = clean_text(
                context_node.get_text(" ", strip=True) if context_node else label
            )
            candidate = f"{label} {context}".lower()
            if not href or not any(
                word in candidate
                for word in (
                    "报名",
                    "申请",
                    "投递",
                    "网申",
                    "应聘",
                    "招聘平台",
                    "招聘网",
                    "apply",
                    "application",
                )
            ):
                continue
            candidate_url = normalize_url(urljoin(base_url, href))
            if urlparse(candidate_url).scheme not in {"http", "https"}:
                # Government CMS templates often expose a ``javascript:;``
                # placeholder next to the real announcement.  Never expose
                # it as an application link in the public record.
                continue
            # An application form or job-list attachment is supporting evidence,
            # not a live application route. The original announcement remains the
            # authoritative link when a public page only provides attachments.
            if re.search(r"\.(?:pdf|docx?|xlsx?|csv|zip|rar)$", urlparse(candidate_url).path, re.IGNORECASE):
                continue
            return candidate_url
        return None

    @staticmethod
    def _published_date_from_meta(soup: BeautifulSoup) -> str | None:
        for selector in (
            "meta[property='article:published_time']",
            "meta[name='PubDate']",
            "meta[name='pubdate']",
            "meta[name='ArticleDate']",
            "meta[name='publishdate']",
            "meta[name='date']",
            "time",
        ):
            element = soup.select_one(selector)
            if not element:
                continue
            value = element.get("content") or element.get("datetime") or element.get_text()
            parsed = extract_published_date(value)
            if parsed:
                return parsed
        return None

    @staticmethod
    def _xml_text(entry: ET.Element, tag: str) -> str:
        element = entry.find(tag)
        return element.text if element is not None and element.text else ""

    @staticmethod
    def _rss_link(entry: ET.Element) -> str:
        link = OfficialSourceCollector._xml_text(entry, "link")
        if link:
            return link
        atom_link = entry.find("{http://www.w3.org/2005/Atom}link")
        if atom_link is not None:
            return atom_link.get("href", "")
        return ""

    @staticmethod
    def _nested_value(value: Any, path: str) -> Any:
        current = value
        for part in path.split("."):
            if not part:
                continue
            if not isinstance(current, dict):
                return None
            current = current.get(part)
        return current


def load_source_registry(path: str) -> list[dict[str, Any]]:
    try:
        with open(path, "r", encoding="utf-8") as handle:
            payload = json.load(handle)
    except OSError as error:
        raise SourceCollectionError(f"Cannot read source registry: {error}") from error
    try:
        return validate_source_registry(payload)
    except ContractValidationError as error:
        raise SourceCollectionError(f"Invalid source registry: {error}") from error


def load_source_registries(paths: list[str]) -> list[dict[str, Any]]:
    """Load a primary registry and optional supplementary source matrices.

    A separate provincial file keeps the active collector configuration compact
    while allowing the nationwide source network to remain versioned and
    inspectable.  IDs are global so a duplicate cannot silently overwrite a
    source in SQLite.
    """
    sources: list[dict[str, Any]] = []
    source_ids: set[str] = set()
    for path in paths:
        for source in load_source_registry(path):
            source_id = str(source["id"])
            if source_id in source_ids:
                raise SourceCollectionError(f"Duplicate source id: {source_id}")
            source_ids.add(source_id)
            sources.append(source)
    return sources
