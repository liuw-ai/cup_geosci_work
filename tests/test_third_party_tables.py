import hashlib
from pathlib import Path

import pytest

from job_hub import third_party_tables
from job_hub.third_party_tables import (
    discover_table_links,
    fetch_and_normalize_third_party_table,
    ingest_third_party_table_file,
    table_candidates,
)


def test_third_party_table_rows_become_private_candidates_with_provenance():
    candidates = table_candidates(
        {
            "batch": {
                "id": "offcn-2027",
                "province": "北京",
                "rows": [["A01", "自然资源局", "地质管理", "地质学", "硕士", 1, "北京岗位!4", "北京"]],
            },
            "download": {
                "sha256": "abc123",
                "bytes": 128,
                "content_type": "application/vnd.ms-excel",
                "captured_at": "2026-10-17T00:00:00Z",
                "retained_path": "runtime/third-party/offcn.xlsx",
            },
        },
        source_id="offcn-discovery",
        third_party_table_url="https://www.offcn.com/table.xlsx",
        official_notice_url="https://gov.example/notice",
        official_application_url="https://gov.example/apply",
        discovery_page_url="https://www.offcn.com/guide/",
    )

    assert len(candidates) == 1
    candidate = candidates[0]
    assert candidate["official_url"] == "https://gov.example/notice"
    assert candidate["third_party_table_url"].endswith("table.xlsx")
    assert candidate["metadata"]["publication_status"] == "candidate_only"
    assert candidate["lead_url"].endswith("#row-北京岗位!4")
    assert candidate["metadata"]["application_url"] == "https://gov.example/apply"
    assert candidate["metadata"]["table_evidence"]["sha256"] == "abc123"
    assert candidate["metadata"]["table_evidence"]["discovery_page_url"] == "https://www.offcn.com/guide/"


def test_table_discovery_rejects_page_outside_registered_source_host():
    with pytest.raises(ValueError, match="discovery page host"):
        discover_table_links(
            '<a href="https://www.offcn.com/table.xlsx">职位表</a>',
            page_url="https://untrusted.example/guide",
            source_id="offcn-discovery",
        )


def test_fetch_retains_official_table_evidence_and_hash(monkeypatch, tmp_path):
    from openpyxl import Workbook

    fixture = tmp_path / "fixture.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "内蒙古岗位"
    sheet.append(["岗位代码", "用人单位", "岗位名称", "专业要求", "学历要求", "招录人数", "工作地点"])
    sheet.append(["A01", "自然资源局", "地质管理", "地质学", "硕士", 1, "呼和浩特"])
    workbook.save(fixture)
    body = fixture.read_bytes()
    monkeypatch.setattr(
        third_party_tables,
        "_bounded_download",
        lambda url, timeout, max_bytes: (
            body,
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        ),
    )
    monkeypatch.setattr(
        third_party_tables,
        "_robots_permit",
        lambda url, timeout: {
            "robots_url": "https://www.offcn.com/robots.txt",
            "robots_status": 200,
            "robots_allowed": True,
            "robots_checked_at": "2026-10-18T00:00:00Z",
            "robots_detail": "robots.txt permits access",
        },
    )

    payload = fetch_and_normalize_third_party_table(
        source_id="offcn-discovery",
        third_party_table_url="https://www.offcn.com/2026/positions.xlsx",
        batch_id="offcn-test-2026",
        position_type="civil_service",
        province="内蒙古",
        official_notice_url="https://www.gov.cn/notice/2026",
        official_table_url="https://www.gov.cn/notice/2026/positions.xlsx",
        discovery_page_url="https://www.offcn.com/guide/2026",
        storage_dir=tmp_path,
    )

    assert payload["provenance"]["official_table_url"] == (
        "https://www.gov.cn/notice/2026/positions.xlsx"
    )
    assert payload["download"]["sha256"] == hashlib.sha256(body).hexdigest()
    assert payload["download"]["robots_allowed"] is True
    retained = Path(payload["download"]["retained_path"])
    assert retained.is_file()
    assert retained.read_bytes() == body
    assert payload["candidates"]
    assert payload["candidates"][0]["official_table_url"] == (
        "https://www.gov.cn/notice/2026/positions.xlsx"
    )


def test_fetch_refuses_robots_denied_before_table_download(monkeypatch):
    downloaded = False

    def fail_download(*args, **kwargs):
        nonlocal downloaded
        downloaded = True
        raise AssertionError("robots denial must stop before downloading")

    monkeypatch.setattr(
        third_party_tables,
        "_robots_permit",
        lambda url, timeout: (_ for _ in ()).throw(
            third_party_tables.ThirdPartyTableAccessLimited("robots denied")
        ),
    )
    monkeypatch.setattr(third_party_tables, "_bounded_download", fail_download)

    with pytest.raises(third_party_tables.ThirdPartyTableAccessLimited, match="robots denied"):
        fetch_and_normalize_third_party_table(
            source_id="offcn-discovery",
            third_party_table_url="https://www.offcn.com/2026/positions.xlsx",
            batch_id="offcn-robots-blocked",
            position_type="civil_service",
            province="北京",
        )
    assert downloaded is False


def test_operator_provided_table_file_uses_same_private_evidence_contract(tmp_path):
    from openpyxl import Workbook

    table = tmp_path / "provided.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["岗位代码", "用人单位", "岗位名称", "专业要求", "学历要求", "招录人数", "工作地点"])
    sheet.append(["B01", "自然资源局", "地质调查", "资源勘查工程", "本科", 1, "北京"])
    workbook.save(table)

    payload = ingest_third_party_table_file(
        table_path=table,
        source_id="offcn-discovery",
        third_party_table_url="https://www.offcn.com/2026/positions.xlsx",
        batch_id="offcn-provided-2026",
        position_type="civil_service",
        province="北京",
        discovery_page_url="https://www.offcn.com/guide/2026",
        storage_dir=tmp_path / "retained",
    )

    assert payload["download"]["capture_mode"] == "operator_provided_file"
    assert payload["download"]["sha256"] == hashlib.sha256(table.read_bytes()).hexdigest()
    assert payload["candidates"]
    assert payload["candidates"][0]["official_url"] is None
    assert payload["candidates"][0]["official_application_url"] is None


def test_operator_provided_table_file_still_requires_registered_source_url(tmp_path):
    table = tmp_path / "provided.csv"
    table.write_text("岗位代码,用人单位,岗位名称,专业要求,学历要求,招录人数,工作地点\n", encoding="utf-8")

    with pytest.raises(ValueError, match="host is not allowed"):
        ingest_third_party_table_file(
            table_path=table,
            source_id="offcn-discovery",
            third_party_table_url="https://untrusted.example/table.xlsx",
            batch_id="offcn-provided-invalid",
            position_type="civil_service",
            province="北京",
        )


def test_huatu_asset_host_is_registered_for_real_position_tables():
    adapters = third_party_tables.load_table_adapters()
    assert "u3.huatu.com" in adapters["huatu-discovery"]["allowed_hosts"]

    payload = ingest_third_party_table_file(
        table_path=Path(__file__).resolve().parents[1]
        / "runtime"
        / "third-party-captures"
        / "huatu"
        / "2026-10-16-gansu-cmxy"
        / "huatu-attachment-1-02248460f53feb09.xlsx",
        source_id="huatu-discovery",
        third_party_table_url="https://u3.huatu.com/uploads/soft/261016/df.xlsx",
        batch_id="huatu-asset-host-contract",
        position_type="public_institution",
        province="甘肃",
        discovery_page_url="https://sydw.huatu.com/2026/1016/2875151.html",
    )
    assert payload["download"]["capture_mode"] == "operator_provided_file"
    assert payload["candidates"]


def test_table_candidate_persists_discovery_source_for_funnel_audit(tmp_path):
    from openpyxl import Workbook
    from job_hub.db import Database

    table = tmp_path / "provided.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["岗位代码", "用人单位", "岗位名称", "专业要求", "学历要求", "招录人数", "工作地点"])
    sheet.append(["B02", "自然资源局", "地质调查", "地质学", "本科", 1, "北京"])
    workbook.save(table)

    payload = ingest_third_party_table_file(
        table_path=table,
        source_id="offcn-discovery",
        third_party_table_url="https://www.offcn.com/2026/positions.xlsx",
        batch_id="offcn-funnel-2026",
        position_type="civil_service",
        province="北京",
    )
    database = Database(tmp_path / "isolated.sqlite3")
    database.initialize()
    stored = database.create_candidate_lead(payload["candidates"][0])

    assert stored["discovery_source_id"] == "offcn-discovery"
    assert stored["third_party_table_url"] == "https://www.offcn.com/2026/positions.xlsx"
    assert stored["official_url"] is None
    assert stored["verification_status"] == "candidate"


def test_table_rows_from_one_url_keep_position_level_identity(tmp_path):
    from openpyxl import Workbook
    from job_hub.db import Database

    table = tmp_path / "provided.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["岗位代码", "用人单位", "岗位名称", "专业要求", "学历要求", "招聘人数", "工作地点"])
    sheet.append(["B03", "某省地质院", "地质调查岗", "地质学", "本科", 1, "北京"])
    sheet.append(["B04", "某省地质院", "水文地质岗", "地质工程", "硕士", 1, "北京"])
    workbook.save(table)

    payload = ingest_third_party_table_file(
        table_path=table,
        source_id="offcn-discovery",
        third_party_table_url="https://www.offcn.com/2026/positions.xlsx",
        batch_id="offcn-row-identity-2026",
        position_type="public_institution",
        province="北京",
    )
    database = Database(tmp_path / "isolated.sqlite3")
    database.initialize()
    for candidate in payload["candidates"]:
        database.create_candidate_lead(candidate)

    assert len(database.list_candidate_leads(discovery_source_id="offcn-discovery")) == 2
