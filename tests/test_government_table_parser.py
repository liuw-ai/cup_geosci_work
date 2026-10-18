import json
from pathlib import Path

from job_hub.government_positions import load_position_registry
from job_hub.government_table_parser import (
    load_position_table,
    load_position_table_with_report,
    parse_position_table_manifest,
    parse_position_table_rows,
)


def test_batch_parser_maps_civil_service_aliases_and_keeps_rejected_rows() -> None:
    result = parse_position_table_rows(
        [
            {
                "职位代码": "1001",
                "招录机关": "某市自然资源局",
                "招考职位": "地质资源管理",
                "专业条件": "地质学、地质工程",
                "学历学位": "硕士研究生",
                "招录人数": "2",
                "工作地点": "北京",
            },
            {
                "职位代码": "1002",
                "招录机关": "某市自然资源局",
                "招考职位": "综合管理",
                "专业条件": "",
                "学历学位": "本科",
                "招录人数": "1",
            },
        ],
        batch_id="beijing-gwy-2027",
        source_id="beijing-civil-service-2027",
        position_type="civil_service",
        province="北京",
        official_notice_url="https://gwy.example.gov.cn/notice/1",
        official_attachment_url="https://gwy.example.gov.cn/notice/1.xlsx",
    )
    assert result["metrics"] == {
        "input_rows": 2,
        "accepted_rows": 1,
        "rejected_rows": 1,
        "field_contract": "position_code, employer, title, major, degree, headcount",
    }
    assert result["batch"]["rows"][0][0:3] == ["1001", "某市自然资源局", "地质资源管理"]
    assert result["batch"]["rows"][0][-1] == "北京"
    assert result["rejected_rows"][0]["missing"] == ["major"]


def test_batch_parser_produces_registry_compatible_position_batch() -> None:
    result = parse_position_table_rows(
        [
            {
                "岗位代码": "A01",
                "招聘单位": "某省地质调查院",
                "岗位名称": "地质调查岗",
                "专业要求": "地质工程",
                "学历要求": "硕士研究生及以上",
                "招聘人数": 1,
                "工作地点": "长沙",
            }
        ],
        batch_id="hunan-2027",
        source_id="hunan-public-2027",
        position_type="public_institution",
        province="湖南",
        official_notice_url="https://example.gov.cn/notice/1",
        official_attachment_url="https://example.gov.cn/notice/1.xlsx",
        location_evidence_policy="row_exact",
    )
    payload = {
        "version": 1,
        "as_of": "2026-10-11",
        "records": [],
        "position_batches": [result["batch"]],
        "source_opening_dates": {},
        "source_opening_at": {},
        "source_assessments": [],
    }
    registry_path = Path(".tmp-government-parser-registry.json")
    registry_path.write_text(__import__("json").dumps(payload, ensure_ascii=False), encoding="utf-8")
    try:
        registry = load_position_registry(registry_path)
        assert registry["records"][0]["position_code"] == "A01"
        assert registry["records"][0]["location"] == "长沙"
    finally:
        registry_path.unlink(missing_ok=True)


def test_load_position_table_accepts_json_rows(tmp_path) -> None:
    path = tmp_path / "positions.json"
    path.write_text(
        '{"rows":[{"岗位代码":"1","岗位名称":"地质岗"}]}',
        encoding="utf-8",
    )
    assert load_position_table(path)[0]["岗位代码"] == "1"


def test_excel_parser_reads_all_data_sheets_and_preserves_physical_rows(tmp_path) -> None:
    from openpyxl import Workbook

    workbook = Workbook()
    cover = workbook.active
    cover.title = "说明"
    cover.append(["此页不是职位表"])
    first = workbook.create_sheet("北京岗位")
    first.append(["招聘计划"])
    first.append([])
    headers = ["岗位代码", "招聘单位", "岗位名称", "专业要求", "学历要求", "招聘人数", "工作地点"]
    first.append(headers)
    first.append(["A01", "地质调查院", "地质岗", "地质学", "硕士", 1, "北京"])
    first.append([])
    first.append(["A02", "自然资源局", "地质工程岗", "地质工程", "硕士", 2, "北京"])
    second = workbook.create_sheet("北京分支岗位")
    second.append(headers)
    second.append(["B01", "地勘中心", "资源勘查岗", "资源勘查工程", "本科", 1, "北京"])
    path = tmp_path / "positions.xlsx"
    workbook.save(path)

    loaded = load_position_table_with_report(path)
    assert len(loaded["rows"]) == 3
    assert loaded["skipped_sheets"] == [{"sheet_name": "说明", "reason": "header_not_found"}]
    result = parse_position_table_rows(
        loaded["rows"],
        batch_id="multi-sheet",
        source_id="government-test",
        position_type="public_institution",
        province="北京",
        official_notice_url="https://example.gov.cn/notice",
        official_attachment_url="https://example.gov.cn/positions.xlsx",
    )
    assert [row[6] for row in result["batch"]["rows"]] == [
        "北京岗位!4；岗位代码A01",
        "北京岗位!6；岗位代码A02",
        "北京分支岗位!2；岗位代码B01",
    ]


def test_excel_parser_accepts_adjacent_two_row_headers(tmp_path) -> None:
    from openpyxl import Workbook

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "岗位列表"
    sheet.append(["岗位代码", "招聘单位", "招聘岗位类别", "招聘人数", "招聘专业", "", "", ""])
    sheet.append(["", "", "", "", "专业", "专业代码", "学历", "其他条件"])
    sheet.append(["YJ2601", "甘肃财贸职业学院", "专任教师", 1, "地质学", "0709", "博士研究生", "无"])
    path = tmp_path / "two-row-header.xlsx"
    workbook.save(path)

    loaded = load_position_table_with_report(path)
    assert loaded["sheets"][0]["header_row"] == 1
    assert loaded["rows"][0]["__row_number"] == 3

    result = parse_position_table_rows(
        loaded["rows"],
        batch_id="gansu-cmxy",
        source_id="huatu-discovery",
        position_type="public_institution",
        province="甘肃",
        official_notice_url="https://www.gscmxy.edu.cn/info/1010/4255.htm",
        official_attachment_url="https://www.gscmxy.edu.cn/system/_content/download.jsp?urltype=news.DownloadAttachUrl&owner=1796077919&wbfileid=4761750",
        location="甘肃省兰州市",
    )
    assert result["metrics"]["accepted_rows"] == 1
    assert result["batch"]["rows"][0][0:6] == [
        "YJ2601",
        "甘肃财贸职业学院",
        "专任教师",
        "地质学",
        "博士研究生",
        1,
    ]


def test_excel_parser_expands_merged_employer_cells(tmp_path) -> None:
    from openpyxl import Workbook

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "岗位列表"
    sheet.append(["岗位代码", "招聘单位", "岗位名称", "专业要求", "学历要求", "招聘人数", "工作地点"])
    sheet.append(["A01", "某省地质院", "地质岗", "地质学", "硕士", 1, "甘肃"])
    sheet.append(["A02", None, "地质调查岗", "地质工程", "博士", 1, "甘肃"])
    sheet.merge_cells("B2:B3")
    path = tmp_path / "merged-employer.xlsx"
    workbook.save(path)

    loaded = load_position_table_with_report(path)
    result = parse_position_table_rows(
        loaded["rows"],
        batch_id="merged-employer",
        source_id="government-test",
        position_type="public_institution",
        province="甘肃",
        official_notice_url="https://example.gov.cn/notice",
        official_attachment_url="https://example.gov.cn/table.xlsx",
    )
    assert result["metrics"]["accepted_rows"] == 2
    assert result["batch"]["rows"][1][1] == "某省地质院"


def test_batch_parser_rejects_explicit_cross_province_sheet() -> None:
    result = parse_position_table_rows(
        [{
            "__sheet_name": "天津岗位",
            "__row_number": 8,
            "岗位代码": "B01",
            "招聘单位": "地勘中心",
            "岗位名称": "地质岗",
            "专业要求": "地质学",
            "学历要求": "硕士",
            "招聘人数": 1,
            "工作地点": "天津",
        }],
        batch_id="beijing-only",
        source_id="government-test",
        position_type="public_institution",
        province="北京",
        official_notice_url="https://example.gov.cn/notice",
        official_attachment_url="https://example.gov.cn/positions.xlsx",
    )
    assert result["metrics"]["accepted_rows"] == 0
    assert result["rejected_rows"][0]["reason"] == "province_mismatch"
    assert result["rejected_rows"][0]["row_number"] == 8


def test_excel_parser_reports_missing_header_in_requested_sheet(tmp_path) -> None:
    from openpyxl import Workbook
    import pytest

    workbook = Workbook()
    workbook.active.title = "说明"
    workbook.active.append(["不是职位表"])
    path = tmp_path / "no-positions.xlsx"
    workbook.save(path)

    with pytest.raises(ValueError, match="header"):
        load_position_table(path, sheet_name="说明")


def test_position_table_manifest_batches_multiple_local_official_artifacts(tmp_path) -> None:
    first = tmp_path / "beijing.json"
    first.write_text(
        '{"rows":[{"岗位代码":"BJ-1","用人单位":"北京市自然资源局",'
        '"岗位名称":"地质管理","专业要求":"地质学",'
        '"学历要求":"硕士研究生","招聘人数":1,"工作地点":"北京"}]}',
        encoding="utf-8",
    )
    second = tmp_path / "hunan.json"
    second.write_text(
        '{"rows":[{"岗位代码":"HN-1","用人单位":"湖南省地质院",'
        '"岗位名称":"地质调查","专业要求":"地质工程",'
        '"学历要求":"博士","招聘人数":2,"工作地点":"长沙"}]}',
        encoding="utf-8",
    )
    manifest = tmp_path / "manifest.json"
    manifest.write_text(
        json.dumps(
            {
                "version": 1,
                "tables": [
                    {
                        "path": "beijing.json",
                        "batch_id": "beijing-2027",
                        "source_id": "beijing-hrss",
                        "position_type": "public_institution",
                        "province": "北京",
                        "official_notice_url": "https://rsj.example.gov.cn/bj",
                        "official_attachment_url": "https://rsj.example.gov.cn/bj.json",
                    },
                    {
                        "path": "hunan.json",
                        "batch_id": "hunan-2027",
                        "source_id": "hunan-hrss",
                        "position_type": "public_institution",
                        "province": "湖南",
                        "official_notice_url": "https://rst.example.gov.cn/hn",
                        "official_attachment_url": "https://rst.example.gov.cn/hn.json",
                    },
                ],
            },
            ensure_ascii=False,
        ),
        encoding="utf-8",
    )

    result = parse_position_table_manifest(manifest)
    assert result["metrics"] == {
        "table_count": 2,
        "accepted_rows": 2,
        "rejected_rows": 0,
        "field_contract": "position_code, employer, title, major, degree, headcount",
    }
    assert [batch["id"] for batch in result["batches"]] == [
        "beijing-2027",
        "hunan-2027",
    ]
    assert all(report["load_report"]["sheets"] for report in result["reports"])


def test_position_table_manifest_rejects_duplicate_batch_ids(tmp_path) -> None:
    table = tmp_path / "table.json"
    table.write_text('{"rows":[]}', encoding="utf-8")
    manifest = tmp_path / "manifest.json"
    base = {
        "path": "table.json",
        "batch_id": "duplicate",
        "source_id": "source",
        "position_type": "civil_service",
        "province": "北京",
        "official_notice_url": "https://example.gov.cn/notice",
        "official_attachment_url": "https://example.gov.cn/table.json",
    }
    manifest.write_text(
        json.dumps({"version": 1, "tables": [base, base]}, ensure_ascii=False),
        encoding="utf-8",
    )
    import pytest

    with pytest.raises(ValueError, match="duplicate manifest batch_id"):
        parse_position_table_manifest(manifest)


def test_pdf_position_table_loader_reuses_conservative_attachment_row_parser(tmp_path, monkeypatch) -> None:
    import pypdf

    class FakePage:
        def extract_text(self, extraction_mode=None):
            return (
                "岗位代码    招聘单位    岗位名称    招聘人数    学历    专业\n"
                "202602101  某省地质院  地质技术人员  1  硕士  地质工程"
            )

    class FakeReader:
        pages = [FakePage()]

    monkeypatch.setattr(pypdf, "PdfReader", lambda path: FakeReader())
    path = tmp_path / "positions.pdf"
    path.write_bytes(b"%PDF-fake")

    loaded = load_position_table_with_report(path)

    assert loaded["metadata"]["parser"] == "pypdf"
    assert loaded["sheets"][0]["sheet_name"] == "page-1"
    assert loaded["rows"][0]["__row_number"] == 2
    assert loaded["rows"][0]["职位代码"] == "202602101"
    assert loaded["rows"][0]["专业要求"] == "地质工程"


def test_docx_position_table_loader_preserves_table_and_physical_row(tmp_path) -> None:
    from docx import Document

    document = Document()
    table = document.add_table(rows=2, cols=6)
    headers = ["岗位代码", "招聘单位", "岗位名称", "专业要求", "学历要求", "招聘人数"]
    values = ["D-001", "某地质院", "地质岗", "地质学", "硕士", "1"]
    for index, value in enumerate(headers):
        table.rows[0].cells[index].text = value
    for index, value in enumerate(values):
        table.rows[1].cells[index].text = value
    path = tmp_path / "positions.docx"
    document.save(path)

    loaded = load_position_table_with_report(path)

    assert loaded["metadata"] == {
        "parser": "python-docx",
        "table_count": 1,
        "extraction_mode": "tables",
    }
    assert loaded["rows"][0]["__sheet_name"] == "table-1"
    assert loaded["rows"][0]["__row_number"] == 2
    assert loaded["rows"][0]["岗位名称"] == "地质岗"


def test_image_position_table_ocr_is_low_confidence_and_review_only(tmp_path, monkeypatch) -> None:
    from types import SimpleNamespace

    def fake_run(*args, **kwargs):
        return SimpleNamespace(
            stdout=(
                "202602101 地质技术人员 硕士 地质工程\n"
            ).encode("utf-8")
        )

    monkeypatch.setattr("job_hub.government_table_parser.subprocess.run", fake_run)
    path = tmp_path / "positions.png"
    path.write_bytes(b"not-a-real-image")

    loaded = load_position_table_with_report(path)

    assert loaded["metadata"]["parser"] == "tesseract"
    assert loaded["metadata"]["extraction_confidence"] == "low"
    assert "review_only" in loaded["metadata"]["publication_policy"]
    assert loaded["rows"][0]["__extraction_confidence"] == "low"
